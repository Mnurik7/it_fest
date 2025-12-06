from flask import Flask, render_template, request, jsonify, session, redirect, url_for, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import os
from dotenv import load_dotenv
import google.generativeai as genai
import json
from datetime import datetime
import re
import time
import logging
import tempfile
from typing import List, Dict

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)
# Импорт всех модулей ml_models для предзагрузки
try:
    from ml_models.ai_assistant import get_assistant
    from ml_models import (
        FraudDetectionModel,
        AntifraudModel,
        AntifraudModelEnhanced,
        AntiFraudModelWrapper,
        EnhancedAntiFraudModelWrapper,
        get_model,
        get_enhanced_model,
        MLOpsPipeline,
        RealtimeProcessor
    )
    # Проверка доступности библиотек ML
    try:
        import xgboost as xgb
        XGBOOST_AVAILABLE = True
    except ImportError:
        XGBOOST_AVAILABLE = False
    
    try:
        import lightgbm as lgb
        LIGHTGBM_AVAILABLE = True
    except ImportError:
        LIGHTGBM_AVAILABLE = False
    
    ML_MODELS_AVAILABLE = True
except ImportError as e:
    print(f"⚠️  Предупреждение: Не удалось загрузить ml_models: {e}")
    ML_MODELS_AVAILABLE = False
    XGBOOST_AVAILABLE = False
    LIGHTGBM_AVAILABLE = False
    # Создаем заглушки для работы без ml_models
    FraudDetectionModel = None
    get_assistant = None

import pandas as pd
import numpy as np

# Импорт модулей AI-Procure
from procure_models.tender_analyzer import TenderAnalyzer
from procure_models.supplier_matcher import SupplierMatcher
from procure_models.risk_analyzer import RiskAnalyzer
from procure_models.report_generator import ReportGenerator
from procure_models.beginner_support import BeginnerSupport
from procure_models.data_collector import DataCollector
from procure_models.price_analyzer import PriceAnalyzer
from procure_models.metrics_tracker import MetricsTracker
from procure_models.supplier_database import SupplierDatabase
from procure_models.document_parser import DocumentParser

# Импорт новых модулей AI-Procure
try:
    from procure_models.notification_system import NotificationSystem
except ImportError:
    NotificationSystem = None

try:
    from procure_models.cache_manager import CacheManager
except ImportError:
    CacheManager = None

try:
    from procure_models.async_processor import AsyncProcessor
except ImportError:
    AsyncProcessor = None

try:
    from procure_models.batch_processor import BatchProcessor
except ImportError:
    BatchProcessor = None

try:
    from procure_models.metrics_validator import MetricsValidator
except ImportError:
    MetricsValidator = None

try:
    from procure_models.risk_confirmation import RiskConfirmation
except ImportError:
    RiskConfirmation = None

try:
    from procure_models.prevention_tracker import PreventionTracker
except ImportError:
    PreventionTracker = None

# Импорт модулей AI-Scrum Master
from scrum_models.project_manager import ProjectManager
from scrum_models.task_creator import TaskCreator
from scrum_models.task_decomposer import TaskDecomposer
from scrum_models.meeting_assistant import MeetingAssistant
from scrum_models.meeting_integration_manager import MeetingIntegrationManager
from scrum_models.deadline_tracker import DeadlineTracker
from scrum_models.sprint_manager import SprintManager
from scrum_models.analytics_engine import AnalyticsEngine
from scrum_models.integration_manager import IntegrationManager
# Jira интеграция удалена
from scrum_models.metrics_tracker import ScrumMetricsTracker

# Импорт модулей AI-Business Analyst
from business_analyst_models.document_analyzer import DocumentAnalyzer
from business_analyst_models.requirement_extractor import RequirementExtractor
from business_analyst_models.document_generator import DocumentGenerator
from business_analyst_models.structured_dialog import StructuredDialogManager, DialogStage
from business_analyst_models.metrics_tracker import MetricsTracker as BAMetricsTracker
from business_analyst_models.process_improver import ProcessImprover

# Импорт модулей AI-Code Review Assistant
from code_review_models.architecture_loader import ArchitectureLoader
from code_review_models.code_analyzer import CodeAnalyzer
from code_review_models.review_generator import ReviewGenerator
from code_review_models.report_generator import ReportGenerator as CodeReviewReportGenerator
from code_review_models.gitlab_integration import GitLabIntegration
from code_review_models.educational_support import EducationalSupport
from code_review_models.language_detector import LanguageDetector

load_dotenv()

app = Flask(__name__)
# Используем постоянный secret_key для сохранения сессий между перезапусками
# В продакшене это должно быть в переменных окружения
app.secret_key = os.getenv('FLASK_SECRET_KEY', 'helpdesk-ai-support-platform-secret-key-2025')
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['PERMANENT_SESSION_LIFETIME'] = 86400  # 24 часа

# Настройка Gemini API
api_key = os.getenv('GEMINI_API_KEY')
if not api_key:
    raise ValueError("GEMINI_API_KEY не найден в переменных окружения. Создайте файл .env с вашим API ключом.")

genai.configure(api_key=api_key)
model = genai.GenerativeModel(os.getenv('GEMINI_MODEL', 'gemini-2.5-flash'))

# Инициализация модулей AI-Procure
tender_analyzer = TenderAnalyzer(model)
supplier_matcher = SupplierMatcher(model)
risk_analyzer = RiskAnalyzer(model)
report_generator = ReportGenerator(model)
beginner_support = BeginnerSupport()

# Инициализация новых модулей (опционально)
try:
    data_collector = DataCollector() if DataCollector else None
except:
    data_collector = None
    logger.warning("DataCollector недоступен")

try:
    price_analyzer = PriceAnalyzer() if PriceAnalyzer else None
except:
    price_analyzer = None
    logger.warning("PriceAnalyzer недоступен")

try:
    metrics_tracker = MetricsTracker() if MetricsTracker else None
except:
    metrics_tracker = None
    logger.warning("MetricsTracker недоступен")

try:
    supplier_database = SupplierDatabase() if SupplierDatabase else None
except:
    supplier_database = None
    logger.warning("SupplierDatabase недоступен")

try:
    document_parser = DocumentParser() if DocumentParser else None
except:
    document_parser = None
    logger.warning("DocumentParser недоступен")

# Инициализация новых модулей AI-Procure
try:
    notification_system = NotificationSystem() if NotificationSystem else None
except:
    notification_system = None
    logger.warning("NotificationSystem недоступен")

try:
    cache_manager = CacheManager() if CacheManager else None
except:
    cache_manager = None
    logger.warning("CacheManager недоступен")

try:
    async_processor = AsyncProcessor(max_workers=4) if AsyncProcessor else None
except:
    async_processor = None
    logger.warning("AsyncProcessor недоступен")

try:
    batch_processor = BatchProcessor(async_processor) if BatchProcessor and async_processor else None
except:
    batch_processor = None
    logger.warning("BatchProcessor недоступен")

try:
    metrics_validator = MetricsValidator() if MetricsValidator else None
except:
    metrics_validator = None
    logger.warning("MetricsValidator недоступен")

try:
    risk_confirmation = RiskConfirmation() if RiskConfirmation else None
except:
    risk_confirmation = None
    logger.warning("RiskConfirmation недоступен")

try:
    prevention_tracker = PreventionTracker() if PreventionTracker else None
except:
    prevention_tracker = None
    logger.warning("PreventionTracker недоступен")

# Инициализация модулей AI-Scrum Master
project_manager = ProjectManager(model)
task_creator = TaskCreator(model)
task_decomposer = TaskDecomposer(model)
meeting_assistant = MeetingAssistant(model)
deadline_tracker = DeadlineTracker()
sprint_manager = SprintManager(model)
analytics_engine = AnalyticsEngine(model)
integration_manager = IntegrationManager()
scrum_metrics_tracker = ScrumMetricsTracker()
meeting_integration_manager = MeetingIntegrationManager(model)

# Инициализация модулей AI-Business Analyst
document_analyzer = DocumentAnalyzer(model)
requirement_extractor = RequirementExtractor(model)
document_generator = DocumentGenerator(model)
process_improver = ProcessImprover(model)
ba_metrics_tracker = BAMetricsTracker()

try:
    from business_analyst_models.autocomplete_helper import AutocompleteHelper
    autocomplete_helper = AutocompleteHelper(model)
except ImportError:
    autocomplete_helper = None

# Словарь для хранения активных диалогов (в продакшене лучше использовать Redis)
active_dialogs = {}

# Инициализация модулей AI-Code Review Assistant
language_detector = LanguageDetector()
architecture_loader = ArchitectureLoader(model)
code_analyzer = CodeAnalyzer(model)
review_generator = ReviewGenerator(model)
code_review_report_generator = CodeReviewReportGenerator(model)
gitlab_integration = GitLabIntegration()
educational_support = EducationalSupport(model)

# JSON база данных пользователей
DB_FILE = 'data/users.json'

def load_users():
    """Загрузить пользователей из JSON файла"""
    if not os.path.exists(DB_FILE):
        # Создать директорию если не существует
        os.makedirs(os.path.dirname(DB_FILE), exist_ok=True)
        # Инициализировать базу с тестовыми пользователями
        initial_users = [
            {
                "username": "admin",
                "password_hash": generate_password_hash('admin123'),
                "role": "admin",
                "created_at": "2024-01-01T00:00:00"
            },
            {
                "username": "user",
                "password_hash": generate_password_hash('user123'),
                "role": "user",
                "created_at": "2024-01-01T00:00:00"
            }
        ]
        save_users(initial_users)
        return initial_users
    
    try:
        with open(DB_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except (json.JSONDecodeError, FileNotFoundError):
        return []

def save_users(users_data):
    """Сохранить пользователей в JSON файл"""
    os.makedirs(os.path.dirname(DB_FILE), exist_ok=True)
    with open(DB_FILE, 'w', encoding='utf-8') as f:
        json.dump(users_data, f, ensure_ascii=False, indent=2)

def get_user_by_username(username):
    """Получить пользователя по имени"""
    users = load_users()
    for user in users:
        if user.get('username') == username:
            return user
    return None

def verify_password(username, password):
    """Проверить пароль пользователя"""
    user = get_user_by_username(username)
    if not user:
        return False
    return check_password_hash(user.get('password_hash'), password)

@app.route('/')
def index():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('index.html', username=session.get('username'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        
        if verify_password(username, password):
            user = get_user_by_username(username)
            session.permanent = True  # Делаем сессию постоянной
            session['username'] = username
            session['role'] = user.get('role', 'user')
            return redirect(url_for('index'))
        else:
            return render_template('login.html', error='Неверное имя пользователя или пароль')
    
    if 'username' in session:
        return redirect(url_for('index'))
    
    return render_template('login.html')

@app.route('/login/operator', methods=['GET', 'POST'])
def operator_login():
    """Вход для операторов"""
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        operator_type = request.form.get('operator_type')
        
        # Проверяем учетные данные оператора
        if username == 'admin' and password == 'admin123' and operator_type:
            session.permanent = True
            session['username'] = username
            session['role'] = 'operator'
            session['operator_type'] = operator_type
            return redirect(url_for('operator_dashboard'))
        else:
            return render_template('operator_login.html', error='Неверные учетные данные или не выбран тип оператора')
    
    if 'username' in session and session.get('role') == 'operator':
        return redirect(url_for('operator_dashboard'))
    
    return render_template('operator_login.html')

@app.route('/logout')
def logout():
    session.pop('username', None)
    session.pop('role', None)
    session.pop('operator_type', None)
    if session.get('role') == 'operator':
        return redirect(url_for('operator_login'))
    return redirect(url_for('login'))

@app.route('/ai-procure')
def ai_procure():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('ai_procure.html', username=session.get('username'))

@app.route('/ai-scrum')
def ai_scrum():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('ai_scrum.html', username=session.get('username'))

@app.route('/ai-business-analyst')
def ai_business_analyst():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('ai_business_analyst.html', username=session.get('username'))

@app.route('/confluence-docs')
def confluence_docs():
    """Страница для просмотра сохраненных документов Confluence"""
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('confluence_docs.html', username=session.get('username'))

@app.route('/ai-code-review')
def ai_code_review():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('ai_code_review.html', username=session.get('username'))

@app.route('/it-help-desk')
def it_help_desk():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('it_help_desk.html', username=session.get('username'))

@app.route('/retail-help-desk')
def retail_help_desk():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('retail_help_desk.html', username=session.get('username'))

@app.route('/medical-help-desk')
def medical_help_desk():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('medical_help_desk.html', username=session.get('username'))

@app.route('/legal-help-desk')
def legal_help_desk():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('legal_help_desk.html', username=session.get('username'))

@app.route('/telecom-help-desk')
def telecom_help_desk():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('telecom_help_desk.html', username=session.get('username'))

@app.route('/auto-help-desk')
def auto_help_desk():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('auto_help_desk.html', username=session.get('username'))

@app.route('/help-desk')
def general_help_desk():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('general_help_desk.html', username=session.get('username'))

@app.route('/profile')
def profile():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('profile.html', username=session.get('username'))

@app.route('/operator-chat')
def operator_chat():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('operator_chat.html', username=session.get('username'))

@app.route('/operator/dashboard')
def operator_dashboard():
    """Панель оператора"""
    if 'username' not in session or session.get('role') != 'operator':
        return redirect(url_for('operator_login'))
    return render_template('operator_dashboard.html', 
                         username=session.get('username'),
                         operator_type=session.get('operator_type'))

@app.route('/fraud-detection')
def fraud_detection():
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('fraud_detection.html', username=session.get('username'))

# IT Help Desk API
@app.route('/api/it-help-desk/chat', methods=['POST'])
def api_it_help_desk_chat():
    """API для обработки сообщений IT Help Desk чата"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Системный промпт для IT Help Desk
        system_prompt = """Ты - профессиональный IT-специалист и эксперт по технической поддержке. 
Твоя задача - помогать пользователям решать IT-проблемы, отвечать на технические вопросы и предоставлять квалифицированную помощь.

Отвечай на русском и казахском языках (в зависимости от языка вопроса пользователя).
Будь дружелюбным, профессиональным и точным.
Давай конкретные и практичные решения.
Если проблема требует code review или работы в команде, предложи перейти на соответствующие страницы.

Темы, о которых ты можешь говорить:
- Решение технических проблем с оборудованием
- Настройка программного обеспечения
- Сетевые проблемы и подключение
- Безопасность IT-инфраструктуры
- Обслуживание и обновление систем
- Резервное копирование и восстановление данных
- Устранение неполадок
- Консультации по IT-решениям
- Интеграция систем
- Облачные сервисы

Если пользователь спрашивает о:
- Ревью кода, проверке кода, анализе кода → предложи перейти на "Code Review Assistant"
- Управлении проектами, спринтах, задачах, встречах → предложи перейти на "AI Scrum Master"
- Работе в команде, планировании → предложи перейти на "AI Scrum Master"

ВАЖНО: Если ты не можешь решить проблему пользователя или проблема требует вмешательства оператора/специалиста, обязательно предложи создать заявку оператору. 
В этом случае скажи: "К сожалению, я не могу решить эту проблему автоматически. Предлагаю создать заявку для оператора IT Help Desk, который свяжется с вами для решения проблемы."

Отвечай кратко, но информативно. Используй эмодзи для лучшего восприятия."""
        
        # Формируем контекст из истории чата
        conversation_context = ""
        if chat_history:
            conversation_context = "\n\nИстория диалога:\n"
            for msg in chat_history[-10:]:  # Последние 10 сообщений
                role = "Пользователь" if msg.get('role') == 'user' else "Ассистент"
                content = msg.get('content', '')
                conversation_context += f"{role}: {content}\n"
        
        # Формируем полный промпт
        full_prompt = f"{system_prompt}{conversation_context}\n\nТекущий вопрос пользователя: {user_message}\n\nОтвет:"
        
        # Инициализация Gemini
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'response': 'Извините, AI сервис временно недоступен. Пожалуйста, обратитесь к IT-специалисту.'
            }), 500
        
        genai.configure(api_key=api_key)
        
        # Создаем модель
        model_name = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
        chat_model = genai.GenerativeModel(model_name)
        
        # Генерируем ответ
        try:
            response = chat_model.generate_content(
                full_prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.7,
                    top_p=0.8,
                    top_k=40,
                    max_output_tokens=2048,
                )
            )
            ai_response = response.text.strip() if response.text else 'Извините, не удалось получить ответ.'
        except Exception as e:
            logger.error(f"Ошибка при генерации ответа: {e}")
            ai_response = f"Извините, произошла ошибка при обработке запроса. Пожалуйста, попробуйте еще раз или обратитесь к IT-специалисту."
        
        # Проверяем, нужно ли предложить создание заявки
        suggest_ticket = False
        lower_response = ai_response.lower()
        if any(keyword in lower_response for keyword in [
            'не могу решить', 'не могу помочь', 'создать заявку', 
            'оператор', 'специалист', 'не могу решить эту проблему',
            'требуется вмешательство', 'нужна помощь оператора',
            'к сожалению, я не могу'
        ]):
            suggest_ticket = True
        
        return jsonify({
            'response': ai_response,
            'timestamp': datetime.now().isoformat(),
            'suggest_ticket': suggest_ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка в API IT Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'response': 'Извините, произошла ошибка. Пожалуйста, попробуйте еще раз.'
        }), 500

# Retail Help Desk API
@app.route('/api/retail-help-desk/chat', methods=['POST'])
def api_retail_help_desk_chat():
    """API для обработки сообщений Retail Help Desk чата"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Системный промпт для Retail Help Desk
        system_prompt = """Ты - профессиональный специалист по розничной торговле и поддержке клиентов. 
Твоя задача - помогать клиентам с вопросами о товарах, обработкой жалоб, возвратами и поддержкой продаж.

Отвечай на русском и казахском языках (в зависимости от языка вопроса пользователя).
Будь дружелюбным, профессиональным и вежливым.
Давай конкретные и практичные решения.
Всегда стремись помочь клиенту решить проблему.

Темы, о которых ты можешь говорить:
- Обработка жалоб клиентов
- Возвраты товаров и денежные средства
- Консультации по товарам (характеристики, наличие, применение)
- Поддержка продаж (помощь в выборе товара)
- Информация о доставке и сроках
- Гарантийное обслуживание
- Программы лояльности и скидки
- Работа с претензиями
- Консультации по использованию товаров
- Помощь в оформлении заказов

ВАЖНО: Если ты не можешь решить проблему пользователя или проблема требует вмешательства оператора/специалиста, обязательно предложи создать заявку оператору. 
В этом случае скажи: "К сожалению, я не могу решить эту проблему автоматически. Предлагаю создать заявку для оператора Retail Help Desk, который свяжется с вами для решения проблемы."

Отвечай кратко, но информативно. Используй эмодзи для лучшего восприятия."""
        
        # Формируем контекст из истории чата
        conversation_context = ""
        if chat_history:
            conversation_context = "\n\nИстория диалога:\n"
            for msg in chat_history[-10:]:  # Последние 10 сообщений
                role = "Пользователь" if msg.get('role') == 'user' else "Ассистент"
                content = msg.get('content', '')
                conversation_context += f"{role}: {content}\n"
        
        # Формируем полный промпт
        full_prompt = f"{system_prompt}{conversation_context}\n\nТекущий вопрос пользователя: {user_message}\n\nОтвет:"
        
        # Инициализация Gemini
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'response': 'Извините, AI сервис временно недоступен. Пожалуйста, обратитесь к оператору.'
            }), 500
        
        genai.configure(api_key=api_key)
        
        # Создаем модель
        model_name = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
        chat_model = genai.GenerativeModel(model_name)
        
        # Генерируем ответ
        try:
            response = chat_model.generate_content(
                full_prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.7,
                    top_p=0.8,
                    top_k=40,
                    max_output_tokens=2048,
                )
            )
            ai_response = response.text.strip() if response.text else 'Извините, не удалось получить ответ.'
        except Exception as e:
            logger.error(f"Ошибка при генерации ответа: {e}")
            ai_response = f"Извините, произошла ошибка при обработке запроса. Пожалуйста, попробуйте еще раз или обратитесь к оператору."
        
        # Проверяем, нужно ли предложить создание заявки
        suggest_ticket = False
        lower_response = ai_response.lower()
        if any(keyword in lower_response for keyword in [
            'не могу решить', 'не могу помочь', 'создать заявку', 
            'оператор', 'специалист', 'не могу решить эту проблему',
            'требуется вмешательство', 'нужна помощь оператора',
            'к сожалению, я не могу'
        ]):
            suggest_ticket = True
        
        return jsonify({
            'response': ai_response,
            'timestamp': datetime.now().isoformat(),
            'suggest_ticket': suggest_ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка в API Retail Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'response': 'Извините, произошла ошибка. Пожалуйста, попробуйте еще раз.'
        }), 500

@app.route('/api/retail-help-desk/create-ticket', methods=['POST'])
def api_retail_help_desk_create_ticket():
    """API для создания заявки Retail Help Desk"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        user_name = session.get('username', 'Unknown')
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Путь к JSON файлу
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        os.makedirs(data_dir, exist_ok=True)
        tickets_file = os.path.join(data_dir, 'retail_help_desk_tickets.json')
        
        # Читаем существующие заявки
        if os.path.exists(tickets_file):
            with open(tickets_file, 'r', encoding='utf-8') as f:
                tickets_data = json.load(f)
        else:
            tickets_data = {
                "tickets": [],
                "last_id": 0,
                "departments": {
                    "retail_help_desk": {
                        "name": "Retail Help Desk",
                        "description": "Отдел поддержки розничных клиентов"
                    }
                }
            }
        
        # Создаем новую заявку
        ticket_id = tickets_data['last_id'] + 1
        ticket = {
            "id": ticket_id,
            "user": user_name,
            "message": user_message,
            "chat_history": chat_history[-5:] if chat_history else [],  # Последние 5 сообщений
            "department": "retail_help_desk",
            "status": "open",
            "priority": "medium",
            "created_at": datetime.now().isoformat(),
            "updated_at": datetime.now().isoformat(),
            "assigned_to": None,
            "resolved_at": None
        }
        
        tickets_data['tickets'].append(ticket)
        tickets_data['last_id'] = ticket_id
        
        # Сохраняем в файл
        with open(tickets_file, 'w', encoding='utf-8') as f:
            json.dump(tickets_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Создана заявка Retail Help Desk #{ticket_id} от пользователя {user_name}")
        
        return jsonify({
            'success': True,
            'ticket_id': ticket_id,
            'message': f'Заявка #{ticket_id} успешно создана и отправлена в отдел Retail Help Desk. Оператор свяжется с вами в ближайшее время.',
            'ticket': ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка при создании заявки Retail Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при создании заявки. Пожалуйста, попробуйте еще раз.'
        }), 500

# Medical Help Desk API
@app.route('/api/medical-help-desk/chat', methods=['POST'])
def api_medical_help_desk_chat():
    """API для обработки сообщений Medical Help Desk чата"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Системный промпт для Medical Help Desk
        system_prompt = """Ты - профессиональный медицинский специалист и помощник по поддержке пациентов. 
Твоя задача - помогать пациентам с медицинскими вопросами, записью на приём, консультациями и экстренными обращениями.

ВАЖНО: Ты НЕ можешь ставить диагнозы или назначать лечение. Твоя роль - информационная поддержка и направление к специалистам.

Отвечай на русском и казахском языках (в зависимости от языка вопроса пользователя).
Будь дружелюбным, профессиональным, внимательным и сочувствующим.
Давай конкретные и практичные советы.
Всегда подчеркивай важность обращения к врачу при серьезных симптомах.

Темы, о которых ты можешь говорить:
- Помощь пациентам (общая информация о здоровье)
- Запись на приём к врачу
- Консультации по симптомам (информационные, не диагностические)
- Экстренные обращения (когда обращаться за срочной помощью)
- Информация о медицинских услугах
- Подготовка к приёму врача
- Общие вопросы о здоровье
- Информация о лекарствах (общая, не назначение)
- Профилактика заболеваний
- Вопросы о медицинских документах

ВАЖНО: Если ты не можешь решить проблему пользователя или проблема требует вмешательства врача/специалиста, обязательно предложи создать заявку оператору. 
В этом случае скажи: "К сожалению, я не могу решить эту проблему автоматически. Предлагаю создать заявку для оператора Medical Help Desk, который свяжется с вами для решения проблемы."

При экстренных ситуациях (сильная боль, кровотечение, потеря сознания и т.д.) всегда рекомендуй немедленно обратиться в скорую помощь.

Отвечай кратко, но информативно. Используй эмодзи для лучшего восприятия."""
        
        # Формируем контекст из истории чата
        conversation_context = ""
        if chat_history:
            conversation_context = "\n\nИстория диалога:\n"
            for msg in chat_history[-10:]:  # Последние 10 сообщений
                role = "Пользователь" if msg.get('role') == 'user' else "Ассистент"
                content = msg.get('content', '')
                conversation_context += f"{role}: {content}\n"
        
        # Формируем полный промпт
        full_prompt = f"{system_prompt}{conversation_context}\n\nТекущий вопрос пользователя: {user_message}\n\nОтвет:"
        
        # Инициализация Gemini
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'response': 'Извините, AI сервис временно недоступен. Пожалуйста, обратитесь к оператору.'
            }), 500
        
        genai.configure(api_key=api_key)
        
        # Создаем модель
        model_name = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
        chat_model = genai.GenerativeModel(model_name)
        
        # Генерируем ответ
        try:
            response = chat_model.generate_content(
                full_prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.7,
                    top_p=0.8,
                    top_k=40,
                    max_output_tokens=2048,
                )
            )
            ai_response = response.text.strip() if response.text else 'Извините, не удалось получить ответ.'
        except Exception as e:
            logger.error(f"Ошибка при генерации ответа: {e}")
            ai_response = f"Извините, произошла ошибка при обработке запроса. Пожалуйста, попробуйте еще раз или обратитесь к оператору."
        
        # Проверяем, нужно ли предложить создание заявки
        suggest_ticket = False
        lower_response = ai_response.lower()
        if any(keyword in lower_response for keyword in [
            'не могу решить', 'не могу помочь', 'создать заявку', 
            'оператор', 'специалист', 'не могу решить эту проблему',
            'требуется вмешательство', 'нужна помощь оператора',
            'к сожалению, я не могу', 'врач', 'приём'
        ]):
            suggest_ticket = True
        
        return jsonify({
            'response': ai_response,
            'timestamp': datetime.now().isoformat(),
            'suggest_ticket': suggest_ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка в API Medical Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'response': 'Извините, произошла ошибка. Пожалуйста, попробуйте еще раз.'
        }), 500

@app.route('/api/medical-help-desk/create-ticket', methods=['POST'])
def api_medical_help_desk_create_ticket():
    """API для создания заявки Medical Help Desk"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        user_name = session.get('username', 'Unknown')
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Путь к JSON файлу
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        os.makedirs(data_dir, exist_ok=True)
        tickets_file = os.path.join(data_dir, 'medical_help_desk_tickets.json')
        
        # Читаем существующие заявки
        if os.path.exists(tickets_file):
            with open(tickets_file, 'r', encoding='utf-8') as f:
                tickets_data = json.load(f)
        else:
            tickets_data = {
                "tickets": [],
                "last_id": 0,
                "departments": {
                    "medical_help_desk": {
                        "name": "Medical Help Desk",
                        "description": "Отдел поддержки медицинских запросов"
                    }
                }
            }
        
        # Создаем новую заявку
        ticket_id = tickets_data['last_id'] + 1
        ticket = {
            "id": ticket_id,
            "user": user_name,
            "message": user_message,
            "chat_history": chat_history[-5:] if chat_history else [],  # Последние 5 сообщений
            "department": "medical_help_desk",
            "status": "open",
            "priority": "high",  # Медицинские заявки имеют высокий приоритет
            "created_at": datetime.now().isoformat(),
            "updated_at": datetime.now().isoformat(),
            "assigned_to": None,
            "resolved_at": None
        }
        
        tickets_data['tickets'].append(ticket)
        tickets_data['last_id'] = ticket_id
        
        # Сохраняем в файл
        with open(tickets_file, 'w', encoding='utf-8') as f:
            json.dump(tickets_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Создана заявка Medical Help Desk #{ticket_id} от пользователя {user_name}")
        
        return jsonify({
            'success': True,
            'ticket_id': ticket_id,
            'message': f'Заявка #{ticket_id} успешно создана и отправлена в отдел Medical Help Desk. Оператор свяжется с вами в ближайшее время.',
            'ticket': ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка при создании заявки Medical Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при создании заявки. Пожалуйста, попробуйте еще раз.'
        }), 500

# Legal Help Desk API
@app.route('/api/legal-help-desk/chat', methods=['POST'])
def api_legal_help_desk_chat():
    """API для обработки сообщений Legal Help Desk чата"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Системный промпт для Legal Help Desk
        system_prompt = """Ты - профессиональный юридический консультант и помощник. 
Твоя задача - помогать пользователям с юридическими вопросами, консультациями, обработкой правовых запросов, документооборотом и срочными юридическими вопросами.

ВАЖНО: Ты предоставляешь общую юридическую информацию и консультации, но не можешь заменять профессиональную юридическую помощь адвоката. Всегда подчеркивай необходимость консультации с квалифицированным юристом для конкретных дел.

Отвечай на русском и казахском языках (в зависимости от языка вопроса пользователя).
Будь дружелюбным, профессиональным, точным и внимательным.
Давай полезную юридическую информацию, но всегда подчеркивай необходимость консультации с юристом для конкретных ситуаций.
В срочных юридических вопросах направляй к специалистам.

Темы, о которых ты можешь говорить:
- Общие юридические консультации
- Обработка правовых запросов
- Документооборот и оформление документов
- Срочные юридические вопросы
- Информация о правах и обязанностях
- Консультации по договорам (общая информация)
- Трудовое право (общая информация)
- Гражданское право (общая информация)
- Семейное право (общая информация)
- Корпоративное право (общая информация)
- Налоговое право (общая информация)
- Помощь в подготовке документов (общая информация)

ВАЖНО: Если ты не можешь решить проблему пользователя или проблема требует вмешательства юриста/оператора, обязательно предложи создать заявку оператору. 
В этом случае скажи: "К сожалению, я не могу решить эту проблему автоматически. Предлагаю создать заявку для оператора Legal Help Desk, который свяжется с вами для решения проблемы."

В срочных юридических ситуациях (угроза прав, срочные сроки) направляй к специалистам: "Это срочная юридическая ситуация! Рекомендую немедленно обратиться к квалифицированному юристу или в Legal Help Desk для получения профессиональной помощи."

Отвечай кратко, но информативно. Используй эмодзи для лучшего восприятия."""
        
        # Формируем контекст из истории чата
        conversation_context = ""
        if chat_history:
            conversation_context = "\n\nИстория диалога:\n"
            for msg in chat_history[-10:]:  # Последние 10 сообщений
                role = "Пользователь" if msg.get('role') == 'user' else "Ассистент"
                content = msg.get('content', '')
                conversation_context += f"{role}: {content}\n"
        
        # Формируем полный промпт
        full_prompt = f"{system_prompt}{conversation_context}\n\nТекущий вопрос пользователя: {user_message}\n\nОтвет:"
        
        # Инициализация Gemini
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'response': 'Извините, AI сервис временно недоступен. Пожалуйста, обратитесь к оператору.'
            }), 500
        
        genai.configure(api_key=api_key)
        
        # Создаем модель
        model_name = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
        chat_model = genai.GenerativeModel(model_name)
        
        # Генерируем ответ
        try:
            response = chat_model.generate_content(
                full_prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.7,
                    top_p=0.8,
                    top_k=40,
                    max_output_tokens=2048,
                )
            )
            ai_response = response.text.strip() if response.text else 'Извините, не удалось получить ответ.'
        except Exception as e:
            logger.error(f"Ошибка при генерации ответа: {e}")
            ai_response = f"Извините, произошла ошибка при обработке запроса. Пожалуйста, попробуйте еще раз или обратитесь к оператору."
        
        # Проверяем, нужно ли предложить создание заявки
        suggest_ticket = False
        lower_response = ai_response.lower()
        if any(keyword in lower_response for keyword in [
            'не могу решить', 'не могу помочь', 'создать заявку', 
            'оператор', 'специалист', 'не могу решить эту проблему',
            'требуется вмешательство', 'нужна помощь оператора',
            'к сожалению, я не могу', 'юрист', 'консультация юриста',
            'квалифицированный юрист'
        ]):
            suggest_ticket = True
        
        return jsonify({
            'response': ai_response,
            'timestamp': datetime.now().isoformat(),
            'suggest_ticket': suggest_ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка в API Legal Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'response': 'Извините, произошла ошибка. Пожалуйста, попробуйте еще раз.'
        }), 500

@app.route('/api/legal-help-desk/create-ticket', methods=['POST'])
def api_legal_help_desk_create_ticket():
    """API для создания заявки Legal Help Desk"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        user_name = session.get('username', 'Unknown')
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Путь к JSON файлу
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        os.makedirs(data_dir, exist_ok=True)
        tickets_file = os.path.join(data_dir, 'legal_help_desk_tickets.json')
        
        # Читаем существующие заявки
        if os.path.exists(tickets_file):
            with open(tickets_file, 'r', encoding='utf-8') as f:
                tickets_data = json.load(f)
        else:
            tickets_data = {
                "tickets": [],
                "last_id": 0,
                "departments": {
                    "legal_help_desk": {
                        "name": "Legal Help Desk",
                        "description": "Отдел юридической поддержки"
                    }
                }
            }
        
        # Создаем новую заявку
        ticket_id = tickets_data['last_id'] + 1
        ticket = {
            "id": ticket_id,
            "user": user_name,
            "message": user_message,
            "chat_history": chat_history[-5:] if chat_history else [],  # Последние 5 сообщений
            "department": "legal_help_desk",
            "status": "open",
            "priority": "high",  # Юридические заявки имеют высокий приоритет
            "created_at": datetime.now().isoformat(),
            "updated_at": datetime.now().isoformat(),
            "assigned_to": None,
            "resolved_at": None
        }
        
        tickets_data['tickets'].append(ticket)
        tickets_data['last_id'] = ticket_id
        
        # Сохраняем в файл
        with open(tickets_file, 'w', encoding='utf-8') as f:
            json.dump(tickets_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Создана заявка Legal Help Desk #{ticket_id} от пользователя {user_name}")
        
        return jsonify({
            'success': True,
            'ticket_id': ticket_id,
            'message': f'Заявка #{ticket_id} успешно создана и отправлена в отдел Legal Help Desk. Оператор свяжется с вами в ближайшее время.',
            'ticket': ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка при создании заявки Legal Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при создании заявки. Пожалуйста, попробуйте еще раз.'
        }), 500

# Telecom Help Desk API
@app.route('/api/telecom-help-desk/chat', methods=['POST'])
def api_telecom_help_desk_chat():
    """API для обработки сообщений Telecom Help Desk чата"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Системный промпт для Telecom Help Desk
        system_prompt = """Ты - профессиональный специалист по телекоммуникациям и технической поддержке связи. 
Твоя задача - помогать пользователям с проблемами связи, интернета, настройкой оборудования и вопросами биллинга.

Отвечай на русском и казахском языках (в зависимости от языка вопроса пользователя).
Будь дружелюбным, профессиональным и точным.
Давай конкретные и практичные решения.
Всегда стремись помочь пользователю решить проблему.

Темы, о которых ты можешь говорить:
- Проблемы с связью (мобильная связь, стационарная связь)
- Проблемы с интернетом (медленная скорость, обрывы, подключение)
- Настройка оборудования (роутеры, модемы, телефоны)
- Биллинг и вопросы по счетам
- Настройка Wi-Fi
- Проблемы с сигналом
- Консультации по тарифам
- Техническая диагностика
- Обновление оборудования
- Консультации по услугам связи

ВАЖНО: Если ты не можешь решить проблему пользователя или проблема требует вмешательства оператора/специалиста, обязательно предложи создать заявку оператору. 
В этом случае скажи: "К сожалению, я не могу решить эту проблему автоматически. Предлагаю создать заявку для оператора Telecom Help Desk, который свяжется с вами для решения проблемы."

В критических ситуациях (полное отсутствие связи, срочные проблемы) направляй к специалистам: "Это критическая ситуация! Рекомендую немедленно обратиться в Telecom Help Desk для получения срочной помощи."

Отвечай кратко, но информативно. Используй эмодзи для лучшего восприятия."""
        
        # Формируем контекст из истории чата
        conversation_context = ""
        if chat_history:
            conversation_context = "\n\nИстория диалога:\n"
            for msg in chat_history[-10:]:  # Последние 10 сообщений
                role = "Пользователь" if msg.get('role') == 'user' else "Ассистент"
                content = msg.get('content', '')
                conversation_context += f"{role}: {content}\n"
        
        # Формируем полный промпт
        full_prompt = f"{system_prompt}{conversation_context}\n\nТекущий вопрос пользователя: {user_message}\n\nОтвет:"
        
        # Инициализация Gemini
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'response': 'Извините, AI сервис временно недоступен. Пожалуйста, обратитесь к оператору.'
            }), 500
        
        genai.configure(api_key=api_key)
        
        # Создаем модель
        model_name = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
        chat_model = genai.GenerativeModel(model_name)
        
        # Генерируем ответ
        try:
            response = chat_model.generate_content(
                full_prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.7,
                    top_p=0.8,
                    top_k=40,
                    max_output_tokens=2048,
                )
            )
            ai_response = response.text.strip() if response.text else 'Извините, не удалось получить ответ.'
        except Exception as e:
            logger.error(f"Ошибка при генерации ответа: {e}")
            ai_response = f"Извините, произошла ошибка при обработке запроса. Пожалуйста, попробуйте еще раз или обратитесь к оператору."
        
        # Проверяем, нужно ли предложить создание заявки
        suggest_ticket = False
        lower_response = ai_response.lower()
        if any(keyword in lower_response for keyword in [
            'не могу решить', 'не могу помочь', 'создать заявку', 
            'оператор', 'специалист', 'не могу решить эту проблему',
            'требуется вмешательство', 'нужна помощь оператора',
            'к сожалению, я не могу', 'техник', 'настройка оборудования',
            'критическая ситуация'
        ]):
            suggest_ticket = True
        
        return jsonify({
            'response': ai_response,
            'timestamp': datetime.now().isoformat(),
            'suggest_ticket': suggest_ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка в API Telecom Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'response': 'Извините, произошла ошибка. Пожалуйста, попробуйте еще раз.'
        }), 500

@app.route('/api/telecom-help-desk/create-ticket', methods=['POST'])
def api_telecom_help_desk_create_ticket():
    """API для создания заявки Telecom Help Desk"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        user_name = session.get('username', 'Unknown')
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Путь к JSON файлу
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        os.makedirs(data_dir, exist_ok=True)
        tickets_file = os.path.join(data_dir, 'telecom_help_desk_tickets.json')
        
        # Читаем существующие заявки
        if os.path.exists(tickets_file):
            with open(tickets_file, 'r', encoding='utf-8') as f:
                tickets_data = json.load(f)
        else:
            tickets_data = {
                "tickets": [],
                "last_id": 0,
                "departments": {
                    "telecom_help_desk": {
                        "name": "Telecom Help Desk",
                        "description": "Отдел телекоммуникационной поддержки"
                    }
                }
            }
        
        # Создаем новую заявку
        ticket_id = tickets_data['last_id'] + 1
        ticket = {
            "id": ticket_id,
            "user": user_name,
            "message": user_message,
            "chat_history": chat_history[-5:] if chat_history else [],  # Последние 5 сообщений
            "department": "telecom_help_desk",
            "status": "open",
            "priority": "medium",
            "created_at": datetime.now().isoformat(),
            "updated_at": datetime.now().isoformat(),
            "assigned_to": None,
            "resolved_at": None
        }
        
        tickets_data['tickets'].append(ticket)
        tickets_data['last_id'] = ticket_id
        
        # Сохраняем в файл
        with open(tickets_file, 'w', encoding='utf-8') as f:
            json.dump(tickets_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Создана заявка Telecom Help Desk #{ticket_id} от пользователя {user_name}")
        
        return jsonify({
            'success': True,
            'ticket_id': ticket_id,
            'message': f'Заявка #{ticket_id} успешно создана и отправлена в отдел Telecom Help Desk. Оператор свяжется с вами в ближайшее время.',
            'ticket': ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка при создании заявки Telecom Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при создании заявки. Пожалуйста, попробуйте еще раз.'
        }), 500

# Auto Help Desk API
@app.route('/api/auto-help-desk/chat', methods=['POST'])
def api_auto_help_desk_chat():
    """API для обработки сообщений Auto Help Desk чата"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Системный промпт для Auto Help Desk
        system_prompt = """Ты - профессиональный автомобильный консультант и помощник. 
Твоя задача - помогать пользователям с технической помощью, записью на сервис, консультациями по ремонту и экстренной помощью.

Отвечай на русском и казахском языках (в зависимости от языка вопроса пользователя).
Будь дружелюбным, профессиональным и точным.
Давай конкретные и практичные решения.
Всегда стремись помочь пользователю решить проблему.

Темы, о которых ты можешь говорить:
- Техническая помощь (диагностика проблем, советы по эксплуатации)
- Запись на сервис (информация о сервисных центрах, доступное время)
- Консультации по ремонту (общая информация о ремонте, рекомендации)
- Экстренная помощь (что делать в экстренных ситуациях на дороге)
- Обслуживание автомобиля (ТО, замена масла, фильтров)
- Диагностика проблем (помощь в определении неисправностей)
- Консультации по запчастям (общая информация)
- Советы по эксплуатации (как правильно использовать автомобиль)
- Информация о гарантии (общая информация)
- Консультации по выбору сервиса

ВАЖНО: Если ты не можешь решить проблему пользователя или проблема требует вмешательства механика/оператора, обязательно предложи создать заявку оператору. 
В этом случае скажи: "К сожалению, я не могу решить эту проблему автоматически. Предлагаю создать заявку для оператора Auto Help Desk, который свяжется с вами для решения проблемы."

В экстренных ситуациях (авария, поломка на дороге) направляй к экстренным службам: "Это экстренная ситуация! Немедленно вызовите эвакуатор или экстренные службы. Также можете создать заявку в Auto Help Desk для получения помощи."

Отвечай кратко, но информативно. Используй эмодзи для лучшего восприятия."""
        
        # Формируем контекст из истории чата
        conversation_context = ""
        if chat_history:
            conversation_context = "\n\nИстория диалога:\n"
            for msg in chat_history[-10:]:  # Последние 10 сообщений
                role = "Пользователь" if msg.get('role') == 'user' else "Ассистент"
                content = msg.get('content', '')
                conversation_context += f"{role}: {content}\n"
        
        # Формируем полный промпт
        full_prompt = f"{system_prompt}{conversation_context}\n\nТекущий вопрос пользователя: {user_message}\n\nОтвет:"
        
        # Инициализация Gemini
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'response': 'Извините, AI сервис временно недоступен. Пожалуйста, обратитесь к оператору.'
            }), 500
        
        genai.configure(api_key=api_key)
        
        # Создаем модель
        model_name = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
        chat_model = genai.GenerativeModel(model_name)
        
        # Генерируем ответ
        try:
            response = chat_model.generate_content(
                full_prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.7,
                    top_p=0.8,
                    top_k=40,
                    max_output_tokens=2048,
                )
            )
            ai_response = response.text.strip() if response.text else 'Извините, не удалось получить ответ.'
        except Exception as e:
            logger.error(f"Ошибка при генерации ответа: {e}")
            ai_response = f"Извините, произошла ошибка при обработке запроса. Пожалуйста, попробуйте еще раз или обратитесь к оператору."
        
        # Проверяем, нужно ли предложить создание заявки
        suggest_ticket = False
        lower_response = ai_response.lower()
        if any(keyword in lower_response for keyword in [
            'не могу решить', 'не могу помочь', 'создать заявку', 
            'оператор', 'специалист', 'не могу решить эту проблему',
            'требуется вмешательство', 'нужна помощь оператора',
            'к сожалению, я не могу', 'механик', 'сервис',
            'экстренная ситуация'
        ]):
            suggest_ticket = True
        
        return jsonify({
            'response': ai_response,
            'timestamp': datetime.now().isoformat(),
            'suggest_ticket': suggest_ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка в API Auto Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'response': 'Извините, произошла ошибка. Пожалуйста, попробуйте еще раз.'
        }), 500

@app.route('/api/auto-help-desk/create-ticket', methods=['POST'])
def api_auto_help_desk_create_ticket():
    """API для создания заявки Auto Help Desk"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        user_name = session.get('username', 'Unknown')
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Путь к JSON файлу
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        os.makedirs(data_dir, exist_ok=True)
        tickets_file = os.path.join(data_dir, 'auto_help_desk_tickets.json')
        
        # Читаем существующие заявки
        if os.path.exists(tickets_file):
            with open(tickets_file, 'r', encoding='utf-8') as f:
                tickets_data = json.load(f)
        else:
            tickets_data = {
                "tickets": [],
                "last_id": 0,
                "departments": {
                    "auto_help_desk": {
                        "name": "Auto Help Desk",
                        "description": "Отдел автомобильной поддержки"
                    }
                }
            }
        
        # Создаем новую заявку
        ticket_id = tickets_data['last_id'] + 1
        ticket = {
            "id": ticket_id,
            "user": user_name,
            "message": user_message,
            "chat_history": chat_history[-5:] if chat_history else [],  # Последние 5 сообщений
            "department": "auto_help_desk",
            "status": "open",
            "priority": "medium",
            "created_at": datetime.now().isoformat(),
            "updated_at": datetime.now().isoformat(),
            "assigned_to": None,
            "resolved_at": None
        }
        
        tickets_data['tickets'].append(ticket)
        tickets_data['last_id'] = ticket_id
        
        # Сохраняем в файл
        with open(tickets_file, 'w', encoding='utf-8') as f:
            json.dump(tickets_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Создана заявка Auto Help Desk #{ticket_id} от пользователя {user_name}")
        
        return jsonify({
            'success': True,
            'ticket_id': ticket_id,
            'message': f'Заявка #{ticket_id} успешно создана и отправлена в отдел Auto Help Desk. Оператор свяжется с вами в ближайшее время.',
            'ticket': ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка при создании заявки Auto Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при создании заявки. Пожалуйста, попробуйте еще раз.'
        }), 500

# General Help Desk API - Универсальный Help Desk с определением направления
@app.route('/api/help-desk/chat', methods=['POST'])
def api_general_help_desk_chat():
    """API для обработки сообщений общего Help Desk чата с определением направления"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Системный промпт для общего Help Desk
        system_prompt = """Ты - универсальный AI-ассистент Help Desk, который помогает пользователям определить, в каком направлении им нужна помощь, и направляет их к нужному специализированному Help Desk или помогает напрямую.

Твоя задача:
1. Анализировать запрос пользователя
2. Определять, к какому направлению относится вопрос (IT, Retail, Medical, Legal, Telecom, Auto)
3. Либо помогать напрямую, либо направлять в соответствующий Help Desk
4. Если не можешь помочь - предлагать создать заявку оператору

Доступные направления:
- IT Help Desk: технические проблемы, оборудование, ПО, сети, безопасность
- Retail Help Desk: жалобы клиентов, возвраты, консультации по товарам, продажи
- Medical Help Desk: медицинские вопросы, запись на приём, консультации, экстренные обращения
- Legal Help Desk: юридические консультации, правовые запросы, документооборот
- Telecom Help Desk: проблемы с связью, интернетом, настройка оборудования, биллинг
- Auto Help Desk: техническая помощь, запись на сервис, консультации по ремонту, экстренная помощь

Отвечай на русском и казахском языках (в зависимости от языка вопроса пользователя).
Будь дружелюбным, профессиональным и точным.

ВАЖНО: 
- Если вопрос явно относится к одному из направлений, предложи перейти в соответствующий Help Desk и дай краткую помощь
- Если вопрос общий или неясный, помоги напрямую
- Если не можешь помочь, предложи создать заявку оператору

Формат ответа:
- Если нужно направить в специализированный Help Desk, скажи: "Ваш вопрос относится к [направление]. Рекомендую перейти в [название] Help Desk для получения специализированной помощи. [Краткая помощь]"
- Если помогаешь напрямую, просто помоги
- Если не можешь помочь, скажи: "К сожалению, я не могу решить эту проблему автоматически. Предлагаю создать заявку для оператора Help Desk, который свяжется с вами для решения проблемы."

Отвечай кратко, но информативно. Используй эмодзи для лучшего восприятия."""
        
        # Формируем контекст из истории чата
        conversation_context = ""
        if chat_history:
            conversation_context = "\n\nИстория диалога:\n"
            for msg in chat_history[-10:]:  # Последние 10 сообщений
                role = "Пользователь" if msg.get('role') == 'user' else "Ассистент"
                content = msg.get('content', '')
                conversation_context += f"{role}: {content}\n"
        
        # Формируем полный промпт
        full_prompt = f"{system_prompt}{conversation_context}\n\nТекущий вопрос пользователя: {user_message}\n\nОтвет:"
        
        # Инициализация Gemini
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'response': 'Извините, AI сервис временно недоступен. Пожалуйста, обратитесь к оператору.'
            }), 500
        
        genai.configure(api_key=api_key)
        
        # Создаем модель
        model_name = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
        chat_model = genai.GenerativeModel(model_name)
        
        # Генерируем ответ
        try:
            response = chat_model.generate_content(
                full_prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.7,
                    top_p=0.8,
                    top_k=40,
                    max_output_tokens=2048,
                )
            )
            ai_response = response.text.strip() if response.text else 'Извините, не удалось получить ответ.'
        except Exception as e:
            logger.error(f"Ошибка при генерации ответа: {e}")
            ai_response = f"Извините, произошла ошибка при обработке запроса. Пожалуйста, попробуйте еще раз или обратитесь к оператору."
        
        # Определяем, нужно ли предложить создание заявки или направление
        suggest_ticket = False
        suggested_department = None
        
        lower_response = ai_response.lower()
        
        # Проверяем упоминания направлений
        department_keywords = {
            'it_help_desk': ['it help desk', 'it-помощь', 'техническая поддержка', 'it поддержка'],
            'retail_help_desk': ['retail help desk', 'розничная', 'товар', 'возврат', 'продаж'],
            'medical_help_desk': ['medical help desk', 'медицинск', 'врач', 'приём', 'здоров'],
            'legal_help_desk': ['legal help desk', 'юридическ', 'правов', 'документ'],
            'telecom_help_desk': ['telecom help desk', 'телеком', 'связь', 'интернет', 'биллинг'],
            'auto_help_desk': ['auto help desk', 'автомобил', 'сервис', 'ремонт', 'машина']
        }
        
        for dept, keywords in department_keywords.items():
            if any(keyword in lower_response for keyword in keywords):
                suggested_department = dept
                break
        
        # Проверяем, нужно ли предложить создание заявки
        if any(keyword in lower_response for keyword in [
            'не могу решить', 'не могу помочь', 'создать заявку', 
            'оператор', 'специалист', 'не могу решить эту проблему',
            'требуется вмешательство', 'нужна помощь оператора',
            'к сожалению, я не могу'
        ]):
            suggest_ticket = True
        
        return jsonify({
            'response': ai_response,
            'timestamp': datetime.now().isoformat(),
            'suggest_ticket': suggest_ticket,
            'suggested_department': suggested_department
        })
    
    except Exception as e:
        logger.error(f"Ошибка в API General Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'response': 'Извините, произошла ошибка. Пожалуйста, попробуйте еще раз.'
        }), 500

@app.route('/api/help-desk/create-ticket', methods=['POST'])
def api_general_help_desk_create_ticket():
    """API для создания заявки общего Help Desk"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        department = data.get('department', 'general_help_desk')  # По умолчанию общий Help Desk
        user_name = session.get('username', 'Unknown')
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Определяем файл для сохранения заявки
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        os.makedirs(data_dir, exist_ok=True)
        
        # Если указан конкретный отдел, используем его файл, иначе общий
        if department and department != 'general_help_desk':
            tickets_file = os.path.join(data_dir, f'{department}_tickets.json')
        else:
            tickets_file = os.path.join(data_dir, 'general_help_desk_tickets.json')
        
        # Читаем существующие заявки
        if os.path.exists(tickets_file):
            with open(tickets_file, 'r', encoding='utf-8') as f:
                tickets_data = json.load(f)
        else:
            dept_name = department.replace('_', ' ').title() if department else 'General Help Desk'
            tickets_data = {
                "tickets": [],
                "last_id": 0,
                "departments": {
                    department or "general_help_desk": {
                        "name": dept_name,
                        "description": "Отдел общей поддержки"
                    }
                }
            }
        
        # Создаем новую заявку
        ticket_id = tickets_data['last_id'] + 1
        ticket = {
            "id": ticket_id,
            "user": user_name,
            "message": user_message,
            "chat_history": chat_history[-5:] if chat_history else [],  # Последние 5 сообщений
            "department": department or "general_help_desk",
            "status": "open",
            "priority": "medium",
            "created_at": datetime.now().isoformat(),
            "updated_at": datetime.now().isoformat(),
            "assigned_to": None,
            "resolved_at": None
        }
        
        tickets_data['tickets'].append(ticket)
        tickets_data['last_id'] = ticket_id
        
        # Сохраняем в файл
        with open(tickets_file, 'w', encoding='utf-8') as f:
            json.dump(tickets_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Создана заявка General Help Desk #{ticket_id} от пользователя {user_name} в отдел {department}")
        
        return jsonify({
            'success': True,
            'ticket_id': ticket_id,
            'message': f'Заявка #{ticket_id} успешно создана и отправлена в отдел Help Desk. Оператор свяжется с вами в ближайшее время.',
            'ticket': ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка при создании заявки General Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при создании заявки. Пожалуйста, попробуйте еще раз.'
        }), 500

# Profile API - Получение истории чатов и запросов
@app.route('/api/profile/history', methods=['GET'])
def api_profile_history():
    """API для получения истории чатов и запросов пользователя"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        username = session.get('username')
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        
        # Собираем все тикеты пользователя из всех отделов
        all_tickets = []
        ticket_files = [
            'it_help_desk_tickets.json',
            'retail_help_desk_tickets.json',
            'medical_help_desk_tickets.json',
            'legal_help_desk_tickets.json',
            'telecom_help_desk_tickets.json',
            'auto_help_desk_tickets.json',
            'general_help_desk_tickets.json'
        ]
        
        for ticket_file in ticket_files:
            file_path = os.path.join(data_dir, ticket_file)
            if os.path.exists(file_path):
                with open(file_path, 'r', encoding='utf-8') as f:
                    tickets_data = json.load(f)
                    user_tickets = [t for t in tickets_data.get('tickets', []) if t.get('user') == username]
                    all_tickets.extend(user_tickets)
        
        # Сортируем по дате создания (новые первыми)
        all_tickets.sort(key=lambda x: x.get('created_at', ''), reverse=True)
        
        # Собираем историю чатов из JSON файлов
        chat_history_file = os.path.join(data_dir, 'chat_history.json')
        chat_sessions = []
        if os.path.exists(chat_history_file):
            with open(chat_history_file, 'r', encoding='utf-8') as f:
                chat_data = json.load(f)
                user_sessions = [s for s in chat_data.get('sessions', []) if s.get('user') == username]
                chat_sessions = sorted(user_sessions, key=lambda x: x.get('last_updated', ''), reverse=True)
        
        return jsonify({
            'success': True,
            'tickets': all_tickets,
            'chat_sessions': chat_sessions,
            'total_tickets': len(all_tickets),
            'total_chat_sessions': len(chat_sessions)
        })
    
    except Exception as e:
        logger.error(f"Ошибка при получении истории профиля: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при загрузке истории.'
        }), 500

# Operator Chat API - Чат с операторами
@app.route('/api/operator-chat/send', methods=['POST'])
def api_operator_chat_send():
    """API для отправки сообщения оператору"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        ticket_id = data.get('ticket_id')  # Обязательно - привязка к тикету
        department = data.get('department')  # Опционально - для прямого доступа к БД
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        if not ticket_id:
            return jsonify({'error': 'ticket_id обязателен для отправки сообщения'}), 400
        
        # Преобразуем ticket_id в число, если это строка
        try:
            ticket_id = int(ticket_id)
        except (ValueError, TypeError):
            return jsonify({'error': 'Некорректный ticket_id'}), 400
        
        username = session.get('username', 'Unknown')
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        os.makedirs(data_dir, exist_ok=True)
        
        # Маппинг отделов к файлам БД
        department_file_map = {
            'it_help_desk': 'it_help_desk_tickets.json',
            'retail_help_desk': 'retail_help_desk_tickets.json',
            'medical_help_desk': 'medical_help_desk_tickets.json',
            'legal_help_desk': 'legal_help_desk_tickets.json',
            'telecom_help_desk': 'telecom_help_desk_tickets.json',
            'auto_help_desk': 'auto_help_desk_tickets.json',
            'general_help_desk': 'general_help_desk_tickets.json'
        }
        
        # Если указан department, используем его для прямого доступа к нужной БД
        if department and department in department_file_map:
            ticket_file = department_file_map[department]
            file_path = os.path.join(data_dir, ticket_file)
            
            if os.path.exists(file_path):
                with open(file_path, 'r', encoding='utf-8') as f:
                    tickets_data = json.load(f)
                
                # Ищем тикет в текущем файле
                for i, ticket in enumerate(tickets_data.get('tickets', [])):
                    # Сравниваем ID как числа
                    ticket_id_in_file = ticket.get('id')
                    if isinstance(ticket_id_in_file, str):
                        try:
                            ticket_id_in_file = int(ticket_id_in_file)
                        except ValueError:
                            continue
                    
                    if ticket_id_in_file == ticket_id and ticket.get('user') == username:
                        # Добавляем сообщение пользователя в operator_chat_history
                        if 'operator_chat_history' not in ticket:
                            ticket['operator_chat_history'] = []
                        
                        ticket['operator_chat_history'].append({
                            'role': 'user',
                            'content': user_message,
                            'timestamp': datetime.now().isoformat()
                        })
                        
                        ticket['updated_at'] = datetime.now().isoformat()
                        
                        # Обновляем тикет в данных
                        tickets_data['tickets'][i] = ticket
                        
                        # Сохраняем обновленные данные
                        with open(file_path, 'w', encoding='utf-8') as f:
                            json.dump(tickets_data, f, ensure_ascii=False, indent=2)
                        
                        logger.info(f"Сообщение добавлено в тикет #{ticket_id} от пользователя {username} в файл {ticket_file}")
                        
                        return jsonify({
                            'success': True,
                            'message_id': ticket_id,
                            'message': 'Сообщение отправлено оператору. Ответ будет доступен в чате.'
                        })
        
        # Если department не указан или тикет не найден, ищем во всех БД
        ticket_files = list(department_file_map.values())
        
        for ticket_file in ticket_files:
            file_path = os.path.join(data_dir, ticket_file)
            if os.path.exists(file_path):
                with open(file_path, 'r', encoding='utf-8') as f:
                    tickets_data = json.load(f)
                
                # Ищем тикет в текущем файле
                for i, ticket in enumerate(tickets_data.get('tickets', [])):
                    # Сравниваем ID как числа
                    ticket_id_in_file = ticket.get('id')
                    if isinstance(ticket_id_in_file, str):
                        try:
                            ticket_id_in_file = int(ticket_id_in_file)
                        except ValueError:
                            continue
                    
                    if ticket_id_in_file == ticket_id and ticket.get('user') == username:
                        # Добавляем сообщение пользователя в operator_chat_history
                        if 'operator_chat_history' not in ticket:
                            ticket['operator_chat_history'] = []
                        
                        ticket['operator_chat_history'].append({
                            'role': 'user',
                            'content': user_message,
                            'timestamp': datetime.now().isoformat()
                        })
                        
                        ticket['updated_at'] = datetime.now().isoformat()
                        
                        # Обновляем тикет в данных
                        tickets_data['tickets'][i] = ticket
                        
                        # Сохраняем обновленные данные
                        with open(file_path, 'w', encoding='utf-8') as f:
                            json.dump(tickets_data, f, ensure_ascii=False, indent=2)
                        
                        logger.info(f"Сообщение добавлено в тикет #{ticket_id} от пользователя {username} в файл {ticket_file}")
                        
                        return jsonify({
                            'success': True,
                            'message_id': ticket_id,
                            'message': 'Сообщение отправлено оператору. Ответ будет доступен в чате.'
                        })
        
        # Если тикет не найден
        return jsonify({'error': f'Тикет #{ticket_id} не найден или не принадлежит пользователю {username}'}), 404
        
    except Exception as e:
        logger.error(f"Ошибка при отправке сообщения оператору: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при отправке сообщения.'
        }), 500

@app.route('/api/operator-chat/messages', methods=['GET'])
def api_operator_chat_messages():
    """API для получения сообщений чата с операторами"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        username = session.get('username')
        ticket_id = request.args.get('ticket_id', type=int)  # Получаем ticket_id из параметров запроса
        department = request.args.get('department', type=str)  # Получаем department из параметров запроса
        
        if not ticket_id:
            return jsonify({
                'success': True,
                'messages': []
            })
        
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        
        # Маппинг отделов к файлам БД
        department_file_map = {
            'it_help_desk': 'it_help_desk_tickets.json',
            'retail_help_desk': 'retail_help_desk_tickets.json',
            'medical_help_desk': 'medical_help_desk_tickets.json',
            'legal_help_desk': 'legal_help_desk_tickets.json',
            'telecom_help_desk': 'telecom_help_desk_tickets.json',
            'auto_help_desk': 'auto_help_desk_tickets.json',
            'general_help_desk': 'general_help_desk_tickets.json'
        }
        
        # Если указан department, используем его для прямого доступа к нужной БД
        if department and department in department_file_map:
            ticket_file = department_file_map[department]
            file_path = os.path.join(data_dir, ticket_file)
            
            if os.path.exists(file_path):
                with open(file_path, 'r', encoding='utf-8') as f:
                    tickets_data = json.load(f)
                    for ticket in tickets_data.get('tickets', []):
                        if ticket.get('id') == ticket_id and ticket.get('user') == username:
                            # Преобразуем тикет в сообщение для клиента
                            operator_messages = ticket.get('operator_messages', [])
                            operator_chat_history = ticket.get('operator_chat_history', [])
                            
                            message = {
                                'id': ticket.get('id'),
                                'user': username,
                                'message': ticket.get('message', ''),
                                'ticket_id': ticket.get('id'),
                                'department': ticket.get('department'),
                                'type': 'user',
                                'status': 'pending' if not operator_messages else 'answered',
                                'created_at': ticket.get('created_at', datetime.now().isoformat()),
                                'operator_messages': operator_messages,
                                'operator_chat_history': operator_chat_history,
                                'operator_response': operator_messages[-1].get('response') if operator_messages else None,
                                'answered_at': operator_messages[-1].get('answered_at') if operator_messages else None
                            }
                            
                            return jsonify({
                                'success': True,
                                'messages': [message]
                            })
        
        # Если department не указан или не найден, ищем во всех БД
        ticket_files = list(department_file_map.values())
        
        for ticket_file in ticket_files:
            file_path = os.path.join(data_dir, ticket_file)
            if os.path.exists(file_path):
                with open(file_path, 'r', encoding='utf-8') as f:
                    tickets_data = json.load(f)
                    for ticket in tickets_data.get('tickets', []):
                        if ticket.get('id') == ticket_id and ticket.get('user') == username:
                            # Преобразуем тикет в сообщение для клиента
                            operator_messages = ticket.get('operator_messages', [])
                            operator_chat_history = ticket.get('operator_chat_history', [])
                            
                            message = {
                                'id': ticket.get('id'),
                                'user': username,
                                'message': ticket.get('message', ''),
                                'ticket_id': ticket.get('id'),
                                'department': ticket.get('department'),
                                'type': 'user',
                                'status': 'pending' if not operator_messages else 'answered',
                                'created_at': ticket.get('created_at', datetime.now().isoformat()),
                                'operator_messages': operator_messages,
                                'operator_chat_history': operator_chat_history,
                                'operator_response': operator_messages[-1].get('response') if operator_messages else None,
                                'answered_at': operator_messages[-1].get('answered_at') if operator_messages else None
                            }
                            
                            return jsonify({
                                'success': True,
                                'messages': [message]
                            })
        
        # Тикет не найден
        return jsonify({
            'success': True,
            'messages': []
        })
    
    except Exception as e:
        logger.error(f"Ошибка при получении сообщений оператора: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при загрузке сообщений.'
        }), 500

# Operator Dashboard API
@app.route('/api/operator/messages', methods=['GET'])
def api_operator_messages():
    """API для получения всех сообщений для оператора"""
    if 'username' not in session or session.get('role') != 'operator':
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        operator_type = session.get('operator_type')
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        
        messages = []
        
        # 1. Загружаем тикеты из соответствующего JSON файла отдела
        ticket_file_map = {
            'it_help_desk': 'it_help_desk_tickets.json',
            'retail_help_desk': 'retail_help_desk_tickets.json',
            'medical_help_desk': 'medical_help_desk_tickets.json',
            'legal_help_desk': 'legal_help_desk_tickets.json',
            'telecom_help_desk': 'telecom_help_desk_tickets.json',
            'auto_help_desk': 'auto_help_desk_tickets.json',
            'general_help_desk': 'general_help_desk_tickets.json'
        }
        
        # Загружаем тикеты из файла соответствующего отдела
        if operator_type in ticket_file_map:
            ticket_file = os.path.join(data_dir, ticket_file_map[operator_type])
            if os.path.exists(ticket_file):
                with open(ticket_file, 'r', encoding='utf-8') as f:
                    tickets_data = json.load(f)
                    for ticket in tickets_data.get('tickets', []):
                        # Преобразуем тикет в сообщение для оператора
                        # Показываем все тикеты, не только открытые
                        operator_messages = ticket.get('operator_messages', [])
                        # Для обратной совместимости проверяем старое поле operator_response
                        if ticket.get('operator_response') and not operator_messages:
                            operator_messages = [{
                                'response': ticket.get('operator_response'),
                                'operator_name': ticket.get('operator_name', 'admin'),
                                'answered_at': ticket.get('answered_at', datetime.now().isoformat())
                            }]
                        
                        message = {
                            'id': ticket.get('id'),
                            'user': ticket.get('user', 'Unknown'),
                            'message': ticket.get('message', ''),
                            'ticket_id': ticket.get('id'),
                            'type': 'user',
                            'status': 'pending' if not operator_messages else 'answered',
                            'created_at': ticket.get('created_at', datetime.now().isoformat()),
                            'operator_messages': operator_messages,  # Массив ответов
                            'operator_response': operator_messages[-1].get('response') if operator_messages else None,  # Последний ответ для совместимости
                            'operator_name': operator_messages[-1].get('operator_name') if operator_messages else None,
                            'answered_at': operator_messages[-1].get('answered_at') if operator_messages else None,
                            'department': ticket.get('department'),
                            'priority': ticket.get('priority', 'medium'),
                            'chat_history': ticket.get('chat_history', []),  # История с AI
                            'operator_chat_history': ticket.get('operator_chat_history', [])  # История чата с оператором
                        }
                        messages.append(message)
        
        # 2. Также загружаем сообщения из operator_chat.json
        operator_chat_file = os.path.join(data_dir, 'operator_chat.json')
        if os.path.exists(operator_chat_file):
            with open(operator_chat_file, 'r', encoding='utf-8') as f:
                chat_data = json.load(f)
                all_messages = chat_data.get('messages', [])
                
                # Фильтруем по типу оператора
                for msg in all_messages:
                    if msg.get('ticket_id'):
                        # Проверяем отдел тикета
                        ticket = get_ticket_by_id(msg.get('ticket_id'))
                        if ticket and ticket.get('department') == operator_type:
                            # Проверяем, не добавлен ли уже этот тикет
                            if not any(m.get('ticket_id') == msg.get('ticket_id') for m in messages):
                                messages.append(msg)
                    elif operator_type == 'general_help_desk':
                        # General оператор видит все сообщения без тикетов
                        messages.append(msg)
        
        # Сортируем по дате (новые первыми)
        messages.sort(key=lambda x: x.get('created_at', ''), reverse=True)
        
        return jsonify({
            'success': True,
            'messages': messages
        })
    
    except Exception as e:
        logger.error(f"Ошибка при получении сообщений оператора: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при загрузке сообщений.'
        }), 500

@app.route('/api/operator/respond', methods=['POST'])
def api_operator_respond():
    """API для отправки ответа оператора"""
    if 'username' not in session or session.get('role') != 'operator':
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        message_id = data.get('message_id')
        response = data.get('response', '').strip()
        ticket_id = data.get('ticket_id')  # ID тикета, если есть
        
        print(f"[DEBUG] api_operator_respond вызван: message_id={message_id}, ticket_id={ticket_id}, response={response[:50]}")
        logger.info(f"[DEBUG] api_operator_respond вызван: message_id={message_id}, ticket_id={ticket_id}, response={response[:50]}")
        
        if not message_id or not response:
            return jsonify({'error': 'ID сообщения и ответ обязательны'}), 400
        
        username = session.get('username')
        operator_type = session.get('operator_type')
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        os.makedirs(data_dir, exist_ok=True)
        
        updated = False
        
        # Проверяем наличие ticket_id
        if not ticket_id:
            error_msg = f"ticket_id не предоставлен. message_id: {message_id}, response: {response[:50]}"
            print(f"[ERROR] {error_msg}")
            logger.error(error_msg)
            return jsonify({'error': 'ticket_id обязателен для отправки ответа'}), 400
        
        # Получаем тикет по ID
        ticket = get_ticket_by_id(ticket_id)
        if not ticket:
            error_msg = f"Тикет #{ticket_id} не найден в базе данных. message_id: {message_id}"
            print(f"[ERROR] {error_msg}")
            logger.error(error_msg)
            return jsonify({'error': f'Тикет #{ticket_id} не найден'}), 404
        
        # Определяем файл тикета
        ticket_file_map = {
            'it_help_desk': 'it_help_desk_tickets.json',
            'retail_help_desk': 'retail_help_desk_tickets.json',
            'medical_help_desk': 'medical_help_desk_tickets.json',
            'legal_help_desk': 'legal_help_desk_tickets.json',
            'telecom_help_desk': 'telecom_help_desk_tickets.json',
            'auto_help_desk': 'auto_help_desk_tickets.json',
            'general_help_desk': 'general_help_desk_tickets.json'
        }
        
        dept = ticket.get('department')
        info_msg = f"Тикет #{ticket_id} найден. Отдел: {dept}, Пользователь: {ticket.get('user')}, Оператор: {username}, Тип оператора: {operator_type}"
        print(f"[INFO] {info_msg}")
        logger.info(info_msg)
        
        if not dept or dept not in ticket_file_map:
            error_msg = f"Отдел '{dept}' не найден в ticket_file_map для тикета #{ticket_id}"
            print(f"[ERROR] {error_msg}")
            logger.error(error_msg)
            return jsonify({'error': f'Неизвестный отдел: {dept}'}), 400
        
        ticket_file = os.path.join(data_dir, ticket_file_map[dept])
        print(f"[INFO] Файл тикета: {ticket_file}, существует: {os.path.exists(ticket_file)}")
        
        if not os.path.exists(ticket_file):
            error_msg = f"Файл {ticket_file} не существует"
            print(f"[ERROR] {error_msg}")
            logger.error(error_msg)
            return jsonify({'error': f'Файл базы данных для отдела {dept} не найден'}), 404
        
        # Читаем данные из файла
        try:
            with open(ticket_file, 'r', encoding='utf-8') as f:
                tickets_data = json.load(f)
        except Exception as e:
            logger.error(f"Ошибка при чтении файла {ticket_file}: {e}")
            return jsonify({'error': 'Ошибка при чтении базы данных'}), 500
        
        # Преобразуем ticket_id в число для сравнения
        ticket_id_int = ticket_id
        if isinstance(ticket_id, str):
            try:
                ticket_id_int = int(ticket_id)
            except ValueError:
                logger.error(f"Не удалось преобразовать ticket_id '{ticket_id}' в число")
                return jsonify({'error': 'Неверный формат ticket_id'}), 400
        
        # Обновляем тикет
        updated = False
        for i, t in enumerate(tickets_data.get('tickets', [])):
            # Сравниваем ID как числа
            ticket_id_in_file = t.get('id')
            if isinstance(ticket_id_in_file, str):
                try:
                    ticket_id_in_file = int(ticket_id_in_file)
                except ValueError:
                    continue
            
            if ticket_id_in_file == ticket_id_int:
                # Инициализируем массив ответов, если его нет
                if 'operator_messages' not in t:
                    t['operator_messages'] = []
                    # Для обратной совместимости переносим старый ответ в массив
                    if t.get('operator_response'):
                        t['operator_messages'].append({
                            'response': t.get('operator_response'),
                            'operator_name': t.get('operator_name', username),
                            'answered_at': t.get('answered_at', datetime.now().isoformat())
                        })
                
                # Добавляем новый ответ в массив
                t['operator_messages'].append({
                    'response': response,
                    'operator_name': username,
                    'answered_at': datetime.now().isoformat()
                })
                
                # Добавляем ответ в историю чата тикета
                if 'operator_chat_history' not in t:
                    t['operator_chat_history'] = []
                
                new_message = {
                    'role': 'operator',
                    'content': response,
                    'operator_name': username,
                    'timestamp': datetime.now().isoformat()
                }
                t['operator_chat_history'].append(new_message)
                
                info_msg = f"Добавлено сообщение в operator_chat_history тикета #{ticket_id}: role={new_message['role']}, content={response[:50]}..."
                print(f"[INFO] {info_msg}")
                logger.info(info_msg)
                
                # Для обратной совместимости обновляем старые поля
                t['operator_response'] = response
                t['operator_name'] = username
                t['status'] = 'answered'
                t['answered_at'] = datetime.now().isoformat()
                t['updated_at'] = datetime.now().isoformat()
                
                # Обновляем тикет в данных
                tickets_data['tickets'][i] = t
                updated = True
                
                info_msg = f"Тикет #{ticket_id} обновлен. operator_chat_history: {len(t['operator_chat_history'])} сообщений, operator_messages: {len(t.get('operator_messages', []))} сообщений"
                print(f"[INFO] {info_msg}")
                logger.info(info_msg)
                break
        
        # Сохраняем обновленные данные
        if updated:
            try:
                print(f"[SAVE] Начинаем сохранение в файл: {ticket_file}")
                
                # Сохраняем данные в файл
                with open(ticket_file, 'w', encoding='utf-8') as f:
                    json.dump(tickets_data, f, ensure_ascii=False, indent=2)
                    f.flush()
                    if hasattr(f, 'fileno'):
                        try:
                            os.fsync(f.fileno())
                        except:
                            pass
                
                print(f"[SAVE] Файл записан: {ticket_file}")
                
                # Проверяем, что файл действительно обновился
                if os.path.exists(ticket_file):
                    with open(ticket_file, 'r', encoding='utf-8') as f:
                        verify_data = json.load(f)
                        verify_ticket = None
                        for t in verify_data.get('tickets', []):
                            ticket_id_verify = t.get('id')
                            if isinstance(ticket_id_verify, str):
                                try:
                                    ticket_id_verify = int(ticket_id_verify)
                                except ValueError:
                                    continue
                            if ticket_id_verify == ticket_id_int:
                                verify_ticket = t
                                break
                        
                        if verify_ticket:
                            operator_chat_history_count = len(verify_ticket.get('operator_chat_history', []))
                            operator_messages_count = len(verify_ticket.get('operator_messages', []))
                            success_msg = f"Оператор {username} ответил на тикет #{ticket_id} в отделе {dept}. Файл: {ticket_file}. operator_chat_history: {operator_chat_history_count} сообщений, operator_messages: {operator_messages_count} сообщений."
                            print(f"[SUCCESS] {success_msg}")
                            logger.info(success_msg)
                            
                            # Выводим последнее сообщение для проверки
                            if verify_ticket.get('operator_chat_history'):
                                last_msg = verify_ticket['operator_chat_history'][-1]
                                print(f"[VERIFY] Последнее сообщение в operator_chat_history: role={last_msg.get('role')}, content={last_msg.get('content')[:50]}")
                        else:
                            error_msg = f"Тикет #{ticket_id} не найден в файле {ticket_file} после сохранения!"
                            print(f"[ERROR] {error_msg}")
                            logger.error(error_msg)
                
                return jsonify({
                    'success': True,
                    'message': 'Ответ отправлен клиенту',
                    'ticket_id': ticket_id,
                    'department': dept,
                    'file': ticket_file
                })
            except Exception as e:
                error_msg = f"Ошибка при сохранении файла {ticket_file}: {e}"
                print(f"[ERROR] {error_msg}")
                logger.error(error_msg, exc_info=True)
                return jsonify({'error': f'Ошибка при сохранении ответа: {str(e)}'}), 500
        else:
            logger.error(f"❌ Тикет #{ticket_id} не найден в файле {ticket_file} после поиска. "
                       f"Проверено {len(tickets_data.get('tickets', []))} тикетов.")
            return jsonify({'error': f'Тикет #{ticket_id} не найден в базе данных отдела {dept}'}), 404
        
        return jsonify({
            'success': True,
            'message': 'Ответ отправлен клиенту'
        })
    
    except Exception as e:
        logger.error(f"Ошибка при отправке ответа оператора: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при отправке ответа.'
        }), 500

@app.route('/api/operator/analyze', methods=['POST'])
def api_operator_analyze():
    """API для анализа сообщения клиента с помощью AI"""
    if 'username' not in session or session.get('role') != 'operator':
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        message = data.get('message', '').strip()
        ticket_id = data.get('ticket_id')
        
        if not message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Получаем контекст тикета если есть
        context = ""
        if ticket_id:
            ticket = get_ticket_by_id(ticket_id)
            if ticket:
                context = f"\n\nКонтекст тикета:\nСообщение клиента: {ticket.get('message', '')}\n"
                if ticket.get('chat_history'):
                    context += "История чата:\n"
                    for msg in ticket.get('chat_history', [])[-5:]:
                        role = "Клиент" if msg.get('role') == 'user' else "AI"
                        context += f"{role}: {msg.get('content', '')}\n"
        
        # Системный промпт для анализа
        system_prompt = f"""Ты - опытный оператор службы поддержки. Проанализируй сообщение клиента и предоставь:
1. Краткий анализ проблемы
2. Рекомендации по решению
3. Приоритет (низкий, средний, высокий, критический)
4. Нужна ли помощь мастера/специалиста

Отвечай кратко и профессионально на русском языке."""
        
        full_prompt = f"{system_prompt}{context}\n\nСообщение клиента: {message}\n\nАнализ:"
        
        # Инициализация Gemini
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'analysis': 'AI сервис временно недоступен.'
            }), 500
        
        genai.configure(api_key=api_key)
        model_name = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
        chat_model = genai.GenerativeModel(model_name)
        
        try:
            response = chat_model.generate_content(
                full_prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.7,
                    top_p=0.8,
                    top_k=40,
                    max_output_tokens=1024,
                )
            )
            analysis = response.text.strip() if response.text else 'Не удалось получить анализ.'
        except Exception as e:
            logger.error(f"Ошибка при генерации анализа: {e}")
            analysis = 'Произошла ошибка при анализе сообщения.'
        
        return jsonify({
            'success': True,
            'analysis': analysis
        })
    
    except Exception as e:
        logger.error(f"Ошибка при анализе сообщения: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'analysis': 'Извините, произошла ошибка при анализе.'
        }), 500

@app.route('/api/operator/call-master', methods=['POST'])
def api_operator_call_master():
    """API для вызова мастера через AI анализ"""
    if 'username' not in session or session.get('role') != 'operator':
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        message_id = data.get('message_id')
        ticket_id = data.get('ticket_id')
        
        if not message_id:
            return jsonify({'error': 'ID сообщения обязателен'}), 400
        
        # Получаем сообщение и контекст
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        operator_chat_file = os.path.join(data_dir, 'operator_chat.json')
        
        user_message = ""
        context = ""
        
        if os.path.exists(operator_chat_file):
            with open(operator_chat_file, 'r', encoding='utf-8') as f:
                chat_data = json.load(f)
                for msg in chat_data.get('messages', []):
                    if msg.get('id') == message_id:
                        user_message = msg.get('message', '')
                        break
        
        if ticket_id:
            ticket = get_ticket_by_id(ticket_id)
            if ticket:
                context = f"\n\nКонтекст тикета:\nСообщение: {ticket.get('message', '')}\n"
                if ticket.get('chat_history'):
                    context += "История:\n"
                    for msg in ticket.get('chat_history', [])[-5:]:
                        context += f"{msg.get('role', 'user')}: {msg.get('content', '')}\n"
        
        # AI анализ для определения необходимости вызова мастера
        system_prompt = """Ты - опытный оператор службы поддержки. Проанализируй ситуацию клиента и определи:
1. Нужен ли вызов мастера/специалиста?
2. Какой тип мастера нужен?
3. Срочность ситуации
4. Рекомендации для мастера

Отвечай кратко и профессионально на русском языке."""
        
        full_prompt = f"{system_prompt}{context}\n\nСообщение клиента: {user_message}\n\nАнализ:"
        
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'recommendation': 'AI сервис временно недоступен.'
            }), 500
        
        genai.configure(api_key=api_key)
        model_name = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
        chat_model = genai.GenerativeModel(model_name)
        
        try:
            response = chat_model.generate_content(
                full_prompt,
                generation_config=genai.types.GenerationConfig(
                    temperature=0.7,
                    top_p=0.8,
                    top_k=40,
                    max_output_tokens=1024,
                )
            )
            recommendation = response.text.strip() if response.text else 'Не удалось получить рекомендацию.'
        except Exception as e:
            logger.error(f"Ошибка при генерации рекомендации: {e}")
            recommendation = 'Произошла ошибка при анализе.'
        
        return jsonify({
            'success': True,
            'recommendation': recommendation,
            'should_call_master': 'мастер' in recommendation.lower() or 'специалист' in recommendation.lower()
        })
    
    except Exception as e:
        logger.error(f"Ошибка при вызове мастера: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'recommendation': 'Извините, произошла ошибка.'
        }), 500

def get_ticket_by_id(ticket_id):
    """Получить тикет по ID из всех отделов"""
    data_dir = os.path.join(os.path.dirname(__file__), 'data')
    ticket_files = [
        'it_help_desk_tickets.json',
        'retail_help_desk_tickets.json',
        'medical_help_desk_tickets.json',
        'legal_help_desk_tickets.json',
        'telecom_help_desk_tickets.json',
        'auto_help_desk_tickets.json',
        'general_help_desk_tickets.json'
    ]
    
    # Преобразуем ticket_id в число для сравнения
    ticket_id_int = ticket_id
    if isinstance(ticket_id, str):
        try:
            ticket_id_int = int(ticket_id)
        except ValueError:
            return None
    
    for ticket_file in ticket_files:
        file_path = os.path.join(data_dir, ticket_file)
        if os.path.exists(file_path):
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    tickets_data = json.load(f)
                    for ticket in tickets_data.get('tickets', []):
                        ticket_id_in_file = ticket.get('id')
                        # Преобразуем ID из файла в число для сравнения
                        if isinstance(ticket_id_in_file, str):
                            try:
                                ticket_id_in_file = int(ticket_id_in_file)
                            except ValueError:
                                continue
                        
                        if ticket_id_in_file == ticket_id_int:
                            return ticket
            except Exception as e:
                logger.error(f"Ошибка при чтении файла {ticket_file}: {e}")
                continue
    return None

@app.route('/api/it-help-desk/create-ticket', methods=['POST'])
def api_it_help_desk_create_ticket():
    """API для создания заявки IT Help Desk"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        chat_history = data.get('chat_history', [])
        user_name = session.get('username', 'Unknown')
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Путь к JSON файлу
        data_dir = os.path.join(os.path.dirname(__file__), 'data')
        os.makedirs(data_dir, exist_ok=True)
        tickets_file = os.path.join(data_dir, 'it_help_desk_tickets.json')
        
        # Читаем существующие заявки
        if os.path.exists(tickets_file):
            with open(tickets_file, 'r', encoding='utf-8') as f:
                tickets_data = json.load(f)
        else:
            tickets_data = {
                "tickets": [],
                "last_id": 0,
                "departments": {
                    "it_help_desk": {
                        "name": "IT Help Desk",
                        "description": "Отдел технической поддержки IT"
                    }
                }
            }
        
        # Создаем новую заявку
        ticket_id = tickets_data['last_id'] + 1
        ticket = {
            "id": ticket_id,
            "user": user_name,
            "message": user_message,
            "chat_history": chat_history[-5:] if chat_history else [],  # Последние 5 сообщений
            "department": "it_help_desk",
            "status": "open",
            "priority": "medium",
            "created_at": datetime.now().isoformat(),
            "updated_at": datetime.now().isoformat(),
            "assigned_to": None,
            "resolved_at": None
        }
        
        tickets_data['tickets'].append(ticket)
        tickets_data['last_id'] = ticket_id
        
        # Сохраняем в файл
        os.makedirs('data', exist_ok=True)
        with open(tickets_file, 'w', encoding='utf-8') as f:
            json.dump(tickets_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Создана заявка IT Help Desk #{ticket_id} от пользователя {user_name}")
        
        return jsonify({
            'success': True,
            'ticket_id': ticket_id,
            'message': f'Заявка #{ticket_id} успешно создана и отправлена в отдел IT Help Desk. Оператор свяжется с вами в ближайшее время.',
            'ticket': ticket
        })
    
    except Exception as e:
        logger.error(f"Ошибка при создании заявки IT Help Desk: {e}")
        return jsonify({
            'error': 'Внутренняя ошибка сервера',
            'message': 'Извините, произошла ошибка при создании заявки. Пожалуйста, попробуйте еще раз.'
        }), 500

@app.route('/fraud-chatbot')
def fraud_chatbot():
    """Страница чатбота о мошенничестве"""
    if 'username' not in session:
        return redirect(url_for('login'))
    return render_template('fraud_chatbot.html', username=session.get('username'))

# ==================== Fraud Detection API Endpoints ====================

# Инициализация модели (глобальная переменная)
fraud_model = None
fraud_threshold = 0.3

def init_fraud_model():
    """Инициализация модели fraud detection"""
    global fraud_model
    if not ML_MODELS_AVAILABLE or FraudDetectionModel is None:
        print("⚠️  ML модели недоступны. Fraud Detection будет работать в ограниченном режиме.")
        fraud_model = None
        return
    
    try:
        # FraudDetectionModel уже импортирован в начале файла
        model_path = 'ml_models/saved_models/fraud_model.pkl'
        if os.path.exists(model_path):
            fraud_model = FraudDetectionModel()
            metrics = fraud_model.load_model(model_path)
            print("✓ Fraud Detection модель загружена")
            if metrics:
                print(f"  ✓ Метрики загружены: precision={metrics.get('precision', 0):.4f}, recall={metrics.get('recall', 0):.4f}")
        else:
            print("⚠️  Модель не найдена. Сначала обучите модель: python ml_models/train_model.py")
            # Создаем пустую модель для работы без обучения
            fraud_model = FraudDetectionModel()
            print("✓ Fraud Detection модель инициализирована (без предобучения)")
    except Exception as e:
        print(f"⚠️  Ошибка загрузки модели: {e}")
        import traceback
        traceback.print_exc()
        # Создаем пустую модель в случае ошибки
        try:
            fraud_model = FraudDetectionModel()
            print("✓ Fraud Detection модель инициализирована (после ошибки)")
        except:
            fraud_model = None
            print("⚠️  Не удалось инициализировать Fraud Detection модель")

# Инициализация при старте
print("=" * 60)
print("🚀 Запуск HelpDesk AI Support Platform...")
print("=" * 60)

# Инициализация Fraud Detection модели
init_fraud_model()

# Вывод статуса загрузки модулей
print("\n📦 Статус загрузки модулей:")
print(f"  ✓ Flask приложение инициализировано")
print(f"  ✓ Gemini AI модель настроена")
print(f"  ✓ AI-Procure модули: загружены")
print(f"  ✓ AI-Scrum Master модули: загружены")
print(f"  ✓ AI-Business Analyst модули: загружены")
print(f"  ✓ AI-Code Review Assistant модули: загружены")
if ML_MODELS_AVAILABLE:
    print(f"  ✓ ML Models модули: загружены")
else:
    print(f"  ⚠️  ML Models модули: недоступны")
if fraud_model is not None:
    print(f"  ✓ Fraud Detection модель: {'загружена' if hasattr(fraud_model, 'model') and fraud_model.model is not None else 'инициализирована'}")
else:
    print(f"  ⚠️  Fraud Detection модель: недоступна")

print("=" * 60)
print("✅ Приложение готово к работе!")
print("=" * 60)

@app.route('/api/fraud-detection/threshold', methods=['POST'])
def api_fraud_detection_threshold():
    """Установка порога риска"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        global fraud_threshold
        fraud_threshold = float(data.get('threshold', 0.3))
        return jsonify({'success': True, 'threshold': fraud_threshold})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/fraud-detection/get-threshold', methods=['GET'])
def api_fraud_detection_get_threshold():
    """Получение текущего порога"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    global fraud_threshold
    return jsonify({'success': True, 'threshold': fraud_threshold})

@app.route('/api/fraud-detection/check-all', methods=['POST'])
def api_fraud_detection_check_all():
    """Проверка всех транзакций"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if fraud_model is None:
            return jsonify({'error': 'Модель не загружена. Сначала обучите модель.'}), 500
        
        data = request.json
        threshold = float(data.get('threshold', fraud_threshold))
        
        # Загрузка данных
        transactions_path = 'csv/транзакции в Мобильном интернет Банкинге.csv'
        behavioral_path = 'csv/поведенческие паттерны клиентов.csv'
        
        if not os.path.exists(transactions_path):
            return jsonify({'error': 'Файл транзакций не найден'}), 404
        
        # Загружаем данные через метод модели, который объединяет транзакции с поведенческими данными
        print("  📊 Загрузка данных (транзакции + поведенческие паттерны)...")
        try:
            data_df = fraud_model.load_data(
                transactions_path, 
                behavioral_path if os.path.exists(behavioral_path) else None
            )
            print(f"  ✓ Данные загружены: {len(data_df)} записей")
            if os.path.exists(behavioral_path):
                print(f"  ✓ Использован файл поведенческих данных: {behavioral_path}")
            else:
                print(f"  ⚠️  Файл поведенческих данных не найден: {behavioral_path} (используются только транзакции)")
        except Exception as e:
            import traceback
            traceback.print_exc()
            return jsonify({'error': f'Ошибка при загрузке данных: {str(e)}'}), 500
        
        # Сохраняем общее количество транзакций в файле
        total_in_file = len(data_df)
        
        # Ограничение количества транзакций (настраиваемое через параметр)
        # По умолчанию обрабатываем все транзакции, но можно ограничить для быстрой проверки
        max_transactions = data.get('max_transactions')
        if max_transactions:
            try:
                max_transactions = int(max_transactions)
                if max_transactions > 0 and len(data_df) > max_transactions:
                    data_df = data_df.sample(n=max_transactions, random_state=42).copy()
                    print(f"  ⚠️  Ограничение: проверяется {max_transactions} из {total_in_file} транзакций")
            except (ValueError, TypeError):
                pass  # Игнорируем неверное значение
        
        # Сброс индекса для корректной индексации
        transactions = data_df.reset_index(drop=True)
        
        print(f"  ✓ Загружено транзакций для проверки: {len(transactions)} из {total_in_file} в файле")
        
        # Feature Engineering
        try:
            processed_data = fraud_model.feature_engineering(transactions)
        except Exception as e:
            return jsonify({'error': f'Ошибка при создании признаков: {str(e)}'}), 500
        
        # Подготовка признаков
        try:
            X = fraud_model.prepare_features(processed_data, is_training=False)
        except Exception as e:
            return jsonify({'error': f'Ошибка при подготовке признаков: {str(e)}'}), 500
        
        # Проверка на пустые признаки
        if X.empty or len(X) == 0:
            return jsonify({'error': 'Не удалось создать признаки для предсказания'}), 500
        
        # Проверка признаков на наличие данных (не все нули)
        if hasattr(X, 'values'):
            non_zero_features = (X.values != 0).sum(axis=0)
            non_zero_count = non_zero_features.sum()
            total_features = len(X.columns)
            
            # Диагностическая информация
            feature_info = {
                'total_features': total_features,
                'non_zero_features': int(non_zero_count),
                'feature_names': list(X.columns)[:20],  # Первые 20 для диагностики
                'sample_values': {}
            }
            
            # Берем примеры значений из первых признаков
            for col in X.columns[:5]:
                if len(X[col]) > 0:
                    feature_info['sample_values'][col] = {
                        'min': float(X[col].min()),
                        'max': float(X[col].max()),
                        'mean': float(X[col].mean()),
                        'non_zero_count': int((X[col] != 0).sum())
                    }
            
            if non_zero_count == 0:
                available_columns = list(transactions.columns) if hasattr(transactions, 'columns') else []
                return jsonify({
                    'error': 'Все признаки равны нулю. Проверьте данные и feature engineering.',
                    'hint': 'Возможно, данные не содержат необходимых колонок или feature engineering не создал признаки',
                    'diagnostics': feature_info,
                    'available_columns_in_raw_data': available_columns[:20]
                }), 500
            
            # Предупреждение если слишком мало ненулевых признаков
            if non_zero_count < total_features * 0.1:
                print(f"  ⚠️  Предупреждение: только {non_zero_count} из {total_features} признаков содержат ненулевые значения")
        
        # Проверка, что модель обучена
        if not hasattr(fraud_model, 'model') or fraud_model.model is None:
            return jsonify({
                'error': 'Модель не обучена. Пожалуйста, сначала обучите модель используя train_model.py',
                'hint': 'Запустите: python ml_models/train_model.py'
            }), 500
        
        # Предсказание
        try:
            predictions = fraud_model.predict(X, threshold=threshold)
        except Exception as e:
            return jsonify({'error': f'Ошибка при предсказании: {str(e)}'}), 500
        
        # Формирование результатов
        fraud_list = []
        # Нормализация результатов предсказания в списки
        is_fraud_list = predictions['is_fraud']
        if not isinstance(is_fraud_list, list):
            is_fraud_list = [is_fraud_list]
        
        prob_list = predictions['fraud_probability']
        if not isinstance(prob_list, list):
            prob_list = [prob_list]
        
        # Проверка соответствия размеров
        if len(is_fraud_list) != len(transactions) or len(prob_list) != len(transactions):
            return jsonify({
                'error': f'Несоответствие размеров: транзакций {len(transactions)}, предсказаний {len(is_fraud_list)}'
            }), 500
        
        # Статистика вероятностей для отладки
        if prob_list:
            max_prob = max(prob_list)
            min_prob = min(prob_list)
            avg_prob = sum(prob_list) / len(prob_list)
            # Подсчет транзакций с вероятностью выше порога
            high_risk_count = sum(1 for p in prob_list if p >= threshold)
            # Подсчет транзакций с вероятностью выше 0.1 (для диагностики)
            medium_risk_count = sum(1 for p in prob_list if p >= 0.1)
        else:
            max_prob = min_prob = avg_prob = 0.0
            high_risk_count = 0
            medium_risk_count = 0
        
        # Предупреждение если все вероятности очень низкие
        warning_message = None
        if prob_list and max_prob < 0.1:
            warning_message = f"Все вероятности мошенничества очень низкие (макс: {max_prob:.4f}). Возможно, модель не обучена или данные не подходят для этой модели."
        elif prob_list and high_risk_count == 0 and medium_risk_count > 0:
            warning_message = f"Найдено {medium_risk_count} транзакций с вероятностью > 0.1, но ни одна не превышает порог {threshold}. Рассмотрите возможность снижения порога."
        
        # Итерация по позициям (не по индексам DataFrame)
        for pos in range(len(transactions)):
            if pos < len(is_fraud_list) and pos < len(prob_list):
                row = transactions.iloc[pos]
                # Преобразуем в bool безопасно
                fraud_val = is_fraud_list[pos]
                if isinstance(fraud_val, (bool, int, np.bool_, np.integer)):
                    is_fraud = bool(int(fraud_val))
                else:
                    is_fraud = bool(fraud_val) if fraud_val else False
                prob = float(prob_list[pos])
                
                # Добавляем транзакцию если она помечена как мошенническая
                # (is_fraud уже учитывает threshold)
                if is_fraud:
                    # Формируем базовую информацию о транзакции
                    txn_data = {
                        'id': pos + 1,
                        'customer_id': str(row.get('cst_dim_id', 'N/A')),
                        'amount': float(row.get('amount', 0)) if pd.notna(row.get('amount')) else 0.0,
                        'date': str(row.get('transdate', '')) if pd.notna(row.get('transdate')) else '',
                        'transdatetime': str(row.get('transdatetime', '')) if pd.notna(row.get('transdatetime')) else '',
                        'fraud_probability': round(prob, 4),
                        'is_fraud_predicted': is_fraud,
                        'is_confirmed_fraud': bool(row.get('target', 0) == 1) if 'target' in row and pd.notna(row.get('target')) else False
                    }
                    
                    # Добавляем поведенческие данные если они доступны
                    behavioral_fields = {
                        'monthly_os_changes': 'monthly_os_changes',
                        'monthly_phone_model_changes': 'monthly_phone_model_changes',
                        'last_phone_model': 'last_phone_model_categorical',
                        'last_os_version': 'last_os_categorical',
                        'logins_last_7_days': 'logins_last_7_days',
                        'logins_last_30_days': 'logins_last_30_days',
                        'login_frequency_7d': 'login_frequency_7d',
                        'login_frequency_30d': 'login_frequency_30d',
                        'freq_change_7d_vs_mean': 'freq_change_7d_vs_mean',
                        'logins_7d_over_30d_ratio': 'logins_7d_over_30d_ratio',
                        'avg_login_interval_30d': 'avg_login_interval_30d',
                        'std_login_interval_30d': 'std_login_interval_30d',
                        'burstiness_login_interval': 'burstiness_login_interval',
                        'fano_factor_login_interval': 'fano_factor_login_interval',
                        'zscore_avg_login_interval_7d': 'zscore_avg_login_interval_7d'
                    }
                    
                    behavioral_data = {}
                    for key, col_name in behavioral_fields.items():
                        if col_name in row and pd.notna(row.get(col_name)):
                            value = row.get(col_name)
                            # Преобразуем в числовой формат если возможно
                            try:
                                if isinstance(value, (int, float)):
                                    behavioral_data[key] = float(value)
                                else:
                                    behavioral_data[key] = value
                            except (ValueError, TypeError):
                                behavioral_data[key] = value
                    
                    if behavioral_data:
                        txn_data['behavioral_patterns'] = behavioral_data
                    
                    fraud_list.append(txn_data)
        
        # Статистика
        total = len(transactions)
        fraud_count = len(fraud_list)
        safe_count = total - fraud_count
        
        # Бизнес метрики: правильно заблокированные переводы (после определения fraud_count)
        correctly_blocked = sum(1 for txn in fraud_list if txn.get('is_confirmed_fraud', False))
        false_positives = fraud_count - correctly_blocked
        true_positives = correctly_blocked
        false_negatives = 0  # Не можем вычислить без проверки всех транзакций
        
        statistics = {
            'total_checked': total,
            'fraud_detected': fraud_count,
            'safe_transactions': safe_count,
            'fraud_percentage': round((fraud_count / total * 100) if total > 0 else 0, 2),
            'threshold_used': threshold,
            'total_in_file': total_in_file if 'total_in_file' in locals() else total,
            'business_metrics': {
                'correctly_blocked': correctly_blocked,
                'false_positives': false_positives,
                'true_positives': true_positives,
                'blocking_accuracy': round((correctly_blocked / fraud_count * 100) if fraud_count > 0 else 0, 2)
            },
            'probability_stats': {
                'max': round(max_prob, 4) if prob_list else 0.0,
                'min': round(min_prob, 4) if prob_list else 0.0,
                'avg': round(avg_prob, 4) if prob_list else 0.0,
                'high_risk_count': high_risk_count,
                'medium_risk_count': medium_risk_count if prob_list else 0
            }
        }
        
        response = {
            'success': True,
            'fraud_transactions': fraud_list,
            'statistics': statistics
        }
        
        if warning_message:
            response['warning'] = warning_message
        
        return jsonify(response)
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/fraud-detection/model-info', methods=['GET'])
def api_fraud_detection_model_info():
    """Информация о модели"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if fraud_model is None:
            return jsonify({
                'success': True,
                'model_info': {
                    'status': 'not_loaded',
                    'model_type': None,
                    'feature_count': 0
                }
            })
        
        model_info = {
            'status': 'loaded' if (hasattr(fraud_model, 'model') and fraud_model.model is not None) else 'not_trained',
            'model_type': fraud_model.model_type if hasattr(fraud_model, 'model_type') else 'unknown',
            'feature_count': len(fraud_model.feature_names) if hasattr(fraud_model, 'feature_names') else 0,
            'is_trained': hasattr(fraud_model, 'model') and fraud_model.model is not None
        }
        
        # Загружаем метрики если модель обучена
        if model_info['is_trained']:
            # Проверяем есть ли метрики в объекте модели
            if hasattr(fraud_model, 'metrics') and fraud_model.metrics:
                model_info['metrics'] = fraud_model.metrics
            else:
                # Пытаемся загрузить из файла
                model_path = 'ml_models/saved_models/fraud_model.pkl'
                if os.path.exists(model_path):
                    try:
                        import joblib
                        model_data = joblib.load(model_path)
                        if 'metrics' in model_data:
                            model_info['metrics'] = model_data['metrics']
                            # Сохраняем в объект модели для будущего использования
                            fraud_model.metrics = model_data['metrics']
                    except Exception as e:
                        print(f"  ⚠️  Ошибка при загрузке метрик: {e}")
                        pass
        
        if not model_info['is_trained']:
            model_info['message'] = 'Модель инициализирована, но не обучена. Запустите train_model.py для обучения.'
        
        return jsonify({
            'success': True,
            'model_info': model_info
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/fraud-detection/retrain', methods=['POST'])
def api_fraud_detection_retrain():
    """Переобучение модели"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if fraud_model is None:
            return jsonify({'error': 'Модель не инициализирована'}), 500
        
        data = request.json or {}
        model_type = data.get('model_type', 'xgboost')  # xgboost или lightgbm
        
        # Загрузка данных
        transactions_path = 'csv/транзакции в Мобильном интернет Банкинге.csv'
        behavioral_path = 'csv/поведенческие паттерны клиентов.csv'
        
        if not os.path.exists(transactions_path):
            return jsonify({'error': 'Файл транзакций не найден'}), 404
        
        # Загружаем и обрабатываем данные
        data_df = fraud_model.load_data(transactions_path, behavioral_path if os.path.exists(behavioral_path) else None)
        data_with_features = fraud_model.feature_engineering(data_df)
        X, y = fraud_model.prepare_features(data_with_features, is_training=True)
        
        # Создаем новую модель если тип изменился
        if model_type != fraud_model.model_type:
            # Проверка доступности библиотек
            try:
                import xgboost as xgb
                xgb_available = True
            except ImportError:
                xgb_available = False
            
            try:
                import lightgbm as lgb
                lgb_available = True
            except ImportError:
                lgb_available = False
            
            if model_type == 'xgboost' and not xgb_available:
                return jsonify({'error': 'XGBoost не установлен'}), 400
            if model_type == 'lightgbm' and not lgb_available:
                return jsonify({'error': 'LightGBM не установлен'}), 400
            
            fraud_model.model_type = model_type
            fraud_model.model = None
        
        # Обучение с оптимизацией для Recall
        metrics = fraud_model.train(X, y, test_size=0.2, random_state=42, optimize_recall=True)
        
        # Сохранение модели с метриками
        model_path = 'ml_models/saved_models/fraud_model.pkl'
        fraud_model.save_model(model_path, metrics=metrics)
        # Сохраняем метрики в объект модели для доступа
        fraud_model.metrics = metrics
        
        return jsonify({
            'success': True,
            'message': 'Модель успешно переобучена',
            'metrics': metrics,
            'model_type': model_type
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/fraud-detection/shap-values', methods=['POST'])
def api_fraud_detection_shap():
    """Получение SHAP значений для транзакции"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if fraud_model is None or fraud_model.model is None:
            return jsonify({'error': 'Модель не обучена'}), 500
        
        data = request.json
        transaction_data = data.get('transaction')
        
        if not transaction_data:
            return jsonify({'error': 'Данные транзакции не предоставлены'}), 400
        
        # Подготовка данных
        df = pd.DataFrame([transaction_data])
        processed_data = fraud_model.feature_engineering(df)
        X = fraud_model.prepare_features(processed_data, is_training=False)
        
        # Получение SHAP значений
        shap_result = fraud_model.get_shap_values(X, max_samples=1)
        
        if shap_result and shap_result.get('available'):
            return jsonify({
                'success': True,
                'shap_values': shap_result
            })
        else:
            return jsonify({
                'success': False,
                'message': 'SHAP значения недоступны. Установите библиотеку shap: pip install shap',
                'error': shap_result.get('error') if shap_result else 'SHAP не установлен'
            })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/fraud-detection/execute-action', methods=['POST'])
def api_fraud_detection_execute_action():
    """Выполнение действия с транзакцией (блокировка, верификация и т.д.)"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        if not data:
            return jsonify({'error': 'Данные не предоставлены'}), 400
        
        action = data.get('action')
        transaction_id = data.get('transaction_id')
        customer_id = data.get('customer_id')
        
        print(f"  📋 Выполнение действия: {action}, транзакция: {transaction_id}, клиент: {customer_id}")
        
        if not action:
            return jsonify({'error': 'Не указано действие'}), 400
        if not transaction_id:
            return jsonify({'error': 'Не указан ID транзакции'}), 400
        
        # Сохранение действий в файл (можно заменить на БД)
        actions_file = 'data/fraud_actions.json'
        os.makedirs(os.path.dirname(actions_file), exist_ok=True)
        
        # Загрузка существующих действий
        actions = []
        if os.path.exists(actions_file):
            try:
                with open(actions_file, 'r', encoding='utf-8') as f:
                    actions = json.load(f)
            except:
                actions = []
        
        # Добавление нового действия
        action_record = {
            'transaction_id': transaction_id,
            'customer_id': customer_id,
            'action': action,
            'executed_by': session.get('username', 'unknown'),
            'executed_at': datetime.now().isoformat(),
            'status': 'completed'
        }
        
        actions.append(action_record)
        
        # Сохранение
        with open(actions_file, 'w', encoding='utf-8') as f:
            json.dump(actions, f, ensure_ascii=False, indent=2)
        
        # Определение сообщения в зависимости от действия
        action_messages = {
            'block': f'Транзакция #{transaction_id} заблокирована. Клиент {customer_id} уведомлен.',
            'request_verification': f'Запрос на дополнительную верификацию отправлен клиенту {customer_id} для транзакции #{transaction_id}.',
            'escalate': f'Дело по транзакции #{transaction_id} передано в отдел безопасности. Клиент: {customer_id}.',
            'view_history': f'История транзакций клиента {customer_id} загружена.',
            'monitor': f'Транзакции клиента {customer_id} добавлены в мониторинг.',
            'standard_check': f'Стандартная проверка транзакции #{transaction_id} выполнена.',
            'improve_model': f'Данные о транзакции #{transaction_id} добавлены для улучшения модели.'
        }
        
        message = action_messages.get(action, f'Действие "{action}" выполнено для транзакции #{transaction_id}')
        
        return jsonify({
            'success': True,
            'message': message,
            'transaction_status': action,
            'action_record': action_record
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/fraud-detection/client-history/<customer_id>', methods=['GET'])
def api_fraud_detection_client_history(customer_id):
    """Получение истории транзакций клиента"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        transactions_path = 'csv/транзакции в Мобильном интернет Банкинге.csv'
        
        if not os.path.exists(transactions_path):
            return jsonify({'error': 'Файл транзакций не найден'}), 404
        
        # Определение кодировки
        def detect_encoding(filepath):
            encodings = ['utf-8', 'utf-8-sig', 'windows-1251', 'cp1251', 'latin-1', 'iso-8859-1']
            for enc in encodings:
                try:
                    with open(filepath, 'r', encoding=enc) as f:
                        f.read(10000)
                    return enc
                except (UnicodeDecodeError, UnicodeError):
                    continue
            return 'utf-8'
        
        file_encoding = detect_encoding(transactions_path)
        
        # Загрузка данных
        try:
            transactions = pd.read_csv(transactions_path, encoding=file_encoding, sep=';')
        except:
            try:
                transactions = pd.read_csv(transactions_path, encoding=file_encoding, sep=',')
            except:
                transactions = pd.read_csv(transactions_path, encoding=file_encoding)
        
        # Переименование колонок
        column_mapping = {
            'Уникальный идентификатор клиента': 'cst_dim_id',
            'Дата совершенной транзакции': 'transdate',
            'Дата и время совершенной транзакции': 'transdatetime',
            'Сумма совершенного перевода': 'amount',
            'Уникальный идентификатор транзакции': 'docno',
            'Размеченные транзакции(переводы), где 1 - мошенническая операция , 0 - чистая': 'target'
        }
        
        for old_name, new_name in column_mapping.items():
            if old_name in transactions.columns:
                transactions = transactions.rename(columns={old_name: new_name})
        
        # Очистка и преобразование amount в числовой формат
        if 'amount' in transactions.columns:
            # Удаляем кавычки если есть
            if transactions['amount'].dtype == 'object':
                transactions['amount'] = transactions['amount'].astype(str).str.replace("'", "").str.strip()
            
            # Функция для безопасного преобразования amount
            def safe_convert_amount(value):
                if pd.isna(value) or value == '' or value == 'nan':
                    return 0.0
                try:
                    # Если это строка с несколькими значениями, берем первое
                    if isinstance(value, str) and len(value) > 20:  # Подозрительно длинная строка
                        # Пробуем найти первое число (может быть с точкой)
                        match = re.search(r'\d+\.?\d*', value)
                        if match:
                            return float(match.group())
                        return 0.0
                    return float(value)
                except (ValueError, TypeError):
                    return 0.0
            
            # Применяем безопасное преобразование
            transactions['amount'] = transactions['amount'].apply(safe_convert_amount)
            # Дополнительная проверка через pd.to_numeric с errors='coerce'
            transactions['amount'] = pd.to_numeric(transactions['amount'], errors='coerce').fillna(0.0)
        
        # Фильтрация по клиенту - пробуем разные способы сравнения
        client_transactions = pd.DataFrame()
        
        # Способ 1: Сравнение как строки
        if 'cst_dim_id' in transactions.columns:
            client_transactions = transactions[transactions['cst_dim_id'].astype(str).str.strip() == str(customer_id).strip()].copy()
        
        # Способ 2: Если не найдено, пробуем как число
        if client_transactions.empty:
            try:
                customer_id_int = int(customer_id)
                # Преобразуем cst_dim_id в числовой формат для сравнения
                if 'cst_dim_id' in transactions.columns:
                    transactions['cst_dim_id_numeric'] = pd.to_numeric(transactions['cst_dim_id'], errors='coerce')
                    client_transactions = transactions[transactions['cst_dim_id_numeric'] == customer_id_int].copy()
                    # Удаляем временную колонку
                    if 'cst_dim_id_numeric' in client_transactions.columns:
                        client_transactions = client_transactions.drop('cst_dim_id_numeric', axis=1)
            except (ValueError, TypeError):
                pass
        
        # Способ 3: Пробуем без преобразования типов (если уже числовые)
        if client_transactions.empty and 'cst_dim_id' in transactions.columns:
            try:
                if pd.api.types.is_numeric_dtype(transactions['cst_dim_id']):
                    customer_id_int = int(customer_id)
                    client_transactions = transactions[transactions['cst_dim_id'] == customer_id_int].copy()
            except:
                pass
        
        # Диагностика если не найдено
        if client_transactions.empty:
            # Показываем примеры customer_id для диагностики
            if 'cst_dim_id' in transactions.columns:
                sample_ids = transactions['cst_dim_id'].astype(str).unique()[:10].tolist()
                print(f"  ⚠️  Транзакции для клиента {customer_id} не найдены")
                print(f"  📋 Примеры customer_id в данных: {sample_ids}")
                print(f"  📋 Тип customer_id в данных: {transactions['cst_dim_id'].dtype}")
                print(f"  📋 Тип искомого customer_id: {type(customer_id)}")
            else:
                print(f"  ⚠️  Колонка cst_dim_id не найдена в данных")
                print(f"  📋 Доступные колонки: {list(transactions.columns)}")
            
            return jsonify({
                'success': True,
                'message': f'Транзакции для клиента {customer_id} не найдены',
                'transactions': [],
                'total': 0,
                'debug_info': {
                    'searched_id': customer_id,
                    'sample_ids': sample_ids[:5] if 'sample_ids' in locals() else [],
                    'total_transactions_in_file': len(transactions),
                    'available_columns': list(transactions.columns)[:10]
                }
            })
        
        # Сортировка по дате (последние сначала)
        if 'transdatetime' in client_transactions.columns:
            # Очистка от кавычек
            if client_transactions['transdatetime'].dtype == 'object':
                client_transactions['transdatetime'] = client_transactions['transdatetime'].astype(str).str.replace("'", "")
            client_transactions['transdatetime'] = pd.to_datetime(client_transactions['transdatetime'], errors='coerce')
            client_transactions = client_transactions.sort_values('transdatetime', ascending=False, na_position='last')
        elif 'transdate' in client_transactions.columns:
            # Очистка от кавычек
            if client_transactions['transdate'].dtype == 'object':
                client_transactions['transdate'] = client_transactions['transdate'].astype(str).str.replace("'", "")
            client_transactions['transdate'] = pd.to_datetime(client_transactions['transdate'], errors='coerce')
            client_transactions = client_transactions.sort_values('transdate', ascending=False, na_position='last')
        
        # Ограничение на последние 30 дней - убираем, показываем все транзакции
        # (так как даты в CSV могут быть в будущем или прошлом)
        # Можно добавить опциональный параметр для фильтрации по дате
        
        # Ограничение количества (последние 100 транзакций для показа)
        client_transactions = client_transactions.head(100)
        
        print(f"  📊 Найдено транзакций для клиента {customer_id}: {len(client_transactions)}")
        
        # Форматирование результатов
        transactions_list = []
        for idx, row in client_transactions.iterrows():
            # Получаем дату
            date_value = row.get('transdatetime', row.get('transdate', ''))
            if pd.notna(date_value):
                if isinstance(date_value, pd.Timestamp):
                    date_str = date_value.strftime('%Y-%m-%d %H:%M:%S')
                else:
                    date_str = str(date_value).replace("'", "")
            else:
                date_str = str(date_value) if date_value else ''
            
            # Безопасное получение amount
            amount_value = row.get('amount', 0)
            if pd.notna(amount_value):
                try:
                    amount_float = float(amount_value)
                except (ValueError, TypeError):
                    amount_float = 0.0
            else:
                amount_float = 0.0
            
            transactions_list.append({
                'id': str(row.get('docno', idx)) if pd.notna(row.get('docno')) else str(idx),
                'date': date_str,
                'amount': amount_float,
                'is_fraud': bool(row.get('target', 0) == 1) if 'target' in row and pd.notna(row.get('target')) else False
            })
        
        # Статистика клиента (по всем найденным транзакциям)
        # Безопасное вычисление статистики с обработкой ошибок
        try:
            if 'amount' in client_transactions.columns:
                # Убеждаемся что amount числовой
                client_transactions['amount'] = pd.to_numeric(client_transactions['amount'], errors='coerce').fillna(0.0)
                total_amount = float(client_transactions['amount'].sum())
                avg_amount = float(client_transactions['amount'].mean())
            else:
                total_amount = 0.0
                avg_amount = 0.0
        except Exception as e:
            print(f"  ⚠️  Ошибка при вычислении статистики amount: {e}")
            total_amount = 0.0
            avg_amount = 0.0
        
        # Правильный подсчет мошеннических транзакций
        fraud_count = 0
        if 'target' in client_transactions.columns:
            try:
                # Преобразуем target в числовой формат
                target_numeric = pd.to_numeric(client_transactions['target'], errors='coerce')
                # Подсчитываем количество транзакций где target == 1 (мошенничество)
                fraud_count = int((target_numeric == 1).sum())
            except Exception as e:
                print(f"  ⚠️  Ошибка при подсчете fraud_count: {e}")
                fraud_count = 0
        
        total_all = len(client_transactions)
        
        # Если ограничили до 100, нужно пересчитать статистику по всем
        if total_all >= 100:
            # Получаем все транзакции для статистики
            all_client_transactions = transactions[transactions['cst_dim_id'].astype(str) == str(customer_id)].copy()
            if all_client_transactions.empty:
                try:
                    customer_id_int = int(customer_id)
                    transactions['cst_dim_id_numeric'] = pd.to_numeric(transactions['cst_dim_id'], errors='coerce')
                    all_client_transactions = transactions[transactions['cst_dim_id_numeric'] == customer_id_int].copy()
                    if 'cst_dim_id_numeric' in all_client_transactions.columns:
                        all_client_transactions = all_client_transactions.drop('cst_dim_id_numeric', axis=1)
                except:
                    all_client_transactions = client_transactions.copy()
            
            # Безопасное вычисление статистики для всех транзакций
            try:
                if 'amount' in all_client_transactions.columns:
                    all_client_transactions['amount'] = pd.to_numeric(all_client_transactions['amount'], errors='coerce').fillna(0.0)
                    total_amount = float(all_client_transactions['amount'].sum())
                    avg_amount = float(all_client_transactions['amount'].mean())
            except Exception as e:
                print(f"  ⚠️  Ошибка при вычислении статистики amount для всех транзакций: {e}")
            
            # Правильный подсчет мошеннических транзакций для всех транзакций
            if 'target' in all_client_transactions.columns:
                try:
                    target_numeric = pd.to_numeric(all_client_transactions['target'], errors='coerce')
                    fraud_count = int((target_numeric == 1).sum())
                except Exception as e:
                    print(f"  ⚠️  Ошибка при подсчете fraud_count для всех транзакций: {e}")
            
            total_all = len(all_client_transactions)
        
        return jsonify({
            'success': True,
            'customer_id': customer_id,
            'transactions': transactions_list,
            'total': len(transactions_list),
            'total_all': total_all,  # Всего транзакций у клиента
            'statistics': {
                'total_amount': float(total_amount),
                'fraud_count': int(fraud_count),
                'avg_amount': float(avg_amount),
                'fraud_rate': float(fraud_count / total_all * 100) if total_all > 0 else 0
            }
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/fraud-chatbot/chat', methods=['POST'])
def api_fraud_chatbot_chat():
    """API для обработки сообщений чатбота о мошенничестве"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        user_message = data.get('message', '').strip()
        
        if not user_message:
            return jsonify({'error': 'Сообщение не может быть пустым'}), 400
        
        # Системный промпт для чатбота о мошенничестве
        system_prompt = """Ты - эксперт по финансовой безопасности и противодействию мошенничеству в банковской сфере. 
Твоя задача - помогать пользователям понимать различные виды мошенничества, способы защиты и что делать в случае подозрительных транзакций.

Отвечай на казахском и русском языках (в зависимости от языка вопроса пользователя).
Будь дружелюбным, понятным и профессиональным.
Давай конкретные и практичные советы.
Если не знаешь ответа, честно скажи об этом.

Темы, о которых ты можешь говорить:
- Виды мошенничества в банковской сфере
- Фишинг и социальная инженерия
- Защита от мошенничества
- Что делать при подозрительных транзакциях
- Безопасность мобильного банкинга
- Признаки мошеннических операций
- Как проверить транзакцию
- Процедуры блокировки и верификации

Отвечай кратко, но информативно. Используй эмодзи для лучшего восприятия."""
        
        # Инициализация Gemini
        load_dotenv()
        api_key = os.getenv('GEMINI_API_KEY')
        
        if not api_key:
            return jsonify({
                'error': 'GEMINI_API_KEY не настроен',
                'response': 'Извините, AI сервис временно недоступен. Пожалуйста, обратитесь к специалисту.'
            }), 500
        
        genai.configure(api_key=api_key)
        
        # Создаем модель (используем поддерживаемую версию)
        model_name = os.getenv('GEMINI_MODEL', 'gemini-1.5-flash')
        model = genai.GenerativeModel(model_name)
        
        # Формируем полный промпт
        full_prompt = f"{system_prompt}\n\nВопрос пользователя: {user_message}\n\nОтвет:"
        
        # Генерируем ответ
        response = model.generate_content(
            full_prompt,
            generation_config=genai.types.GenerationConfig(
                temperature=0.7,
                top_p=0.8,
                top_k=40,
                max_output_tokens=1024,
            )
        )
        
        bot_response = response.text.strip() if response.text else "Извините, не удалось сгенерировать ответ. Попробуйте переформулировать вопрос."
        
        return jsonify({
            'success': True,
            'response': bot_response,
            'timestamp': datetime.now().isoformat()
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({
            'error': str(e),
            'response': 'Извините, произошла ошибка при обработке вашего запроса. Попробуйте позже или обратитесь к специалисту.'
        }), 500

# Удалить документ
@app.route('/api/confluence-docs/<doc_id>', methods=['DELETE'])
def api_confluence_doc_delete(doc_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        import json
        doc_file = os.path.join('confluence_docs', f'{doc_id}.json')
        
        # Удаляем файл документа
        if os.path.exists(doc_file):
            os.remove(doc_file)
            
            # Обновляем index.json
            index_file = os.path.join('confluence_docs', 'index.json')
            if os.path.exists(index_file):
                with open(index_file, 'r', encoding='utf-8') as f:
                    index_data = json.load(f)
                
                # Удаляем запись о документе
                index_data = [doc for doc in index_data if doc.get('id') != doc_id]
                
                with open(index_file, 'w', encoding='utf-8') as f:
                    json.dump(index_data, f, ensure_ascii=False, indent=2)
            
            return jsonify({
                'success': True,
                'message': 'Документ успешно удален'
            })
        else:
            return jsonify({'error': 'Документ не найден'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ==================== AI-Procure API Endpoints ====================

# Извлечение параметров тендера
@app.route('/api/ai-procure/extract', methods=['POST'])
def api_procure_extract():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    import time
    start_time = time.time()
    
    try:
        data = request.json
        tender_text = data.get('tender_text', '')
        source_type = data.get('source_type', 'text')
        
        if not tender_text:
            return jsonify({'error': 'Текст тендера не предоставлен'}), 400
        
        tender_params = tender_analyzer.extract_tender_params(tender_text, source_type)
        completeness = tender_analyzer.analyze_tender_completeness(tender_params)
        
        processing_time = time.time() - start_time
        if metrics_tracker:
            metrics_tracker.track_extraction(tender_params, completeness, processing_time)
        
        return jsonify({
            'tender_params': tender_params,
            'completeness': completeness,
            'processing_time': round(processing_time, 2),
            'success': True
        })
    except Exception as e:
        if metrics_tracker:
            metrics_tracker.track_error('extract_error', str(e))
        return jsonify({'error': str(e)}), 500

# Подбор поставщиков
@app.route('/api/ai-procure/match-suppliers', methods=['POST'])
def api_procure_match_suppliers():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        tender_params = data.get('tender_params', {})
        criteria = data.get('criteria', {})
        
        if not tender_params:
            return jsonify({'error': 'Параметры тендера не предоставлены'}), 400
        
        suppliers, summary = supplier_matcher.match_suppliers(tender_params, criteria)
        
        return jsonify({
            'suppliers': suppliers,
            'summary': summary,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Скачивание файлов
@app.route('/uploads/<filename>')
def download_file(filename):
    """Скачивание файлов из папки uploads"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        file_path = os.path.join('uploads', filename)
        if os.path.exists(file_path):
            return send_file(file_path, as_attachment=True)
        else:
            return jsonify({'error': 'Файл не найден'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Анализ рисков
@app.route('/api/ai-procure/analyze-risks', methods=['POST'])
def api_procure_analyze_risks():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        tender_params = data.get('tender_params', {})
        winner_info = data.get('winner_info')
        
        if not tender_params:
            return jsonify({'error': 'Параметры тендера не предоставлены'}), 400
        
        risk_analysis = risk_analyzer.analyze_risks(tender_params, winner_info)
        
        return jsonify({
            'risk_analysis': risk_analysis,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Получить список всех документов Confluence
@app.route('/api/confluence-docs/list', methods=['GET'])
def api_confluence_docs_list():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        import json
        index_file = os.path.join('confluence_docs', 'index.json')
        if os.path.exists(index_file):
            with open(index_file, 'r', encoding='utf-8') as f:
                documents = json.load(f)
            return jsonify({
                'success': True,
                'documents': documents
            })
        else:
            return jsonify({
                'success': True,
                'documents': []
            })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Получить конкретный документ
@app.route('/api/confluence-docs/<doc_id>', methods=['GET'])
def api_confluence_doc(doc_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        import json
        doc_file = os.path.join('confluence_docs', f'{doc_id}.json')
        if os.path.exists(doc_file):
            with open(doc_file, 'r', encoding='utf-8') as f:
                doc_data = json.load(f)
            return jsonify({
                'success': True,
                'document': doc_data
            })
        else:
            return jsonify({'error': 'Документ не найден'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Проверка аффилированности
@app.route('/api/ai-procure/check-affiliation', methods=['POST'])
def api_procure_check_affiliation():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        entities = data.get('entities', [])
        
        if not entities:
            return jsonify({'error': 'Список сущностей не предоставлен'}), 400
        
        affiliation_analysis = risk_analyzer.check_affiliation(entities)
        
        return jsonify({
            'affiliation_analysis': affiliation_analysis,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Генерация комплексного отчёта
@app.route('/api/ai-procure/generate-report', methods=['POST'])
def api_procure_generate_report():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        tender_params = data.get('tender_params', {})
        suppliers = data.get('suppliers', [])
        risk_analysis = data.get('risk_analysis', {})
        completeness = data.get('completeness', {})
        winner = data.get('winner')
        
        if not tender_params:
            return jsonify({'error': 'Параметры тендера не предоставлены'}), 400
        
        report = report_generator.generate_comprehensive_report(
            tender_params, suppliers, risk_analysis, completeness, winner
        )
        
        return jsonify({
            'report': report,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Генерация текста тендера через AI
@app.route('/api/ai-procure/generate-tender-text', methods=['POST'])
def api_procure_generate_tender_text():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        description = data.get('description', '')
        
        if not description:
            return jsonify({'error': 'Описание обязательно'}), 400
        
        tender_text = tender_analyzer.generate_tender_text(description)
        return jsonify({'tender_text': tender_text, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Улучшение описания компании через AI
@app.route('/api/ai-procure/improve-company-description', methods=['POST'])
def api_procure_improve_company_description():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        description = data.get('description', '')
        
        if not description:
            return jsonify({'error': 'Описание обязательно'}), 400
        
        improved_description = tender_analyzer.improve_company_description(description)
        return jsonify({'improved_description': improved_description, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Комплексный анализ тендера (все функции сразу)
@app.route('/api/ai-procure/analyze', methods=['POST'])
def api_procure_analyze():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    import time
    start_time = time.time()
    
    try:
        data = request.json
        tender_text = data.get('tender_text', '')
        source_type = data.get('source_type', 'text')
        include_suppliers = data.get('include_suppliers', True)
        include_risks = data.get('include_risks', True)
        include_report = data.get('include_report', True)
        include_price_analysis = data.get('include_price_analysis', True)
        
        if not tender_text:
            return jsonify({'error': 'Текст тендера не предоставлен'}), 400
        
        # 1. Извлечение параметров
        extract_start = time.time()
        tender_params = tender_analyzer.extract_tender_params(tender_text, source_type)
        completeness = tender_analyzer.analyze_tender_completeness(tender_params)
        extract_time = time.time() - extract_start
        
        # Отслеживаем метрики извлечения
        if metrics_tracker:
            metrics_tracker.track_extraction(tender_params, completeness, extract_time)
        
        result = {
            'tender_params': tender_params,
            'completeness': completeness,
            'success': True
        }
        
        # 2. Анализ цен
        if include_price_analysis and price_analyzer:
            try:
                price_analysis = price_analyzer.analyze_price(tender_params)
                result['price_analysis'] = price_analysis
            except Exception as e:
                logger.error(f"Ошибка при анализе цен: {e}")
                result['price_analysis'] = {'error': str(e)}
        
        # 3. Подбор поставщиков
        if include_suppliers:
            match_start = time.time()
            suppliers, summary = supplier_matcher.match_suppliers(tender_params)
            match_time = time.time() - match_start
            
            # Отслеживаем метрики подбора
            if metrics_tracker:
                metrics_tracker.track_supplier_match(len(suppliers), match_time)
            
            result['suppliers'] = suppliers
            result['suppliers_summary'] = summary
        else:
            suppliers = []
        
        # 4. Анализ рисков
        if include_risks:
            risk_start = time.time()
            risk_analysis = risk_analyzer.analyze_risks(tender_params)
            risk_time = time.time() - risk_start
            
            # Отслеживаем метрики анализа рисков
            if metrics_tracker:
                metrics_tracker.track_risk_analysis(risk_analysis, risk_time)
            
            result['risk_analysis'] = risk_analysis
        else:
            risk_analysis = {}
        
        # 5. Генерация отчёта
        if include_report:
            report_start = time.time()
            report = report_generator.generate_comprehensive_report(
                tender_params, suppliers, risk_analysis, completeness
            )
            report_time = time.time() - report_start
            
            # Отслеживаем метрики генерации отчёта
            if metrics_tracker:
                metrics_tracker.track_report_generation(report_time)
            
            result['report'] = report
        
        total_time = time.time() - start_time
        result['processing_time'] = round(total_time, 2)
        
        return jsonify(result)
    except Exception as e:
        if metrics_tracker:
            metrics_tracker.track_error('analyze_error', str(e))
        return jsonify({'error': str(e)}), 500

# Поддержка новичков - получение рейтинга
@app.route('/api/ai-procure/beginner/rating', methods=['GET'])
def api_procure_beginner_rating():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        user_id = session.get('username', 'anonymous')
        rating = beginner_support.get_beginner_rating(user_id)
        return jsonify({'rating': rating, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Поддержка новичков - обновление рейтинга
@app.route('/api/ai-procure/beginner/update-rating', methods=['POST'])
def api_procure_beginner_update_rating():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        action = data.get('action', 'participate')
        success = data.get('success', False)
        user_id = session.get('username', 'anonymous')
        
        beginner_support.update_rating(user_id, action, success)
        rating = beginner_support.get_beginner_rating(user_id)
        
        return jsonify({'rating': rating, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Поддержка новичков - создание тендера в песочнице
@app.route('/api/ai-procure/beginner/sandbox/create', methods=['POST'])
def api_procure_beginner_sandbox_create():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        tender_params = data.get('tender_params', {})
        user_id = session.get('username', 'anonymous')
        
        sandbox_tender = beginner_support.create_sandbox_tender(tender_params, user_id)
        return jsonify({'sandbox_tender': sandbox_tender, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Поддержка новичков - получение тендеров из песочницы
@app.route('/api/ai-procure/beginner/sandbox/list', methods=['GET'])
def api_procure_beginner_sandbox_list():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        user_id = session.get('username', 'anonymous')
        sandbox_tenders = beginner_support.get_sandbox_tenders(user_id)
        return jsonify({'sandbox_tenders': sandbox_tenders, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Поддержка новичков - отправка попытки в песочнице
@app.route('/api/ai-procure/beginner/sandbox/submit', methods=['POST'])
def api_procure_beginner_sandbox_submit():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        tender_id = data.get('tender_id')
        attempt_data = data.get('attempt_data', {})
        user_id = session.get('username', 'anonymous')
        
        if not tender_id:
            return jsonify({'error': 'ID тендера не указан'}), 400
        
        feedback = beginner_support.submit_sandbox_attempt(tender_id, user_id, attempt_data)
        return jsonify({'feedback': feedback, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Поддержка новичков - пошаговая инструкция
@app.route('/api/ai-procure/beginner/guide', methods=['POST'])
def api_procure_beginner_guide():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        tender_params = data.get('tender_params', {})
        step = data.get('step', 1)
        
        guide = report_generator.generate_beginner_guide(tender_params, step)
        return jsonify({'guide': guide, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Поддержка новичков - генерация документов
@app.route('/api/ai-procure/beginner/generate-documents', methods=['POST'])
def api_procure_beginner_generate_documents():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        tender_params = data.get('tender_params', {})
        document_type = data.get('document_type', 'application')
        
        document = beginner_support.generate_documents(tender_params, document_type)
        return jsonify({'document': document, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Поиск партнёрств для новичков
@app.route('/api/ai-procure/beginner/partnerships', methods=['POST'])
def api_procure_beginner_partnerships():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        beginner_supplier = data.get('beginner_supplier', {})
        tender_params = data.get('tender_params', {})
        
        if not tender_params:
            return jsonify({'error': 'Параметры тендера не предоставлены'}), 400
        
        partners, strategy = supplier_matcher.find_partnership_opportunities(beginner_supplier, tender_params)
        return jsonify({
            'partners': partners,
            'strategy': strategy,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ==================== Новые модули AI-Procure ====================

# Сбор данных из открытых источников
@app.route('/api/ai-procure/collect', methods=['POST'])
def api_procure_collect():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not data_collector:
        return jsonify({'error': 'DataCollector недоступен. Установите зависимости: pip install requests beautifulsoup4'}), 503
    
    try:
        data = request.json
        source_name = data.get('source', 'synthetic')
        filters = data.get('filters', {})
        
        tenders = data_collector.collect_from_source(source_name, filters)
        return jsonify({
            'tenders': tenders,
            'count': len(tenders),
            'source': source_name,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/collect/list', methods=['GET'])
def api_procure_collect_list():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not data_collector:
        return jsonify({'error': 'DataCollector недоступен'}), 503
    
    try:
        filters = request.args.to_dict()
        tenders = data_collector.get_collected_tenders(filters if filters else None)
        return jsonify({
            'tenders': tenders,
            'count': len(tenders),
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/parse-web', methods=['POST'])
def api_procure_parse_web():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not data_collector:
        return jsonify({'error': 'DataCollector недоступен'}), 503
    
    try:
        data = request.json
        url = data.get('url', '')
        
        if not url:
            return jsonify({'error': 'URL не предоставлен'}), 400
        
        text = data_collector.parse_web_page(url)
        if text:
            return jsonify({
                'text': text,
                'url': url,
                'success': True
            })
        else:
            return jsonify({'error': 'Не удалось извлечь текст со страницы'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Анализ цен
@app.route('/api/ai-procure/analyze-price', methods=['POST'])
def api_procure_analyze_price():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not price_analyzer:
        return jsonify({'error': 'PriceAnalyzer недоступен'}), 503
    
    try:
        data = request.json
        tender_params = data.get('tender_params', {})
        market_data = data.get('market_data')
        
        if not tender_params:
            return jsonify({'error': 'Параметры тендера не предоставлены'}), 400
        
        price_analysis = price_analyzer.analyze_price(tender_params, market_data)
        return jsonify({
            'price_analysis': price_analysis,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/compare-prices', methods=['POST'])
def api_procure_compare_prices():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not price_analyzer:
        return jsonify({'error': 'PriceAnalyzer недоступен'}), 503
    
    try:
        data = request.json
        tender_params = data.get('tender_params', {})
        similar_tenders = data.get('similar_tenders', [])
        
        if not tender_params:
            return jsonify({'error': 'Параметры тендера не предоставлены'}), 400
        
        comparison = price_analyzer.compare_with_similar_tenders(tender_params, similar_tenders)
        return jsonify({
            'comparison': comparison,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/price-statistics', methods=['GET'])
def api_procure_price_statistics():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not price_analyzer:
        return jsonify({'error': 'PriceAnalyzer недоступен'}), 503
    
    try:
        category = request.args.get('category')
        stats = price_analyzer.get_price_statistics(category)
        return jsonify({
            'statistics': stats,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Метрики
@app.route('/api/ai-procure/metrics/performance', methods=['GET'])
def api_procure_metrics_performance():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not metrics_tracker:
        return jsonify({'error': 'MetricsTracker недоступен'}), 503
    
    try:
        metrics = metrics_tracker.get_performance_metrics()
        return jsonify({
            'metrics': metrics,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/metrics/risks', methods=['GET'])
def api_procure_metrics_risks():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not metrics_tracker:
        return jsonify({'error': 'MetricsTracker недоступен'}), 503
    
    try:
        metrics = metrics_tracker.get_risk_detection_metrics()
        return jsonify({
            'metrics': metrics,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/metrics/usability', methods=['GET'])
def api_procure_metrics_usability():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not metrics_tracker:
        return jsonify({'error': 'MetricsTracker недоступен'}), 503
    
    try:
        metrics = metrics_tracker.get_usability_metrics()
        return jsonify({
            'metrics': metrics,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/metrics/daily', methods=['GET'])
def api_procure_metrics_daily():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not metrics_tracker:
        return jsonify({'error': 'MetricsTracker недоступен'}), 503
    
    try:
        days = int(request.args.get('days', 7))
        stats = metrics_tracker.get_daily_statistics(days)
        return jsonify({
            'statistics': stats,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/metrics/feedback', methods=['POST'])
def api_procure_metrics_feedback():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not metrics_tracker:
        return jsonify({'error': 'MetricsTracker недоступен'}), 503
    
    try:
        data = request.json
        feedback_type = data.get('type', 'general')
        feedback_data = data.get('data', {})
        
        metrics_tracker.add_user_feedback(feedback_type, feedback_data)
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# База данных поставщиков
@app.route('/api/ai-procure/suppliers/add', methods=['POST'])
def api_procure_suppliers_add():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not supplier_database:
        return jsonify({'error': 'SupplierDatabase недоступен'}), 503
    
    try:
        data = request.json
        supplier_data = data.get('supplier_data', {})
        
        if not supplier_data.get('name'):
            return jsonify({'error': 'Название поставщика обязательно'}), 400
        
        supplier_id = supplier_database.add_supplier(supplier_data)
        return jsonify({
            'supplier_id': supplier_id,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/suppliers/search', methods=['POST'])
def api_procure_suppliers_search():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not supplier_database:
        return jsonify({'error': 'SupplierDatabase недоступен'}), 503
    
    try:
        data = request.json
        criteria = data.get('criteria', {})
        
        suppliers = supplier_database.find_suppliers(criteria)
        return jsonify({
            'suppliers': suppliers,
            'count': len(suppliers),
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/suppliers/<supplier_id>', methods=['GET'])
def api_procure_suppliers_get(supplier_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not supplier_database:
        return jsonify({'error': 'SupplierDatabase недоступен'}), 503
    
    try:
        supplier = supplier_database.get_supplier(supplier_id)
        if supplier:
            return jsonify({
                'supplier': supplier,
                'success': True
            })
        else:
            return jsonify({'error': 'Поставщик не найден'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/suppliers/<supplier_id>/statistics', methods=['GET'])
def api_procure_suppliers_statistics(supplier_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not supplier_database:
        return jsonify({'error': 'SupplierDatabase недоступен'}), 503
    
    try:
        stats = supplier_database.get_supplier_statistics(supplier_id)
        if stats:
            return jsonify({
                'statistics': stats,
                'success': True
            })
        else:
            return jsonify({'error': 'Поставщик не найден'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/suppliers/list', methods=['GET'])
def api_procure_suppliers_list():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not supplier_database:
        return jsonify({'error': 'SupplierDatabase недоступен'}), 503
    
    try:
        suppliers = supplier_database.get_all_suppliers()
        return jsonify({
            'suppliers': suppliers,
            'count': len(suppliers),
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Парсинг документов
@app.route('/api/ai-procure/parse-document', methods=['POST'])
def api_procure_parse_document():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not document_parser:
        return jsonify({'error': 'DocumentParser недоступен'}), 503
    
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'Файл не предоставлен'}), 400
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': 'Файл не выбран'}), 400
        
        # Сохраняем временный файл
        with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(file.filename)[1]) as tmp_file:
            file.save(tmp_file.name)
            result = document_parser.parse_file(tmp_file.name)
            os.unlink(tmp_file.name)  # Удаляем временный файл
        
        return jsonify({
            'result': result,
            'success': result.get('success', False)
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/supported-formats', methods=['GET'])
def api_procure_supported_formats():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not document_parser:
        return jsonify({'error': 'DocumentParser недоступен'}), 503
    
    try:
        formats = document_parser.get_supported_formats()
        return jsonify({
            'formats': formats,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ==================== Новые расширенные модули AI-Procure ====================

# Уведомления
@app.route('/api/ai-procure/notifications/send', methods=['POST'])
def api_procure_notifications_send():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not notification_system:
        return jsonify({'error': 'NotificationSystem недоступен'}), 503
    
    try:
        data = request.json
        notification_type = data.get('type', 'email')
        to = data.get('to', '')
        subject = data.get('subject', '')
        body = data.get('body', '')
        
        if notification_type == 'email':
            result = notification_system.send_email(to, subject, body)
        elif notification_type == 'sms':
            result = notification_system.send_sms(to, body)
        else:
            return jsonify({'error': 'Неизвестный тип уведомления'}), 400
        
        return jsonify({'success': result})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/notifications/history', methods=['GET'])
def api_procure_notifications_history():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not notification_system:
        return jsonify({'error': 'NotificationSystem недоступен'}), 503
    
    try:
        limit = request.args.get('limit', 100, type=int)
        history = notification_system.get_notification_history(limit)
        return jsonify({'history': history, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Кэширование
@app.route('/api/ai-procure/cache/stats', methods=['GET'])
def api_procure_cache_stats():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not cache_manager:
        return jsonify({'error': 'CacheManager недоступен'}), 503
    
    try:
        stats = cache_manager.get_stats()
        return jsonify({'stats': stats, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/cache/clear', methods=['POST'])
def api_procure_cache_clear():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not cache_manager:
        return jsonify({'error': 'CacheManager недоступен'}), 503
    
    try:
        data = request.json or {}
        pattern = data.get('pattern')
        cache_key = data.get('cache_key')
        
        if pattern:
            cache_manager.invalidate(pattern=pattern)
        elif cache_key:
            cache_manager.invalidate(cache_key=cache_key)
        else:
            cache_manager.clear_all()
        
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Асинхронная обработка
@app.route('/api/ai-procure/async/submit', methods=['POST'])
def api_procure_async_submit():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not async_processor:
        return jsonify({'error': 'AsyncProcessor недоступен'}), 503
    
    try:
        data = request.json
        tender_text = data.get('tender_text', '')
        source_type = data.get('source_type', 'text')
        
        if not tender_text:
            return jsonify({'error': 'Текст тендера не предоставлен'}), 400
        
        # Создаём функцию для анализа
        def analyze_tender(text, src_type):
            params = tender_analyzer.extract_tender_params(text, src_type)
            if data.get('include_suppliers', False):
                suppliers, _ = supplier_matcher.match_suppliers(params)
                params['suppliers'] = suppliers
            if data.get('include_risks', False):
                risks = risk_analyzer.analyze_risks(params)
                params['risks'] = risks
            return params
        
        task_id = f"TASK_{datetime.now().strftime('%Y%m%d%H%M%S')}"
        async_processor.submit_task(task_id, analyze_tender, tender_text, source_type)
        
        return jsonify({'task_id': task_id, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/async/status/<task_id>', methods=['GET'])
def api_procure_async_status(task_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not async_processor:
        return jsonify({'error': 'AsyncProcessor недоступен'}), 503
    
    try:
        status = async_processor.get_task_status(task_id)
        if status:
            return jsonify({'status': status, 'success': True})
        else:
            return jsonify({'error': 'Задача не найдена'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Пакетная обработка
@app.route('/api/ai-procure/batch/process', methods=['POST'])
def api_procure_batch_process():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not batch_processor:
        return jsonify({'error': 'BatchProcessor недоступен'}), 503
    
    try:
        data = request.json
        tenders = data.get('tenders', [])
        async_mode = data.get('async', True)
        
        if not tenders:
            return jsonify({'error': 'Список тендеров не предоставлен'}), 400
        
        def analyze_tender(tender_text, source_type):
            params = tender_analyzer.extract_tender_params(tender_text, source_type)
            completeness = tender_analyzer.analyze_tender_completeness(params)
            return {'params': params, 'completeness': completeness}
        
        if async_mode:
            result = batch_processor.process_batch(tenders, analyze_tender)
        else:
            result = batch_processor.process_batch_sync(tenders, analyze_tender)
        
        return jsonify({'result': result, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/batch/status/<batch_id>', methods=['GET'])
def api_procure_batch_status(batch_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not batch_processor:
        return jsonify({'error': 'BatchProcessor недоступен'}), 503
    
    try:
        status = batch_processor.get_batch_status(batch_id)
        if status:
            return jsonify({'status': status, 'success': True})
        else:
            return jsonify({'error': 'Пакет не найден'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Валидация метрик
@app.route('/api/ai-procure/validation/add-reference', methods=['POST'])
def api_procure_validation_add_reference():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not metrics_validator:
        return jsonify({'error': 'MetricsValidator недоступен'}), 503
    
    try:
        data = request.json
        tender_id = data.get('tender_id', '')
        extracted_params = data.get('extracted_params', {})
        ground_truth = data.get('ground_truth', {})
        
        if not tender_id or not ground_truth:
            return jsonify({'error': 'Недостаточно данных'}), 400
        
        result = metrics_validator.add_reference_data(tender_id, extracted_params, ground_truth)
        return jsonify({'success': result})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/validation/validate/<tender_id>', methods=['POST'])
def api_procure_validation_validate(tender_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not metrics_validator:
        return jsonify({'error': 'MetricsValidator недоступен'}), 503
    
    try:
        data = request.json
        extracted_params = data.get('extracted_params', {})
        
        result = metrics_validator.validate_extraction(tender_id, extracted_params)
        return jsonify({'result': result, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/validation/statistics', methods=['GET'])
def api_procure_validation_statistics():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not metrics_validator:
        return jsonify({'error': 'MetricsValidator недоступен'}), 503
    
    try:
        stats = metrics_validator.get_validation_statistics()
        return jsonify({'statistics': stats, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Подтверждение рисков
@app.route('/api/ai-procure/risks/confirm', methods=['POST'])
def api_procure_risks_confirm():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not risk_confirmation:
        return jsonify({'error': 'RiskConfirmation недоступен'}), 503
    
    try:
        data = request.json
        tender_id = data.get('tender_id', '')
        risk_id = data.get('risk_id', '')
        confirmed = data.get('confirmed', False)
        user_feedback = data.get('user_feedback')
        evidence = data.get('evidence')
        
        if not tender_id or not risk_id:
            return jsonify({'error': 'Недостаточно данных'}), 400
        
        result = risk_confirmation.confirm_risk(tender_id, risk_id, confirmed, user_feedback, evidence)
        return jsonify({'success': result})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/risks/confirmation-statistics', methods=['GET'])
def api_procure_risks_confirmation_statistics():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not risk_confirmation:
        return jsonify({'error': 'RiskConfirmation недоступен'}), 503
    
    try:
        stats = risk_confirmation.get_confirmation_statistics()
        return jsonify({'statistics': stats, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Отслеживание предотвращённых случаев
@app.route('/api/ai-procure/prevention/track', methods=['POST'])
def api_procure_prevention_track():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not prevention_tracker:
        return jsonify({'error': 'PreventionTracker недоступен'}), 503
    
    try:
        data = request.json
        tender_id = data.get('tender_id', '')
        risk_analysis = data.get('risk_analysis', {})
        decision = data.get('decision', 'отказ_от_участия')
        reason = data.get('reason')
        estimated_loss = data.get('estimated_loss')
        
        if not tender_id or not risk_analysis:
            return jsonify({'error': 'Недостаточно данных'}), 400
        
        result = prevention_tracker.track_prevention(tender_id, risk_analysis, decision, reason, estimated_loss)
        return jsonify({'success': result})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/prevention/statistics', methods=['GET'])
def api_procure_prevention_statistics():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not prevention_tracker:
        return jsonify({'error': 'PreventionTracker недоступен'}), 503
    
    try:
        period_days = request.args.get('period_days', type=int)
        stats = prevention_tracker.get_prevention_statistics(period_days)
        return jsonify({'statistics': stats, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-procure/prevention/recent', methods=['GET'])
def api_procure_prevention_recent():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not prevention_tracker:
        return jsonify({'error': 'PreventionTracker недоступен'}), 503
    
    try:
        limit = request.args.get('limit', 10, type=int)
        recent = prevention_tracker.get_recent_preventions(limit)
        return jsonify({'preventions': recent, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Старый endpoint для обратной совместимости
@app.route('/api/ai-procure', methods=['POST'])
def api_ai_procure():
    """Старый endpoint, перенаправляет на комплексный анализ"""
    return api_procure_analyze()

# ==================== AI-Scrum Master API Endpoints ====================

# Управление проектами
@app.route('/api/ai-scrum/projects/create', methods=['POST'])
def api_scrum_create_project():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        project = project_manager.create_project(
            name=data.get('name'),
            description=data.get('description', ''),
            owner=session.get('username', 'unknown'),
            deadline=data.get('deadline')
        )
        return jsonify({'project': project, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/projects', methods=['GET'])
def api_scrum_get_projects():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        username = request.args.get('username')
        projects = project_manager.get_projects(username)
        return jsonify({'projects': projects, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/projects/<project_id>', methods=['GET'])
def api_scrum_get_project(project_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        project = project_manager.get_project(project_id)
        if not project:
            return jsonify({'error': 'Проект не найден'}), 404
        
        # Получаем дополнительную информацию
        team_members = project_manager.get_team_members(project_id)
        
        # Загружаем спринты проекта
        sprints = sprint_manager._load_sprints()
        project_sprints = [s for s in sprints if s.get("project_id") == project_id]
        
        # Загружаем задачи проекта (если есть)
        tasks = sprint_manager._load_tasks()
        project_tasks = [t for t in tasks if t.get("project_id") == project_id]
        
        # Статистика проекта
        stats = {
            "team_size": len(team_members),
            "sprints_count": len(project_sprints),
            "active_sprints": len([s for s in project_sprints if s.get("status") == "active"]),
            "tasks_count": len(project_tasks),
            "completed_tasks": len([t for t in project_tasks if t.get("status") == "completed"]),
            "in_progress_tasks": len([t for t in project_tasks if t.get("status") == "in_progress"])
        }
        
        # Анализ дедлайнов
        deadline_risks = deadline_tracker.analyze_risks(project_id)
        
        return jsonify({
            'project': project,
            'team_members': team_members,
            'sprints': project_sprints,
            'tasks': project_tasks[:20],  # Первые 20 задач
            'statistics': stats,
            'deadline_risks': deadline_risks,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Управление командами
@app.route('/api/ai-scrum/teams/create', methods=['POST'])
def api_scrum_create_team():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        team = project_manager.create_team(
            project_id=data.get('project_id'),
            team_name=data.get('team_name'),
            members=data.get('members', [])
        )
        return jsonify({'team': team, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/members/add', methods=['POST'])
def api_scrum_add_member():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        member = project_manager.add_member(
            project_id=data.get('project_id'),
            username=data.get('username'),
            role=data.get('role', 'developer'),
            workload=data.get('workload', 1.0)
        )
        return jsonify({'member': member, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Заполнение полей задачи через AI
@app.route('/api/ai-scrum/tasks/fill-fields', methods=['POST'])
def api_scrum_fill_task_fields():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        text = data.get('text', '')
        project_id = data.get('project_id')
        context = data.get('context')
        
        if not text:
            return jsonify({'error': 'Текст обязателен'}), 400
        
        # Получаем контекст проекта если project_id указан
        if project_id:
            try:
                projects_file = os.path.join('data', 'scrum_projects.json')
                if os.path.exists(projects_file):
                    with open(projects_file, 'r', encoding='utf-8') as f:
                        projects = json.load(f)
                        project = next((p for p in projects if p.get('project_id') == project_id), None)
                        if project:
                            context = {
                                'project_name': project.get('name', ''),
                                'project_description': project.get('description', ''),
                                'team_members': project.get('team_members', [])
                            }
            except Exception as e:
                print(f"Ошибка при загрузке контекста проекта: {e}")
        
        task_fields = task_creator.fill_task_fields(text, project_id, context)
        return jsonify({'task_fields': task_fields, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Создание задач из текста
@app.route('/api/ai-scrum/tasks/create-from-text', methods=['POST'])
def api_scrum_create_tasks_from_text():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        text = data.get('text', '')
        project_id = data.get('project_id')
        context = data.get('context')
        
        if not text or not project_id:
            return jsonify({'error': 'Текст и project_id обязательны'}), 400
        
        tasks_data = task_creator.create_tasks_from_text(text, project_id, context)
        return jsonify({'tasks_data': tasks_data, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/tasks/confirm', methods=['POST'])
def api_scrum_confirm_tasks():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        tasks_data = data.get('tasks_data', {})
        result = task_creator.confirm_and_save_tasks(tasks_data, False)
        return jsonify({'result': result, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Декомпозиция задач
@app.route('/api/ai-scrum/tasks/decompose', methods=['POST'])
def api_scrum_decompose_task():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        decomposition = task_decomposer.decompose_task(
            task=data.get('task', {}),
            team_info=data.get('team_info', {}),
            resources=data.get('resources', {}),
            calendar=data.get('calendar', {})
        )
        return jsonify({'decomposition': decomposition, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/tasks/optimize-allocation', methods=['POST'])
def api_scrum_optimize_allocation():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        result = task_decomposer.optimize_allocation(
            tasks=data.get('tasks', []),
            team=data.get('team', [])
        )
        return jsonify({'result': result, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Ассистент встреч
@app.route('/api/ai-scrum/meetings/analyze', methods=['POST'])
def api_scrum_analyze_meeting():
    """Анализирует транскрипт встречи"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        
        # Анализируем встречу
        analysis = meeting_assistant.analyze_meeting_transcript(
            transcript=data.get('transcript', ''),
            meeting_type=data.get('meeting_type', 'standup'),
            project_id=data.get('project_id'),
            participants=data.get('participants', [])
        )
        
        return jsonify({
            'analysis': analysis,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/meetings/generate-report', methods=['POST'])
def api_scrum_generate_meeting_report():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        analysis = data.get('analysis', {})
        report = meeting_assistant.generate_meeting_report(analysis)
        return jsonify({'report': report, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/meetings', methods=['GET'])
def api_scrum_get_meetings():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        project_id = request.args.get('project_id')
        meetings = meeting_assistant.get_meetings(project_id)
        return jsonify({'meetings': meetings, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Контроль дедлайнов
@app.route('/api/ai-scrum/deadlines/add', methods=['POST'])
def api_scrum_add_deadline():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        deadline = deadline_tracker.add_deadline(
            task_id=data.get('task_id'),
            deadline=data.get('deadline'),
            assignee=data.get('assignee'),
            project_id=data.get('project_id'),
            priority=data.get('priority', 'medium')
        )
        return jsonify({'deadline': deadline, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/deadlines/check', methods=['GET'])
def api_scrum_check_deadlines():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        days_ahead = request.args.get('days_ahead', 7, type=int)
        result = deadline_tracker.check_deadlines(days_ahead)
        return jsonify({'result': result, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/deadlines/reminders', methods=['GET'])
def api_scrum_get_reminders():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        assignee = request.args.get('assignee')
        reminders = deadline_tracker.get_reminders(assignee)
        return jsonify({'reminders': reminders, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/deadlines/risks', methods=['GET'])
def api_scrum_analyze_deadline_risks():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        project_id = request.args.get('project_id')
        risks = deadline_tracker.analyze_risks(project_id)
        return jsonify({'risks': risks, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Управление спринтами
@app.route('/api/ai-scrum/sprints/create', methods=['POST'])
def api_scrum_create_sprint():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        sprint = sprint_manager.create_sprint(
            project_id=data.get('project_id'),
            name=data.get('name'),
            start_date=data.get('start_date'),
            end_date=data.get('end_date'),
            goal=data.get('goal', '')
        )
        return jsonify({'sprint': sprint, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/sprints/<sprint_id>/backlog', methods=['POST'])
def api_scrum_build_backlog(sprint_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        backlog = sprint_manager.build_sprint_backlog(
            sprint_id=sprint_id,
            available_tasks=data.get('available_tasks', []),
            team_capacity=data.get('team_capacity', {})
        )
        return jsonify({'backlog': backlog, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/sprints/<sprint_id>/burndown', methods=['GET'])
def api_scrum_calculate_burndown(sprint_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        burndown = sprint_manager.calculate_burndown(sprint_id)
        return jsonify({'burndown': burndown, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/sprints/<sprint_id>/performance', methods=['GET'])
def api_scrum_analyze_sprint_performance(sprint_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        performance = sprint_manager.analyze_sprint_performance(sprint_id)
        return jsonify({'performance': performance, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Метрики
@app.route('/api/ai-scrum/metrics/performance', methods=['GET'])
def api_scrum_metrics_performance():
    """Получает Performance метрики"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        metrics = scrum_metrics_tracker.get_performance_metrics()
        return jsonify({'metrics': metrics, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/metrics/business', methods=['GET'])
def api_scrum_metrics_business():
    """Получает Business метрики"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        metrics = scrum_metrics_tracker.get_business_metrics()
        return jsonify({'metrics': metrics, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/metrics/usability', methods=['GET'])
def api_scrum_metrics_usability():
    """Получает Usability метрики"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        metrics = scrum_metrics_tracker.get_usability_metrics()
        return jsonify({'metrics': metrics, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/metrics/all', methods=['GET'])
def api_scrum_metrics_all():
    """Получает все метрики"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        metrics = scrum_metrics_tracker.get_all_metrics()
        return jsonify({'metrics': metrics, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Интеграция встреч с AI-ассистентом
@app.route('/api/ai-scrum/meetings/start-with-ai', methods=['POST'])
def api_scrum_start_meeting_with_ai():
    """Запускает встречу с полной интеграцией (подключение, транскрибация, AI-ассистент)"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        result = meeting_integration_manager.start_meeting_with_ai(
            meeting_id=data.get('meeting_id', ''),
            meeting_type=data.get('meeting_type', 'standup'),
            project_id=data.get('project_id'),
            participants=data.get('participants', []),
            provider=data.get('provider', 'google'),
            transcription_provider=data.get('transcription_provider', 'local'),
            meeting_goal=data.get('meeting_goal')
        )
        return jsonify(result)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/ai-scrum/meetings/<meeting_id>/ai-question', methods=['POST'])
def api_scrum_generate_ai_question(meeting_id: str):
    """Генерирует вопрос AI-ассистента для встречи"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        question = meeting_integration_manager.generate_ai_question(meeting_id)
        if question:
            return jsonify({
                'success': True,
                'question': question
            })
        else:
            return jsonify({
                'success': False,
                'message': 'Вопрос не нужен в данный момент'
            })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/ai-scrum/meetings/<meeting_id>/answer', methods=['POST'])
def api_scrum_answer_question(meeting_id: str):
    """Отвечает на вопрос участника встречи"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        question = data.get('question', '')
        
        if not question:
            return jsonify({'success': False, 'error': 'Вопрос не предоставлен'}), 400
        
        result = meeting_integration_manager.answer_question(meeting_id, question)
        return jsonify(result)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/ai-scrum/meetings/<meeting_id>/progress', methods=['GET'])
def api_scrum_get_meeting_progress(meeting_id: str):
    """Получает прогресс встречи"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        result = meeting_integration_manager.get_meeting_progress(meeting_id)
        return jsonify(result)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/ai-scrum/meetings/<meeting_id>/end', methods=['POST'])
def api_scrum_end_meeting(meeting_id: str):
    """Завершает встречу и возвращает сводку"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        result = meeting_integration_manager.end_meeting(meeting_id)
        return jsonify(result)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

# Аналитика
@app.route('/api/ai-scrum/analytics/insights', methods=['POST'])
def api_scrum_generate_insights():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        project_id = data.get('project_id')
        if not project_id:
            return jsonify({'error': 'project_id обязателен'}), 400
        
        # Если данные не предоставлены, аналитика загрузит их сама
        custom_data = data.get('data')
        insights = analytics_engine.generate_insights(
            project_id=project_id,
            data=custom_data
        )
        return jsonify({'insights': insights, 'success': True})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-scrum/analytics/predict-sprint', methods=['POST'])
def api_scrum_predict_sprint():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        prediction = analytics_engine.predict_sprint_completion(
            sprint_id=data.get('sprint_id'),
            current_data=data.get('current_data', {})
        )
        return jsonify({'prediction': prediction, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Интеграции
# Google Meet интеграция
@app.route('/api/ai-scrum/integrations/google-meet/settings', methods=['GET'])
def api_scrum_google_meet_settings():
    """Получает текущие настройки Google Meet"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    import os
    import json
    
    # Проверяем наличие библиотек
    try:
        from google.auth.transport.requests import Request
        from google.oauth2.credentials import Credentials
        from google_auth_oauthlib.flow import Flow
        from googleapiclient.discovery import build
        libraries_available = True
    except ImportError:
        libraries_available = False
    
    settings_file = os.path.join('data', 'google_meet_settings.json')
    if os.path.exists(settings_file):
        try:
            with open(settings_file, 'r', encoding='utf-8') as f:
                settings = json.load(f)
                
                # Проверяем наличие токена авторизации
                token_file = os.path.join('data', 'google_meet_token.json')
                is_authenticated = False
                if os.path.exists(token_file):
                    try:
                        # Инициализируем интеграцию для проверки
                        if settings.get('client_id') and settings.get('client_secret'):
                            if not meeting_integration_manager.google_meet:
                                meeting_integration_manager.initialize_google_meet(
                                    client_id=settings.get('client_id'),
                                    client_secret=settings.get('client_secret'),
                                    redirect_uri=settings.get('redirect_uri', 'http://localhost:5000/api/ai-scrum/integrations/google-meet/callback')
                                )
                            is_authenticated = meeting_integration_manager.google_meet.is_authenticated if meeting_integration_manager.google_meet else False
                    except Exception as e:
                        logger.error(f"Ошибка проверки авторизации: {e}")
                
                safe_settings = {
                    'client_id': settings.get('client_id', ''),
                    'redirect_uri': settings.get('redirect_uri', ''),
                    'enabled': settings.get('enabled', False),
                    'has_secret': bool(settings.get('client_secret')),
                    'connected': is_authenticated
                }
                return jsonify({
                    'settings': safe_settings, 
                    'success': True,
                    'libraries_available': libraries_available
                })
        except Exception as e:
            return jsonify({'error': str(e)}), 500
    
    return jsonify({
        'settings': {
            'client_id': '',
            'redirect_uri': '',
            'enabled': False,
            'has_secret': False,
            'connected': False
        },
        'success': True,
        'libraries_available': libraries_available
    })

@app.route('/api/ai-scrum/integrations/google-meet/settings', methods=['POST'])
def api_scrum_google_meet_save_settings():
    """Сохраняет настройки подключения к Google Meet"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        settings = {
            'client_id': data.get('client_id', ''),
            'client_secret': data.get('client_secret', ''),
            'redirect_uri': data.get('redirect_uri', ''),
            'enabled': data.get('enabled', True),
            'updated_at': datetime.now().isoformat()
        }
        
        settings_file = os.path.join('data', 'google_meet_settings.json')
        os.makedirs(os.path.dirname(settings_file), exist_ok=True)
        
        with open(settings_file, 'w', encoding='utf-8') as f:
            json.dump(settings, f, ensure_ascii=False, indent=2)
        
        # Инициализируем интеграцию
        if settings['client_id'] and settings['client_secret']:
            meeting_integration_manager.initialize_google_meet(
                client_id=settings['client_id'],
                client_secret=settings['client_secret'],
                redirect_uri=settings['redirect_uri']
            )
        
        return jsonify({
            'success': True,
            'message': 'Настройки Google Meet сохранены',
            'enabled': settings['enabled'],
            'connected': False
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/ai-scrum/integrations/google-meet/test', methods=['POST'])
def api_scrum_google_meet_test():
    """Тестирует подключение к Google Meet"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json or {}
        
        if data.get('client_id') and data.get('client_secret'):
            from scrum_models.google_meet_integration import GoogleMeetIntegration
            test_integration = GoogleMeetIntegration(
                client_id=data.get('client_id'),
                client_secret=data.get('client_secret'),
                redirect_uri=data.get('redirect_uri', '')
            )
            result = test_integration.test_connection()
        else:
            if not meeting_integration_manager.google_meet:
                return jsonify({
                    'success': False,
                    'error': 'Настройки Google Meet не указаны'
                }), 400
            
            result = meeting_integration_manager.google_meet.test_connection()
        
        return jsonify(result)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/ai-scrum/integrations/google-meet/authorize', methods=['GET'])
def api_scrum_google_meet_authorize():
    """Получает URL для OAuth авторизации Google Meet"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        # Загружаем настройки
        settings_file = os.path.join('data', 'google_meet_settings.json')
        if not os.path.exists(settings_file):
            return jsonify({
                'success': False,
                'error': 'Настройки Google Meet не найдены. Сначала настройте client_id и client_secret'
            }), 400
        
        with open(settings_file, 'r', encoding='utf-8') as f:
            settings = json.load(f)
        
        client_id = settings.get('client_id')
        client_secret = settings.get('client_secret')
        redirect_uri = settings.get('redirect_uri', 'http://localhost:5000/api/ai-scrum/integrations/google-meet/callback')
        
        if not client_id or not client_secret:
            return jsonify({
                'success': False,
                'error': 'OAuth credentials не настроены. Укажите client_id и client_secret в настройках'
            }), 400
        
        # Логируем настройки для отладки
        logger.info(f"Запрос авторизации Google Meet. Client ID: {client_id[:20]}..., Redirect URI: {redirect_uri}")
        
        # Инициализируем интеграцию если еще не инициализирована
        if not meeting_integration_manager.google_meet:
            meeting_integration_manager.initialize_google_meet(
                client_id=client_id,
                client_secret=client_secret,
                redirect_uri=redirect_uri
            )
        
        # Получаем URL авторизации
        result = meeting_integration_manager.google_meet.get_authorization_url()
        
        if result.get('success'):
            logger.info(f"URL авторизации создан успешно. State: {result.get('state', 'N/A')}")
        else:
            logger.error(f"Ошибка создания URL авторизации: {result.get('error')}")
        
        return jsonify(result)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/ai-scrum/integrations/google-meet/callback', methods=['GET'])
def api_scrum_google_meet_callback():
    """Обрабатывает OAuth callback от Google"""
    try:
        # Логируем все параметры для отладки
        logger.info(f"OAuth callback получен. Параметры: {dict(request.args)}")
        logger.info(f"URL: {request.url}")
        
        code = request.args.get('code')
        state = request.args.get('state')
        error = request.args.get('error')
        error_description = request.args.get('error_description', '')
        
        # Если есть ошибка от Google
        if error:
            error_msg = f'Ошибка авторизации от Google: {error}'
            if error_description:
                error_msg += f' - {error_description}'
            logger.error(error_msg)
            
            # Специальная обработка для redirect_uri_mismatch
            if error == 'redirect_uri_mismatch':
                # Загружаем настройки для получения правильного redirect_uri
                settings_file = 'data/google_meet_settings.json'
                redirect_uri = 'http://localhost:5000/api/ai-scrum/integrations/google-meet/callback'
                if os.path.exists(settings_file):
                    try:
                        with open(settings_file, 'r', encoding='utf-8') as f:
                            settings = json.load(f)
                            redirect_uri = settings.get('redirect_uri', redirect_uri)
                    except:
                        pass
                
                return f'''
                <!DOCTYPE html>
                <html>
                <head>
                    <title>Ошибка: Неверный Redirect URI</title>
                    <meta charset="utf-8">
                    <style>
                        body {{
                            font-family: Arial, sans-serif;
                            display: flex;
                            justify-content: center;
                            align-items: center;
                            height: 100vh;
                            margin: 0;
                            background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
                        }}
                        .container {{
                            background: white;
                            padding: 2rem;
                            border-radius: 10px;
                            box-shadow: 0 10px 25px rgba(0,0,0,0.2);
                            text-align: left;
                            max-width: 600px;
                        }}
                        h1 {{ color: #f5576c; text-align: center; }}
                        p {{ color: #666; }}
                        .error-details {{
                            background: #fee;
                            padding: 1.5rem;
                            border-radius: 5px;
                            margin-top: 1rem;
                            font-size: 0.9rem;
                            line-height: 1.8;
                        }}
                        code {{
                            background: #f3f4f6;
                            padding: 2px 6px;
                            border-radius: 3px;
                            font-family: monospace;
                            color: #dc2626;
                            font-weight: bold;
                        }}
                        .steps {{
                            margin-top: 1rem;
                        }}
                        .steps ol {{
                            margin: 0.5rem 0;
                            padding-left: 1.5rem;
                        }}
                        .steps li {{
                            margin: 0.5rem 0;
                        }}
                        a {{
                            color: #0284C7;
                            text-decoration: underline;
                        }}
                    </style>
                </head>
                <body>
                    <div class="container">
                        <h1>❌ Ошибка: redirect_uri_mismatch</h1>
                        <p style="text-align: center; color: #666;">Redirect URI в Google Cloud Console не совпадает с настройками приложения.</p>
                        <div class="error-details">
                            <strong style="color: #dc2626;">🔧 Как исправить:</strong>
                            <div class="steps">
                                <ol>
                                    <li>Откройте <a href="https://console.cloud.google.com/apis/credentials" target="_blank">Google Cloud Console → Credentials</a></li>
                                    <li>Найдите ваш OAuth 2.0 Client ID и нажмите на него для редактирования</li>
                                    <li>В разделе "Authorized redirect URIs" убедитесь, что добавлен <strong>точно такой же</strong> URI:</li>
                                </ol>
                                <div style="background: #f9fafb; padding: 1rem; border-radius: 5px; margin: 1rem 0; text-align: center;">
                                    <code>{redirect_uri}</code>
                                </div>
                                <ol start="4">
                                    <li>Если этого URI нет — добавьте его и нажмите "Save"</li>
                                    <li>Подождите 1-2 минуты (изменения могут применяться с задержкой)</li>
                                    <li>Попробуйте авторизоваться снова</li>
                                </ol>
                            </div>
                            <div style="margin-top: 1rem; padding-top: 1rem; border-top: 1px solid #ddd;">
                                <strong>⚠️ Важно:</strong> URI должен совпадать <strong>точно</strong>, включая:
                                <ul style="margin: 0.5rem 0; padding-left: 1.5rem;">
                                    <li>Протокол (http:// или https://)</li>
                                    <li>Домен (localhost:5000)</li>
                                    <li>Полный путь (/api/ai-scrum/integrations/google-meet/callback)</li>
                                </ul>
                            </div>
                        </div>
                    </div>
                </body>
                </html>
                ''', 400
            
            # Специальная обработка для access_denied (приложение в режиме тестирования)
            if error == 'access_denied':
                return f'''
                <!DOCTYPE html>
                <html>
                <head>
                    <title>Ошибка: Доступ заблокирован</title>
                    <meta charset="utf-8">
                    <style>
                        body {{
                            font-family: Arial, sans-serif;
                            display: flex;
                            justify-content: center;
                            align-items: center;
                            height: 100vh;
                            margin: 0;
                            background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
                        }}
                        .container {{
                            background: white;
                            padding: 2rem;
                            border-radius: 10px;
                            box-shadow: 0 10px 25px rgba(0,0,0,0.2);
                            text-align: left;
                            max-width: 650px;
                        }}
                        h1 {{ color: #f5576c; text-align: center; }}
                        p {{ color: #666; }}
                        .error-details {{
                            background: #fef3c7;
                            padding: 1.5rem;
                            border-radius: 5px;
                            margin-top: 1rem;
                            font-size: 0.9rem;
                            line-height: 1.8;
                            border-left: 4px solid #f59e0b;
                        }}
                        .steps {{
                            margin-top: 1rem;
                        }}
                        .steps ol {{
                            margin: 0.5rem 0;
                            padding-left: 1.5rem;
                        }}
                        .steps li {{
                            margin: 0.5rem 0;
                        }}
                        a {{
                            color: #0284C7;
                            text-decoration: underline;
                            font-weight: 600;
                        }}
                        .warning {{
                            background: #fee;
                            padding: 1rem;
                            border-radius: 5px;
                            margin-top: 1rem;
                            border-left: 4px solid #dc2626;
                        }}
                        code {{
                            background: #f3f4f6;
                            padding: 2px 6px;
                            border-radius: 3px;
                            font-family: monospace;
                            color: #dc2626;
                        }}
                    </style>
                </head>
                <body>
                    <div class="container">
                        <h1>⚠️ Доступ заблокирован</h1>
                        <p style="text-align: center; color: #666;">Приложение находится в режиме тестирования и доступно только для одобренных тестировщиков.</p>
                        <div class="error-details">
                            <strong style="color: #f59e0b;">🔧 Как исправить (для разработчика):</strong>
                            <div class="steps">
                                <ol>
                                    <li>Откройте <a href="https://console.cloud.google.com/apis/credentials/consent" target="_blank">Google Cloud Console → OAuth consent screen</a></li>
                                    <li>В разделе "<strong>Test users</strong>" (Тестовые пользователи) нажмите "<strong>+ ADD USERS</strong>"</li>
                                    <li>Добавьте email адрес пользователя, которому нужен доступ:<br>
                                        <code>sabithanbauyrzan@gmail.com</code></li>
                                    <li>Нажмите "<strong>ADD</strong>" и сохраните изменения</li>
                                    <li>Пользователь получит доступ в течение нескольких минут</li>
                                </ol>
                            </div>
                            <div class="warning" style="margin-top: 1rem;">
                                <strong>⚠️ Важно:</strong>
                                <ul style="margin: 0.5rem 0; padding-left: 1.5rem;">
                                    <li>В режиме тестирования можно добавить до <strong>100 тестовых пользователей</strong></li>
                                    <li>Для публичного доступа нужно отправить приложение на проверку Google (может занять несколько недель)</li>
                                    <li>Для внутреннего использования в организации можно выбрать тип "Internal" вместо "External"</li>
                                </ul>
                            </div>
                            <div style="margin-top: 1rem; padding-top: 1rem; border-top: 1px solid #ddd;">
                                <strong>📝 Альтернативный вариант:</strong><br>
                                Если вы хотите сделать приложение доступным для всех пользователей без проверки, измените тип приложения на "<strong>Internal</strong>" (только для Google Workspace организаций) или отправьте на проверку Google.
                            </div>
                        </div>
                    </div>
                </body>
                </html>
                ''', 403
            
            return f'''
            <!DOCTYPE html>
            <html>
            <head>
                <title>Ошибка авторизации</title>
                <meta charset="utf-8">
                <style>
                    body {{
                        font-family: Arial, sans-serif;
                        display: flex;
                        justify-content: center;
                        align-items: center;
                        height: 100vh;
                        margin: 0;
                        background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
                    }}
                    .container {{
                        background: white;
                        padding: 2rem;
                        border-radius: 10px;
                        box-shadow: 0 10px 25px rgba(0,0,0,0.2);
                        text-align: center;
                        max-width: 500px;
                    }}
                    h1 {{ color: #f5576c; }}
                    p {{ color: #666; }}
                    .error-details {{
                        background: #fee;
                        padding: 1rem;
                        border-radius: 5px;
                        margin-top: 1rem;
                        font-size: 0.9rem;
                    }}
                </style>
            </head>
            <body>
                <div class="container">
                    <h1>❌ Ошибка авторизации</h1>
                    <p>{error_msg}</p>
                    <div class="error-details">
                        <strong>Что делать:</strong><br>
                        1. Убедитесь, что redirect_uri в Google Cloud Console точно совпадает с настройками<br>
                        2. Проверьте, что приложение имеет доступ к Google Meet API<br>
                        3. Попробуйте авторизоваться снова
                    </div>
                </div>
            </body>
            </html>
            ''', 400
        
        # Проверяем наличие кода
        if not code:
            # Логируем все параметры для отладки
            all_params = dict(request.args)
            logger.warning(f"Authorization code не получен. Все параметры: {all_params}")
            
            return f'''
            <!DOCTYPE html>
            <html>
            <head>
                <title>Ошибка авторизации</title>
                <meta charset="utf-8">
                <style>
                    body {{
                        font-family: Arial, sans-serif;
                        display: flex;
                        justify-content: center;
                        align-items: center;
                        height: 100vh;
                        margin: 0;
                        background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
                    }}
                    .container {{
                        background: white;
                        padding: 2rem;
                        border-radius: 10px;
                        box-shadow: 0 10px 25px rgba(0,0,0,0.2);
                        text-align: center;
                        max-width: 600px;
                    }}
                    h1 {{ color: #f5576c; }}
                    p {{ color: #666; }}
                    .error-details {{
                        background: #fee;
                        padding: 1rem;
                        border-radius: 5px;
                        margin-top: 1rem;
                        font-size: 0.9rem;
                        text-align: left;
                    }}
                    code {{
                        background: #f5f5f5;
                        padding: 2px 6px;
                        border-radius: 3px;
                        font-family: monospace;
                    }}
                </style>
            </head>
            <body>
                <div class="container">
                    <h1>❌ Authorization code не получен</h1>
                    <p>Google не вернул код авторизации в URL.</p>
                    <div class="error-details">
                        <strong>Возможные причины:</strong><br>
                        1. <strong>Неверный redirect_uri</strong> - убедитесь, что в Google Cloud Console указан точно такой же redirect_uri:<br>
                        &nbsp;&nbsp;&nbsp;&nbsp;<code>http://localhost:5000/api/ai-scrum/integrations/google-meet/callback</code><br><br>
                        2. <strong>Проблема с настройками OAuth</strong> - проверьте настройки OAuth 2.0 Client ID в Google Cloud Console<br><br>
                        3. <strong>Пользователь отменил авторизацию</strong> - попробуйте авторизоваться снова<br><br>
                        <strong>Полученные параметры:</strong><br>
                        {all_params if all_params else 'Параметры отсутствуют'}
                    </div>
                </div>
            </body>
            </html>
            ''', 400
        
        # Загружаем настройки
        settings_file = os.path.join('data', 'google_meet_settings.json')
        if not os.path.exists(settings_file):
            return jsonify({
                'success': False,
                'error': 'Настройки Google Meet не найдены'
            }), 400
        
        with open(settings_file, 'r', encoding='utf-8') as f:
            settings = json.load(f)
        
        client_id = settings.get('client_id')
        client_secret = settings.get('client_secret')
        redirect_uri = settings.get('redirect_uri', 'http://localhost:5000/api/ai-scrum/integrations/google-meet/callback')
        
        # Инициализируем интеграцию если еще не инициализирована
        if not meeting_integration_manager.google_meet:
            meeting_integration_manager.initialize_google_meet(
                client_id=client_id,
                client_secret=client_secret,
                redirect_uri=redirect_uri
            )
        
        # Обрабатываем callback
        result = meeting_integration_manager.google_meet.handle_oauth_callback(code, state)
        
        # Возвращаем HTML страницу с результатом
        if result.get('success'):
            return '''
            <!DOCTYPE html>
            <html>
            <head>
                <title>Авторизация успешна</title>
                <meta charset="utf-8">
                <style>
                    body {
                        font-family: Arial, sans-serif;
                        display: flex;
                        justify-content: center;
                        align-items: center;
                        height: 100vh;
                        margin: 0;
                        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
                    }
                    .container {
                        background: white;
                        padding: 2rem;
                        border-radius: 10px;
                        box-shadow: 0 10px 25px rgba(0,0,0,0.2);
                        text-align: center;
                    }
                    h1 { color: #4CAF50; }
                    p { color: #666; }
                </style>
            </head>
            <body>
                <div class="container">
                    <h1>✅ Авторизация успешна!</h1>
                    <p>Вы успешно авторизованы в Google Meet API.</p>
                    <p>Это окно можно закрыть.</p>
                </div>
                <script>
                    setTimeout(function() {
                        window.close();
                    }, 2000);
                </script>
            </body>
            </html>
            '''
        else:
            return f'''
            <!DOCTYPE html>
            <html>
            <head>
                <title>Ошибка авторизации</title>
                <meta charset="utf-8">
                <style>
                    body {{
                        font-family: Arial, sans-serif;
                        display: flex;
                        justify-content: center;
                        align-items: center;
                        height: 100vh;
                        margin: 0;
                        background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
                    }}
                    .container {{
                        background: white;
                        padding: 2rem;
                        border-radius: 10px;
                        box-shadow: 0 10px 25px rgba(0,0,0,0.2);
                        text-align: center;
                    }}
                    h1 {{ color: #f5576c; }}
                    p {{ color: #666; }}
                </style>
            </head>
            <body>
                <div class="container">
                    <h1>❌ Ошибка авторизации</h1>
                    <p>{result.get('error', 'Неизвестная ошибка')}</p>
                </div>
            </body>
            </html>
            ''', 400
    except Exception as e:
        return f'''
        <!DOCTYPE html>
        <html>
        <head>
            <title>Ошибка</title>
            <meta charset="utf-8">
            <style>
                body {{
                    font-family: Arial, sans-serif;
                    display: flex;
                    justify-content: center;
                    align-items: center;
                    height: 100vh;
                    margin: 0;
                    background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
                }}
                .container {{
                    background: white;
                    padding: 2rem;
                    border-radius: 10px;
                    box-shadow: 0 10px 25px rgba(0,0,0,0.2);
                    text-align: center;
                }}
                h1 {{ color: #f5576c; }}
                p {{ color: #666; }}
            </style>
        </head>
        <body>
            <div class="container">
                <h1>❌ Ошибка</h1>
                <p>{str(e)}</p>
            </div>
        </body>
        </html>
        ''', 500

# Старый endpoint для обратной совместимости
@app.route('/api/ai-scrum', methods=['POST'])
def api_ai_scrum():
    """Старый endpoint, перенаправляет на создание задач из текста"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    data = request.json
    query = data.get('query', '')
    
    if not query:
        return jsonify({'error': 'Запрос не предоставлен'}), 400
    
    # Используем новый функционал создания задач
    try:
        # Создаём временный проект если нужно
        project = project_manager.create_project(
            name="Временный проект",
            description="Создан из запроса",
            owner=session.get('username', 'unknown')
        )
        
        tasks_data = task_creator.create_tasks_from_text(query, project["project_id"])
        return jsonify({
            'result': 'Задачи созданы из текста',
            'tasks_data': tasks_data,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ==================== AI-Business Analyst API Endpoints ====================

# Анализ документа
@app.route('/api/ai-business-analyst/analyze-document', methods=['POST'])
def api_ba_analyze_document():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        document_text = data.get('document_text', '')
        document_type = data.get('document_type')
        
        if not document_text:
            return jsonify({'error': 'Текст документа не предоставлен'}), 400
        
        analysis = document_analyzer.analyze_document(document_text, document_type)
        return jsonify({
            'analysis': analysis,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Извлечение требований
@app.route('/api/ai-business-analyst/extract-requirements', methods=['POST'])
def api_ba_extract_requirements():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        document_text = data.get('document_text', '')
        context = data.get('context', {})
        
        if not document_text:
            return jsonify({'error': 'Текст документа не предоставлен'}), 400
        
        requirements = requirement_extractor.extract_requirements(document_text, context)
        
        # Если есть анализ, находим пробелы
        analysis = data.get('analysis', {})
        if analysis:
            gaps = requirement_extractor.identify_gaps(requirements, analysis)
            requirements['gaps'] = gaps
        
        return jsonify({
            'requirements': requirements,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Генерация BRD
@app.route('/api/ai-business-analyst/generate-brd', methods=['POST'])
def api_ba_generate_brd():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        analysis = data.get('analysis', {})
        context = data.get('context', {})
        export_format = data.get('export_format', None)  # 'docx', 'doc', 'pdf', или None для текста
        
        if not analysis:
            return jsonify({'error': 'Анализ документа не предоставлен'}), 400
        
        brd = document_generator.generate_brd(analysis, context)
        
        # Extract BRD content
        brd_content = ''
        if isinstance(brd, dict):
            if 'error' in brd:
                return jsonify({
                    'error': brd.get('error', 'Ошибка генерации BRD'),
                    'success': False
                }), 500
            else:
                brd_content = brd.get('content', str(brd))
        else:
            brd_content = str(brd)
        
        # Если запрошен экспорт в файл
        if export_format:
            os.makedirs('uploads', exist_ok=True)
            
            # 'doc' и 'docx' обрабатываются одинаково (используем docx формат)
            if export_format in ['docx', 'doc']:
                try:
                    from docx import Document
                    doc = Document()
                    
                    # Парсим Markdown и добавляем в документ
                    lines = brd_content.split('\n')
                    for line in lines:
                        if line.startswith('# '):
                            doc.add_heading(line[2:], level=1)
                        elif line.startswith('## '):
                            doc.add_heading(line[3:], level=2)
                        elif line.startswith('### '):
                            doc.add_heading(line[4:], level=3)
                        elif line.startswith('- ') or line.startswith('* '):
                            doc.add_paragraph(line[2:], style='List Bullet')
                        elif line.strip():
                            doc.add_paragraph(line)
                    
                    # Используем .docx для обоих форматов (python-docx не поддерживает старый .doc)
                    file_ext = '.docx' if export_format == 'docx' else '.docx'
                    filename = f'BRD_{datetime.now().strftime("%Y%m%d_%H%M%S")}{file_ext}'
                    filepath = os.path.join('uploads', filename)
                    doc.save(filepath)
                    
                    return jsonify({
                        'brd': {'content': brd_content, 'type': 'BRD'},
                        'filename': filename,
                        'filepath': filepath,
                        'format': export_format,
                        'success': True
                    })
                except ImportError:
                    return jsonify({'error': 'Библиотека python-docx не установлена'}), 500
            
            elif export_format == 'pdf':
                try:
                    from reportlab.lib.pagesizes import letter, A4
                    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
                    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
                    from reportlab.lib.units import inch
                    import re
                    
                    filename = f'BRD_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf'
                    filepath = os.path.join('uploads', filename)
                    
                    doc = SimpleDocTemplate(filepath, pagesize=A4)
                    styles = getSampleStyleSheet()
                    story = []
                    
                    # Создаем кастомные стили
                    title_style = ParagraphStyle(
                        'CustomTitle',
                        parent=styles['Heading1'],
                        fontSize=18,
                        textColor='#8B1E2D',
                        spaceAfter=12,
                    )
                    heading1_style = ParagraphStyle(
                        'CustomHeading1',
                        parent=styles['Heading1'],
                        fontSize=16,
                        textColor='#8B1E2D',
                        spaceAfter=10,
                    )
                    heading2_style = ParagraphStyle(
                        'CustomHeading2',
                        parent=styles['Heading2'],
                        fontSize=14,
                        textColor='#A52A2A',
                        spaceAfter=8,
                    )
                    
                    # Парсим Markdown
                    lines = brd_content.split('\n')
                    for line in lines:
                        if line.startswith('# '):
                            story.append(Paragraph(line[2:], title_style))
                            story.append(Spacer(1, 0.2*inch))
                        elif line.startswith('## '):
                            story.append(Paragraph(line[3:], heading1_style))
                            story.append(Spacer(1, 0.15*inch))
                        elif line.startswith('### '):
                            story.append(Paragraph(line[4:], heading2_style))
                            story.append(Spacer(1, 0.1*inch))
                        elif line.startswith('- ') or line.startswith('* '):
                            story.append(Paragraph('• ' + line[2:], styles['Normal']))
                            story.append(Spacer(1, 0.05*inch))
                        elif line.strip():
                            story.append(Paragraph(line, styles['Normal']))
                            story.append(Spacer(1, 0.1*inch))
                    
                    doc.build(story)
                    
                    return jsonify({
                        'brd': {'content': brd_content, 'type': 'BRD'},
                        'filename': filename,
                        'filepath': filepath,
                        'format': export_format,
                        'success': True
                    })
                except ImportError:
                    return jsonify({'error': 'Библиотека reportlab не установлена. Установите: pip install reportlab'}), 500
                except Exception as e:
                    import traceback
                    traceback.print_exc()
                    return jsonify({'error': f'Ошибка создания PDF: {str(e)}'}), 500
        
        # Возвращаем только текст BRD
        return jsonify({
            'brd': {'content': brd_content, 'type': 'BRD'},
            'success': True
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': f'Ошибка генерации BRD: {str(e)}'}), 500

# Генерация Use Case
@app.route('/api/ai-business-analyst/generate-use-case', methods=['POST'])
def api_ba_generate_use_case():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        if not data:
            return jsonify({'error': 'Данные не предоставлены'}), 400
        
        requirement = data.get('requirement', '')
        actor = data.get('actor', 'Пользователь')
        context = data.get('context', {})
        
        if not requirement or not requirement.strip():
            return jsonify({'error': 'Требование не предоставлено или пустое'}), 400
        
        use_case = document_generator.generate_use_case(requirement, actor, context)
        
        # Ensure use_case is a dict with proper structure
        if isinstance(use_case, dict):
            if 'error' in use_case:
                return jsonify({'error': use_case.get('error', 'Ошибка генерации Use Case')}), 500
        else:
            # If it's not a dict, wrap it
            use_case = {
                'type': 'Use Case',
                'content': str(use_case),
                'actor': actor
            }
        
        return jsonify({
            'use_case': use_case,
            'success': True
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': f'Ошибка генерации Use Case: {str(e)}'}), 500

# Генерация User Stories
@app.route('/api/ai-business-analyst/generate-user-stories', methods=['POST'])
def api_ba_generate_user_stories():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        requirements = data.get('requirements', [])
        roles = data.get('roles', [])
        
        if not requirements:
            return jsonify({'error': 'Требования не предоставлены'}), 400
        
        user_stories = document_generator.generate_user_stories(requirements, roles)
        return jsonify({
            'user_stories': user_stories,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Генерация страницы для Confluence
@app.route('/api/ai-business-analyst/generate-confluence', methods=['POST'])
def api_ba_generate_confluence():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        documents = data.get('documents', [])
        title = data.get('title', 'Бизнес-анализ')
        
        if not documents:
            return jsonify({'error': 'Документы не предоставлены'}), 400
        
        confluence_page = document_generator.generate_confluence_page(documents, title)
        return jsonify({
            'confluence_page': confluence_page,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

def generate_bpmn_xml(bpmn_data):
    """Генерирует BPMN XML из структурированных данных"""
    process_name = bpmn_data.get('process_name', 'Business Process')
    tasks = bpmn_data.get('tasks', [])
    gateways = bpmn_data.get('gateways', [])
    start_event = bpmn_data.get('start_event', {'id': 'start', 'label': 'Start'})
    end_events = bpmn_data.get('end_events', [{'id': 'end', 'label': 'End'}])
    flows = bpmn_data.get('flows', [])
    
    # Генерируем простой BPMN XML
    xml = f'''<?xml version="1.0" encoding="UTF-8"?>
<bpmn:definitions xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL"
                  xmlns:bpmndi="http://www.omg.org/spec/BPMN/20100524/DI"
                  xmlns:dc="http://www.omg.org/spec/DD/20100524/DC"
                  xmlns:di="http://www.omg.org/spec/DD/20100524/DI"
                  id="Definitions_1"
                  targetNamespace="http://bpmn.io/schema/bpmn">
  <bpmn:process id="Process_1" name="{process_name}" isExecutable="false">
    <bpmn:startEvent id="{start_event['id']}" name="{start_event.get('label', 'Start')}" />
    {''.join([f'<bpmn:task id="{task["id"]}" name="{task.get("label", "")}" />' for task in tasks])}
    {''.join([f'<bpmn:exclusiveGateway id="{gw["id"]}" name="{gw.get("label", "")}" />' for gw in gateways])}
    {''.join([f'<bpmn:endEvent id="{evt["id"]}" name="{evt.get("label", "End")}" />' for evt in end_events])}
    {''.join([f'<bpmn:sequenceFlow id="flow_{i}" sourceRef="{flow["from"]}" targetRef="{flow["to"]}" name="{flow.get("label", "")}" />' for i, flow in enumerate(flows)])}
  </bpmn:process>
</bpmn:definitions>'''
    
    return xml

# Генерация BPMN диаграммы
@app.route('/api/ai-business-analyst/generate-bpmn', methods=['POST'])
def api_ba_generate_bpmn():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        process_description = data.get('process_description', '')
        use_cases = data.get('use_cases', [])
        
        if not process_description and not use_cases:
            return jsonify({'error': 'Описание процесса или Use Cases не предоставлены'}), 400
        
        # Генерируем структурированное описание процесса для визуализации
        prompt = f"""Проанализируй следующий бизнес-процесс и создай структурированное описание для BPMN диаграммы:

Описание процесса: {process_description}

Use Cases: {use_cases}

Верни результат в формате JSON со следующей структурой:
{{
    "process_name": "название процесса",
    "start_event": {{"id": "start", "label": "Начало процесса"}},
    "tasks": [
        {{"id": "task1", "label": "Название задачи", "type": "task"}},
        {{"id": "task2", "label": "Название задачи", "type": "task"}}
    ],
    "gateways": [
        {{"id": "gateway1", "label": "Условие", "type": "exclusive"}}
    ],
    "end_events": [
        {{"id": "end1", "label": "Конец процесса"}}
    ],
    "flows": [
        {{"from": "start", "to": "task1", "label": ""}},
        {{"from": "task1", "to": "gateway1", "label": ""}},
        {{"from": "gateway1", "to": "task2", "label": "Да"}},
        {{"from": "gateway1", "to": "end1", "label": "Нет"}},
        {{"from": "task2", "to": "end1", "label": ""}}
    ]
}}

Важно: верни ТОЛЬКО валидный JSON, без дополнительного текста."""
        
        try:
            response = model.generate_content(prompt)
            bpmn_json_text = response.text if hasattr(response, 'text') else str(response)
            
            # Извлекаем JSON из ответа
            import re
            json_match = re.search(r'\{[\s\S]*\}', bpmn_json_text)
            if json_match:
                bpmn_data = json.loads(json_match.group(0))
                
                # Генерируем BPMN XML для визуализации
                bpmn_xml = generate_bpmn_xml(bpmn_data)
                
                return jsonify({
                    'bpmn': bpmn_xml,
                    'bpmn_data': bpmn_data,
                    'success': True
                })
            else:
                # Fallback на текстовое описание
                return jsonify({
                    'bpmn': bpmn_json_text,
                    'success': True
                })
        except Exception as e:
            logger.error(f"Ошибка при генерации BPMN: {e}")
            # Fallback на простое описание
            return jsonify({
                'bpmn': f"Ошибка генерации структурированного BPMN: {str(e)}",
                'success': False
            })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Экспорт BRD в Word/PDF
@app.route('/api/ai-business-analyst/export-brd', methods=['POST'])
def api_ba_export_brd():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        brd_content = data.get('brd_content', '')
        format_type = data.get('format', 'word')  # 'word' or 'pdf'
        
        if not brd_content:
            return jsonify({'error': 'Содержимое BRD не предоставлено'}), 400
        
        # Для Word формата
        if format_type == 'word':
            try:
                from docx import Document
                doc = Document()
                # Парсим Markdown и добавляем в документ
                lines = brd_content.split('\n')
                for line in lines:
                    if line.startswith('# '):
                        doc.add_heading(line[2:], level=1)
                    elif line.startswith('## '):
                        doc.add_heading(line[3:], level=2)
                    elif line.startswith('### '):
                        doc.add_heading(line[4:], level=3)
                    elif line.startswith('- ') or line.startswith('* '):
                        doc.add_paragraph(line[2:], style='List Bullet')
                    elif line.strip():
                        doc.add_paragraph(line)
                
                # Сохраняем временный файл
                filename = f'BRD_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx'
                filepath = os.path.join('uploads', filename)
                os.makedirs('uploads', exist_ok=True)
                doc.save(filepath)
                
                return jsonify({
                    'filename': filename,
                    'filepath': filepath,
                    'success': True
                })
            except ImportError:
                return jsonify({'error': 'Библиотека python-docx не установлена'}), 500
        
        # Для PDF формата (простой вариант - возвращаем текст)
        elif format_type == 'pdf':
            return jsonify({
                'content': brd_content,
                'format': 'pdf',
                'success': True,
                'message': 'PDF экспорт будет реализован позже. Используйте Word формат.'
            })
        
        return jsonify({'error': 'Неподдерживаемый формат'}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Публикация в Confluence (локальное хранение)
@app.route('/api/ai-business-analyst/publish-confluence', methods=['POST'])
def api_ba_publish_confluence():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        import json
        data = request.json
        if not data:
            return jsonify({'error': 'Данные не предоставлены'}), 400
        
        page_content = data.get('page_content', '')
        page_title = data.get('page_title', 'Business Requirements Document')
        space_key = data.get('space_key', 'BA')
        brd_content = data.get('brd_content', '')
        
        # Используем brd_content если page_content пустой
        final_content = page_content or brd_content
        
        if not final_content or len(final_content.strip()) < 10:
            return jsonify({'error': 'Содержимое страницы не предоставлено или слишком короткое'}), 400
        
        if not page_title or not page_title.strip():
            page_title = 'Business Requirements Document'
        
        if not space_key or not space_key.strip():
            space_key = 'BA'
        
        # Сохраняем документ в локальный файл
        confluence_dir = 'confluence_docs'
        os.makedirs(confluence_dir, exist_ok=True)
        
        # Создаем уникальный ID для документа
        doc_id = f"{space_key}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        
        # Сохраняем документ
        doc_data = {
            'id': doc_id,
            'title': page_title,
            'space_key': space_key,
            'content': final_content,
            'created_at': datetime.now().isoformat(),
            'created_by': session.get('username', 'unknown'),
            'type': 'BRD'
        }
        
        # Сохраняем в JSON файл
        doc_file = os.path.join(confluence_dir, f'{doc_id}.json')
        with open(doc_file, 'w', encoding='utf-8') as f:
            json.dump(doc_data, f, ensure_ascii=False, indent=2)
        
        # Обновляем индекс всех документов
        index_file = os.path.join(confluence_dir, 'index.json')
        documents = []
        if os.path.exists(index_file):
            try:
                with open(index_file, 'r', encoding='utf-8') as f:
                    documents = json.load(f)
            except:
                documents = []
        
        # Добавляем новый документ в начало списка
        documents.insert(0, {
            'id': doc_id,
            'title': page_title.strip(),
            'space_key': space_key.strip(),
            'created_at': doc_data['created_at'],
            'created_by': doc_data['created_by'],
            'type': doc_data['type']
        })
        
        # Сохраняем индекс
        with open(index_file, 'w', encoding='utf-8') as f:
            json.dump(documents, f, ensure_ascii=False, indent=2)
        
        return jsonify({
            'success': True,
            'message': f'Страница "{page_title}" успешно опубликована в локальное хранилище Confluence (пространство: {space_key})',
            'doc_id': doc_id,
            'page_url': f'/confluence-docs?doc={doc_id}'
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# Комплексный анализ документа (все функции сразу)
@app.route('/api/ai-business-analyst/analyze', methods=['POST'])
def api_ba_analyze():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        document_text = data.get('document_text', '')
        document_type = data.get('document_type')
        context = data.get('context', {})
        generate_documents = data.get('generate_documents', False)
        
        if not document_text:
            return jsonify({'error': 'Текст документа не предоставлен'}), 400
        
        # 1. Анализ документа
        analysis = document_analyzer.analyze_document(document_text, document_type)
        
        # 2. Извлечение требований
        requirements = requirement_extractor.extract_requirements(document_text, context)
        gaps = requirement_extractor.identify_gaps(requirements, analysis)
        requirements['gaps'] = gaps
        
        result = {
            'analysis': analysis,
            'requirements': requirements,
            'success': True
        }
        
        # 3. Генерация документов (если запрошено)
        if generate_documents:
            # BRD
            try:
                brd = document_generator.generate_brd(analysis, context)
                # Ensure BRD is properly formatted
                if isinstance(brd, dict):
                    if 'error' in brd:
                        result['brd_error'] = brd.get('error', 'Ошибка генерации BRD')
                    else:
                        result['brd'] = brd
                else:
                    result['brd'] = {'content': str(brd), 'type': 'BRD'}
            except Exception as e:
                import traceback
                traceback.print_exc()
                result['brd_error'] = str(e)
            
            # User Stories (если есть функциональные требования)
            functional_reqs = requirements.get('functional', [])
            if functional_reqs:
                user_stories = document_generator.generate_user_stories(
                    functional_reqs[:5],  # Первые 5 требований
                    context.get('roles', [])
                )
                result['user_stories'] = user_stories
        
        return jsonify(result)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# Загрузка и обработка файлов
@app.route('/api/ai-business-analyst/upload-file', methods=['POST'])
def api_ba_upload_file():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'Файл не предоставлен'}), 400
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': 'Файл не выбран'}), 400
        
        filename = secure_filename(file.filename)
        file_ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
        
        # Извлекаем текст из файла
        text = ''
        if file_ext == 'txt':
            text = file.read().decode('utf-8', errors='ignore')
        elif file_ext == 'pdf':
            try:
                try:
                    import PyPDF2  # type: ignore # noqa: F401
                except ImportError:
                    return jsonify({'error': 'Библиотека PyPDF2 не установлена. Установите: pip install PyPDF2'}), 500
                pdf_reader = PyPDF2.PdfReader(file)
                text = ''
                for page in pdf_reader.pages:
                    text += page.extract_text() + '\n'
            except Exception as e:
                return jsonify({'error': f'Ошибка чтения PDF: {str(e)}'}), 500
        elif file_ext in ['doc', 'docx']:
            try:
                try:
                    from docx import Document  # type: ignore
                except ImportError:
                    return jsonify({'error': 'Библиотека python-docx не установлена. Установите: pip install python-docx'}), 500
                doc = Document(file)
                text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
            except Exception as e:
                return jsonify({'error': f'Ошибка чтения DOCX: {str(e)}'}), 500
        else:
            return jsonify({'error': f'Неподдерживаемый формат файла: {file_ext}'}), 400
        
        if not text.strip():
            return jsonify({'error': 'Не удалось извлечь текст из файла'}), 400
        
        # Автоматически анализируем документ
        analysis = None
        try:
            if len(text) > 100:  # Анализируем только если текст достаточно большой
                analysis = document_analyzer.analyze_document(text[:50000])  # Ограничиваем размер
        except Exception as e:
            logger.error(f"Ошибка при анализе документа: {e}")
        
        return jsonify({
            'text': text,
            'analysis': analysis,
            'filename': filename,
            'filename': filename,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Вспомогательные функции для умного чата
def _extract_context_from_history(chat_history: List) -> str:
    """Извлекает ключевую информацию из истории чата"""
    if not chat_history:
        return ""
    
    context = "Уже обсуждено в разговоре:\n"
    keywords = []
    
    for msg in chat_history[-5:]:
        if msg.get('role') == 'user':
            content = msg.get('content', '').lower()
            if 'проект' in content:
                keywords.append('проект')
            if 'требование' in content or 'требования' in content:
                keywords.append('требования')
            if 'система' in content:
                keywords.append('система')
            if 'цель' in content:
                keywords.append('цель')
    
    if keywords:
        context += f"- Обсуждались: {', '.join(set(keywords))}\n"
    
    return context

def _analyze_message_intelligently(message: str, model, chat_history: List) -> str:
    """Умный анализ сообщения пользователя"""
    try:
        prompt = f"""Проанализируй сообщение пользователя:

Сообщение: "{message}"

Определи тип запроса и ключевые темы. Краткий анализ (1-2 предложения):"""
        
        response = model.generate_content(prompt)
        analysis = response.text.strip() if hasattr(response, 'text') else ''
        return analysis if analysis else "Пользователь обращается за помощью."
    except:
        return "Пользователь обращается за помощью."

def _create_smart_summary(analysis_result: Dict, model) -> str:
    """Создает умное резюме анализа"""
    try:
        raw_analysis = analysis_result.get('raw_analysis', '')
        if not raw_analysis:
            return None
        
        prompt = f"""Создай краткое резюме (3-5 пунктов) на основе анализа:

{raw_analysis[:800]}

Выдели самые важные моменты:"""
        
        response = model.generate_content(prompt)
        summary = response.text.strip() if hasattr(response, 'text') else ''
        return summary if summary and len(summary) > 50 else None
    except:
        return None

# Новый структурированный чат-бот для бизнес-аналитика
@app.route('/api/ai-business-analyst/chat', methods=['POST'])
def api_ba_chat():
    """Структурированный чат-бот для диалога с бизнес-аналитиком"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        message = data.get('message', '')
        chat_history = data.get('chat_history', [])
        context = data.get('context', {})
        session_id = data.get('session_id') or f"{session.get('username', 'user')}_{int(time.time())}"
        use_structured_dialog = data.get('use_structured_dialog', True)  # По умолчанию включен структурированный диалог
        
        if not message:
            return jsonify({'error': 'Сообщение не предоставлено'}), 400
        
        # Инициализируем или восстанавливаем структурированный диалог
        if use_structured_dialog:
            if session_id not in active_dialogs:
                active_dialogs[session_id] = StructuredDialogManager(session_id, model=model)
            else:
                # Обновляем модель для умного анализа, если она еще не установлена
                if not active_dialogs[session_id].smart_analyzer:
                    try:
                        from business_analyst_models.smart_answer_analyzer import SmartAnswerAnalyzer
                        active_dialogs[session_id].smart_analyzer = SmartAnswerAnalyzer(model)
                    except:
                        pass
            
            dialog_manager = active_dialogs[session_id]
            
            # Обрабатываем ответ пользователя с использованием AI модели для умного анализа
            response_text, is_complete = dialog_manager.process_answer(message, model=model)
            
            # Если диалог завершен, автоматически генерируем документы
            auto_generated = None
            if is_complete:
                # Генерируем резюме перед автоматической генерацией
                if dialog_manager.current_stage == DialogStage.REVIEW:
                    response_text = dialog_manager.generate_summary()
                elif dialog_manager.current_stage == DialogStage.COMPLETE:
                    # Автоматически запускаем генерацию документов
                    start_time = time.time()
                    collected_data = dialog_manager.get_collected_data_for_brd()
                    
                    # Генерируем BRD
                    brd_result = document_generator.generate_brd(collected_data, context)
                    
                    # Генерируем лидирующие индикаторы
                    leading_indicators = document_generator.generate_leading_indicators(
                        collected_data, 
                        collected_data.get('kpi', [])
                    )
                    
                    duration = time.time() - start_time
                    
                    # Отслеживаем метрики
                    ba_metrics_tracker.track_brd_generation(duration, success=(brd_result.get('content') is not None))
                    ba_metrics_tracker.track_dialog_completion(
                        time.time() - dialog_manager.total_start_time,
                        completed_without_help=True
                    )
                    
                    auto_generated = {
                        'brd': brd_result,
                        'leading_indicators': leading_indicators,
                        'generation_time_seconds': duration
                    }
                    
                    response_text += "\n\n✅ **Диалог завершен!** Все документы автоматически сгенерированы."
            
            # Получаем информацию о текущем этапе
            dialog_stats = dialog_manager.get_dialog_stats()
            
            # Генерируем предложения вариантов ответа через AI
            answer_suggestions = []
            if autocomplete_helper and not is_complete:
                try:
                    context_for_suggestions = {
                        'collected_data': dialog_manager.get_collected_data_for_brd(),
                        'current_stage': dialog_manager.current_stage.value,
                        'conversation_history': dialog_manager.conversation_history[-3:]
                    }
                    suggestions = autocomplete_helper.suggest_answer_options(
                        response_text, 
                        context_for_suggestions
                    )
                    answer_suggestions = suggestions[:3]  # Максимум 3 предложения
                except:
                    pass
            
            return jsonify({
                'response': response_text,
                'success': True,
                'session_id': session_id,
                'dialog_complete': is_complete,
                'current_stage': dialog_manager.current_stage.value,
                'dialog_stats': dialog_stats,
                'auto_generated': auto_generated,
                'answer_suggestions': answer_suggestions
            })
        
        # Умный режим чата с глубоким AI-анализом
        else:
            # Используем умный анализатор для понимания намерений
            try:
                from business_analyst_models.smart_answer_analyzer import SmartAnswerAnalyzer
                smart_analyzer = SmartAnswerAnalyzer(model)
                
                # Анализируем намерение пользователя
                intent_analysis = smart_analyzer.understand_intent(message, {
                    'collected_data': {},
                    'conversation_history': chat_history
                })
                
                # Формируем расширенный контекст с умным анализом
                conversation_context = ""
                if chat_history:
                    conversation_context = "\n\nКОНТЕКСТ ПРЕДЫДУЩЕГО РАЗГОВОРА:\n"
                    for msg in chat_history[-10:]:  # Последние 10 сообщений для лучшего контекста
                        role = msg.get('role', 'user')
                        content = msg.get('content', '')
                        conversation_context += f"{'👤 Пользователь' if role == 'user' else '🤖 Ассистент'}: {content}\n"
                
                # Извлекаем ключевую информацию из истории
                extracted_info = _extract_context_from_history(chat_history)
                
                # Умный системный промпт
                system_prompt = f"""Ты умный AI-Business Analyst, работающий в банке. Ты профессиональный помощник с глубоким пониманием бизнес-процессов.

ТВОЯ РОЛЬ:
- 🔍 Анализировать бизнес-ситуации и документы
- 📋 Извлекать и структурировать требования
- 🎯 Находить пробелы и противоречия
- 📄 Формировать профессиональную документацию
- 🤝 Помогать сотрудникам в любых вопросах бизнес-аналитики
- 💡 Предлагать улучшения и оптимизации

ТВОЙ СТИЛЬ:
- Дружелюбный и профессиональный
- Структурированный и четкий
- Понимающий контекст
- Адаптивный к стилю пользователя
- Проактивный (предлагай следующие шаги)

АНАЛИЗ СООБЩЕНИЯ:
Намерение пользователя: {intent_analysis.get('summary', 'не определено')}
Релевантность: {'Да' if intent_analysis.get('is_relevant', True) else 'Нет'}

{extracted_info}

КОНТЕКСТ ПРОЕКТА:
"""
                
                # Добавляем контекст пользователя
                user_context = ""
                if context:
                    if context.get('customer'):
                        user_context += f"- Заказчик: {context['customer']}\n"
                    if context.get('purpose'):
                        user_context += f"- Цель: {context['purpose']}\n"
                    if context.get('for_confluence'):
                        user_context += "- Результат для Confluence\n"
                
                # Умный анализ сообщения
                message_analysis = _analyze_message_intelligently(message, model, chat_history)
                
                # Формируем умный промпт
                prompt = f"""{system_prompt}{user_context}{conversation_context}

ТЕКУЩЕЕ СООБЩЕНИЕ ПОЛЬЗОВАТЕЛЯ:
"{message}"

АНАЛИЗ СООБЩЕНИЯ:
{message_analysis}

ЗАДАЧА:
1. Понять, что хочет пользователь (даже если формулировка неформальная)
2. Предоставить полезный, релевантный ответ
3. Если нужно - задать уточняющие вопросы
4. Предложить следующие шаги, если уместно
5. Использовать контекст предыдущих сообщений

ВАЖНО:
- Будь умным и понимающим
- Адаптируйся к стилю пользователя
- Если пользователь дает неформальный ответ - пойми его намерение
- Если информация неполная - задай умные уточняющие вопросы
- Предлагай конкретные действия и следующие шаги

СФОРМИРУЙ ОТВЕТ:"""
                
                # Генерируем умный ответ
                response = model.generate_content(prompt)
                response_text = response.text if hasattr(response, 'text') else str(response)
                
                # Если это похоже на документ или проект, проводим глубокий анализ
                if len(message) > 100:
                    try:
                        # Проверяем, нужно ли провести анализ документа
                        if any(keyword in message.lower() for keyword in ['проект', 'система', 'требование', 'задача', 'задача', 'документ']):
                            # Используем документ-анализатор для извлечения структуры
                            analysis_result = document_analyzer.analyze_document(message[:2000])  # Первые 2000 символов
                            if analysis_result and not analysis_result.get('error'):
                                # Формируем умное резюме анализа
                                summary = _create_smart_summary(analysis_result, model)
                                if summary:
                                    response_text += f"\n\n---\n\n## 💡 Ключевые выводы:\n\n{summary}"
                    except Exception as e:
                        logger.error(f"Ошибка при анализе: {e}")
                
                return jsonify({
                    'response': response_text,
                    'success': True,
                    'intent': intent_analysis.get('intent', 'answer')
                })
                
            except Exception as e:
                logger.error(f"Ошибка в умном режиме: {e}")
                # Fallback на простой режим
                pass
            
            # Простой режим (fallback)
            system_prompt = """Ты AI-Business Analyst, работающий в банке. Отвечай умно, профессионально и полезно."""
            
            conversation_context = ""
            if chat_history:
                conversation_context = "\n\nИстория:\n"
                for msg in chat_history[-5:]:
                    role = msg.get('role', 'user')
                    content = msg.get('content', '')
                    conversation_context += f"{'Пользователь' if role == 'user' else 'Ассистент'}: {content}\n"
            
            prompt = f"""{system_prompt}{conversation_context}

Сообщение пользователя: {message}

Ответь умно и полезно:"""
            
            response = model.generate_content(prompt)
            response_text = response.text if hasattr(response, 'text') else str(response)
            
            return jsonify({
                'response': response_text,
                'success': True
            })
    
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# Метрики AI Business Analyst
@app.route('/api/ai-business-analyst/metrics', methods=['GET'])
def api_ba_metrics():
    """Получить все метрики AI Business Analyst"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        metrics_type = request.args.get('type', 'all')  # all, performance, business, usability
        
        if metrics_type == 'performance':
            metrics = ba_metrics_tracker.get_performance_metrics()
        elif metrics_type == 'business':
            metrics = ba_metrics_tracker.get_business_metrics()
        elif metrics_type == 'usability':
            metrics = ba_metrics_tracker.get_usability_metrics()
        else:
            metrics = ba_metrics_tracker.get_all_metrics()
        
        return jsonify({
            'success': True,
            'metrics': metrics,
            'metrics_type': metrics_type
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Генерация лидирующих индикаторов
@app.route('/api/ai-business-analyst/generate-leading-indicators', methods=['POST'])
def api_ba_generate_leading_indicators():
    """Генерировать лидирующие индикаторы"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        project_data = data.get('project_data', {})
        kpi_data = data.get('kpi', [])
        
        start_time = time.time()
        indicators = document_generator.generate_leading_indicators(project_data, kpi_data)
        duration = time.time() - start_time
        
        return jsonify({
            'success': indicators.get('content') is not None,
            'leading_indicators': indicators,
            'generation_time_seconds': duration
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Анализ и совершенствование процессов
@app.route('/api/ai-business-analyst/analyze-process', methods=['POST'])
def api_ba_analyze_process():
    """Анализ бизнес-процесса"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        process_description = data.get('process_description', '')
        current_metrics = data.get('current_metrics')
        
        if not process_description:
            return jsonify({'error': 'Описание процесса не предоставлено'}), 400
        
        analysis = process_improver.analyze_process(process_description, current_metrics)
        
        return jsonify({
            'success': analysis.get('content') is not None,
            'analysis': analysis
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-business-analyst/suggest-improvements', methods=['POST'])
def api_ba_suggest_improvements():
    """Предложить улучшения процесса"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        analysis = data.get('analysis', {})
        priority = data.get('priority', 'high')
        
        improvements = process_improver.suggest_improvements(analysis, priority)
        
        return jsonify({
            'success': True,
            'improvements': improvements,
            'count': len(improvements)
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Автоматическая генерация всех документов после диалога
@app.route('/api/ai-business-analyst/auto-generate-all', methods=['POST'])
def api_ba_auto_generate_all():
    """Автоматически сгенерировать все документы на основе собранных данных"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        session_id = data.get('session_id')
        collected_data = data.get('collected_data', {})
        context = data.get('context', {})
        
        if not session_id and not collected_data:
            return jsonify({'error': 'Недостаточно данных для генерации'}), 400
        
        # Получаем данные из диалога или из запроса
        if session_id and session_id in active_dialogs:
            dialog_manager = active_dialogs[session_id]
            collected_data = dialog_manager.get_collected_data_for_brd()
        
        start_time = time.time()
        
        # Генерируем все документы
        documents = {}
        
        # 1. BRD
        brd = document_generator.generate_brd(collected_data, context)
        documents['brd'] = brd
        
        # 2. Лидирующие индикаторы
        leading_indicators = document_generator.generate_leading_indicators(
            collected_data,
            collected_data.get('kpi', [])
        )
        documents['leading_indicators'] = leading_indicators
        
        # 3. Use Cases (если есть требования)
        if collected_data.get('scope'):
            use_cases = []
            for scope_item in collected_data['scope'][:3]:  # Первые 3 элемента
                use_case = document_generator.generate_use_case(
                    scope_item,
                    'Пользователь',
                    context
                )
                use_cases.append(use_case)
            documents['use_cases'] = use_cases
        
        # 4. User Stories
        requirements = collected_data.get('scope', [])
        if requirements:
            user_stories = document_generator.generate_user_stories(requirements, [])
            documents['user_stories'] = user_stories
        
        # 5. BPMN (генерируется через существующий endpoint логику)
        # Здесь можно добавить автоматическую генерацию BPMN
        
        duration = time.time() - start_time
        
        # Отслеживаем метрики
        ba_metrics_tracker.track_brd_generation(duration, success=True)
        
        return jsonify({
            'success': True,
            'documents': documents,
            'generation_time_seconds': duration,
            'documents_count': len(documents)
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# Автозаполнение через AI
@app.route('/api/ai-business-analyst/autocomplete/suggest-answers', methods=['POST'])
def api_ba_suggest_answers():
    """Предложить варианты ответов на вопрос"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not autocomplete_helper:
        return jsonify({'error': 'Autocomplete helper not available'}), 500
    
    try:
        data = request.json
        question = data.get('question', '')
        context = data.get('context', {})
        
        if not question:
            return jsonify({'error': 'Вопрос не предоставлен'}), 400
        
        suggestions = autocomplete_helper.suggest_answer_options(question, context)
        
        return jsonify({
            'success': True,
            'suggestions': suggestions,
            'count': len(suggestions)
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-business-analyst/autocomplete/fill-from-chat', methods=['POST'])
def api_ba_fill_from_chat():
    """Автоматически заполнить поля BRD из истории чата"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not autocomplete_helper:
        return jsonify({'error': 'Autocomplete helper not available'}), 500
    
    try:
        data = request.json
        chat_history = data.get('chat_history', [])
        collected_data = data.get('collected_data', {})
        session_id = data.get('session_id')
        
        # Если есть session_id, получаем данные из диалога
        if session_id and session_id in active_dialogs:
            dialog_manager = active_dialogs[session_id]
            collected_data = dialog_manager.get_collected_data_for_brd()
            # Формируем историю из диалога
            if not chat_history:
                chat_history = []
                for conv in dialog_manager.conversation_history:
                    chat_history.append({
                        'role': 'user',
                        'content': conv.get('user_answer', '')
                    })
        
        filled_data = autocomplete_helper.autofill_fields_from_chat(chat_history, collected_data)
        
        return jsonify({
            'success': True,
            'filled_data': filled_data
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-business-analyst/autocomplete/field', methods=['POST'])
def api_ba_autocomplete_field():
    """Автозаполнение конкретного поля"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not autocomplete_helper:
        return jsonify({'error': 'Autocomplete helper not available'}), 500
    
    try:
        data = request.json
        field_name = data.get('field_name', '')
        partial_value = data.get('partial_value', '')
        context = data.get('context', {})
        
        if not field_name:
            return jsonify({'error': 'Название поля не предоставлено'}), 400
        
        suggestions = autocomplete_helper.autofill_field(field_name, partial_value, context)
        
        return jsonify({
            'success': True,
            'suggestions': suggestions
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-business-analyst/autocomplete/suggest-scope', methods=['POST'])
def api_ba_suggest_scope():
    """Предложить элементы scope"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not autocomplete_helper:
        return jsonify({'error': 'Autocomplete helper not available'}), 500
    
    try:
        data = request.json
        project_description = data.get('project_description', '')
        existing_scope = data.get('existing_scope', [])
        
        if not project_description:
            # Пытаемся получить из контекста или истории
            session_id = data.get('session_id')
            if session_id and session_id in active_dialogs:
                dialog_manager = active_dialogs[session_id]
                collected = dialog_manager.get_collected_data_for_brd()
                project_description = collected.get('project_description', '') or collected.get('project_goal', '')
        
        suggestions = autocomplete_helper.suggest_scope_items(project_description, existing_scope)
        
        return jsonify({
            'success': True,
            'suggestions': suggestions,
            'count': len(suggestions)
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-business-analyst/autocomplete/suggest-kpi', methods=['POST'])
def api_ba_suggest_kpi():
    """Предложить KPI метрики"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not autocomplete_helper:
        return jsonify({'error': 'Autocomplete helper not available'}), 500
    
    try:
        data = request.json
        project_description = data.get('project_description', '')
        existing_kpi = data.get('existing_kpi', [])
        
        if not project_description:
            session_id = data.get('session_id')
            if session_id and session_id in active_dialogs:
                dialog_manager = active_dialogs[session_id]
                collected = dialog_manager.get_collected_data_for_brd()
                project_description = collected.get('project_description', '') or collected.get('project_goal', '')
        
        suggestions = autocomplete_helper.suggest_kpi_metrics(project_description, existing_kpi)
        
        return jsonify({
            'success': True,
            'suggestions': suggestions,
            'count': len(suggestions)
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-business-analyst/autocomplete/fill-from-text', methods=['POST'])
def api_ba_fill_from_text():
    """Умное заполнение всех полей из текста"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    if not autocomplete_helper:
        return jsonify({'error': 'Autocomplete helper not available'}), 500
    
    try:
        data = request.json
        text = data.get('text', '')
        
        if not text or len(text.strip()) < 10:
            return jsonify({'error': 'Текст слишком короткий (минимум 10 символов)'}), 400
        
        filled_data = autocomplete_helper.smart_fill_from_text(text)
        
        return jsonify({
            'success': True,
            'filled_data': filled_data
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Старый endpoint для обратной совместимости
@app.route('/api/ai-business-analyst', methods=['POST'])
def api_ai_business_analyst():
    """Старый endpoint, перенаправляет на комплексный анализ"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    data = request.json
    query = data.get('query', '')
    
    if not query:
        return jsonify({'error': 'Запрос не предоставлен'}), 400
    
    # Используем новый функционал анализа
    try:
        analysis = document_analyzer.analyze_document(query)
        return jsonify({
            'result': analysis.get('raw_analysis', 'Анализ выполнен'),
            'analysis': analysis,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ==================== AI-Code Review Assistant API Endpoints ====================

# Загрузка архитектуры проекта
@app.route('/api/ai-code-review/upload-architecture', methods=['POST'])
def api_code_review_upload_architecture():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'Файл не предоставлен'}), 400
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': 'Файл не выбран'}), 400
        
        # Сохраняем файл временно
        filename = secure_filename(file.filename)
        upload_dir = 'uploads/architecture'
        os.makedirs(upload_dir, exist_ok=True)
        file_path = os.path.join(upload_dir, filename)
        file.save(file_path)
        
        # Определяем тип файла
        file_type = request.form.get('file_type') or filename.rsplit('.', 1)[1].lower() if '.' in filename else 'txt'
        
        # Загружаем архитектуру
        architecture_data = architecture_loader.load_architecture(file_path, file_type)
        
        # Анализируем архитектуру
        result = {
            'file_path': architecture_data.get('file_path'),
            'file_type': architecture_data.get('file_type'),
            'success': architecture_data.get('success', False)
        }
        
        if architecture_data.get('error'):
            result['error'] = architecture_data.get('error')
        elif architecture_data.get('success'):
            try:
                analysis = architecture_loader.analyze_architecture(architecture_data)
                if analysis.get('error'):
                    result['analysis_error'] = analysis.get('error')
                elif analysis.get('analysis'):
                    result['analysis'] = analysis.get('analysis')
                else:
                    result['analysis_error'] = 'Анализ выполнен, но результат пуст'
            except Exception as e:
                result['analysis_error'] = f'Ошибка при анализе: {str(e)}'
        
        # Удаляем временный файл
        try:
            os.remove(file_path)
        except:
            pass
        
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Анализ кода (отдельный файл)
# Генерация кода через AI
@app.route('/api/ai-code-review/generate-code', methods=['POST'])
def api_code_review_generate_code():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        description = data.get('description', '')
        language = data.get('language', 'Python')
        file_path = data.get('file_path')
        
        if not description:
            return jsonify({'error': 'Описание обязательно'}), 400
        
        result = code_analyzer.generate_code_from_description(description, language, file_path)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Генерация diff через AI
@app.route('/api/ai-code-review/generate-diff', methods=['POST'])
def api_code_review_generate_diff():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        description = data.get('description', '')
        base_code = data.get('base_code')
        
        if not description:
            return jsonify({'error': 'Описание изменений обязательно'}), 400
        
        result = code_analyzer.generate_diff_from_description(description, base_code)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Улучшение описания проблемы через AI
@app.route('/api/ai-code-review/improve-issue-description', methods=['POST'])
def api_code_review_improve_issue_description():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        description = data.get('description', '')
        code_snippet = data.get('code_snippet')
        
        if not description:
            return jsonify({'error': 'Описание проблемы обязательно'}), 400
        
        result = code_analyzer.improve_issue_description(description, code_snippet)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai-code-review/analyze-code', methods=['POST'])
def api_code_review_analyze_code():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        code = data.get('code', '')
        language = data.get('language', '')
        file_path = data.get('file_path')
        
        if not code:
            return jsonify({'error': 'Код не предоставлен'}), 400
        
        # Автоматически определяем язык, если не указан или указан как "Auto"/"Unknown"
        if not language or language in ['Auto', 'Unknown', 'Другой']:
            detected_language = language_detector.detect_language(code, file_path)
            language = detected_language if detected_language != 'Unknown' else 'Python'  # Fallback на Python
        else:
            # Если язык указан, проверяем его корректность
            detected_language = language_detector.detect_language(code, file_path)
            # Если определенный язык отличается от указанного, используем определенный
            if detected_language != 'Unknown' and detected_language != language:
                language = detected_language  # Используем автоматически определенный
        
        analysis = code_analyzer.analyze_code(code, language, file_path)
        analysis['type'] = 'code'
        analysis['detected_language'] = language  # Добавляем определенный язык в ответ
        
        # Генерируем рекомендации
        if analysis.get('success'):
            recommendations = review_generator.generate_recommendations(analysis, 'file')
            analysis['recommendations'] = recommendations
        
        return jsonify(analysis)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Анализ diff
@app.route('/api/ai-code-review/analyze-diff', methods=['POST'])
def api_code_review_analyze_diff():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        diff_text = data.get('diff', '')
        base_code = data.get('base_code')
        new_code = data.get('new_code')
        description = data.get('description')
        
        if not diff_text:
            return jsonify({'error': 'Diff не предоставлен'}), 400
        
        analysis = code_analyzer.analyze_diff(diff_text, base_code, new_code, description)
        analysis['type'] = 'diff'
        
        return jsonify(analysis)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Анализ всего проекта
@app.route('/api/ai-code-review/analyze-project', methods=['POST'])
def api_code_review_analyze_project():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        files = data.get('files', [])  # [{'path': '...', 'code': '...', 'language': '...'}]
        architecture = data.get('architecture')
        
        if not files:
            return jsonify({'error': 'Файлы не предоставлены'}), 400
        
        # Автоматически определяем язык для каждого файла, если не указан
        for file_info in files:
            if not file_info.get('language') or file_info.get('language') in ['Auto', 'Unknown', 'Другой']:
                detected_language = language_detector.detect_language(
                    file_info.get('code', ''),
                    file_info.get('path')
                )
                file_info['language'] = detected_language if detected_language != 'Unknown' else 'Unknown'
        
        analysis = code_analyzer.analyze_project(files, architecture)
        analysis['type'] = 'project'
        analysis['files_with_languages'] = [
            {'path': f.get('path', 'unknown'), 'language': f.get('language', 'unknown')}
            for f in files
        ]
        
        # Генерируем рекомендации
        if analysis.get('success'):
            recommendations = review_generator.generate_recommendations(analysis, 'project')
            analysis['recommendations'] = recommendations
        
        return jsonify(analysis)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Генерация итогового отчёта
@app.route('/api/ai-code-review/generate-report', methods=['POST'])
def api_code_review_generate_report():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        analyses = data.get('analyses', [])
        merge_recommendation = data.get('merge_recommendation')
        mr_info = data.get('mr_info')
        
        if not analyses:
            return jsonify({'error': 'Анализы не предоставлены'}), 400
        
        report = code_review_report_generator.generate_report(
            analyses, merge_recommendation, mr_info
        )
        
        return jsonify(report)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Обновление статуса MR
@app.route('/api/ai-code-review/mr/update-status', methods=['POST'])
def api_code_review_update_mr_status():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        mr_id = data.get('mr_id')
        status = data.get('status')
        labels = data.get('labels', [])
        review_data = data.get('review_data')
        
        if not mr_id or not status:
            return jsonify({'error': 'mr_id и status обязательны'}), 400
        
        result = gitlab_integration.update_mr_status(mr_id, status, labels, review_data)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Получение статуса MR
@app.route('/api/ai-code-review/mr/<mr_id>', methods=['GET'])
def api_code_review_get_mr(mr_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        result = gitlab_integration.get_mr_status(mr_id)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Список MR
@app.route('/api/ai-code-review/mr/list', methods=['GET'])
def api_code_review_list_mrs():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        status_filter = request.args.get('status')
        mrs = gitlab_integration.list_mrs(status_filter)
        return jsonify({'mrs': mrs, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Поиск проектов GitLab
@app.route('/api/ai-code-review/gitlab/projects', methods=['GET'])
def api_code_review_gitlab_projects():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        search = request.args.get('search', '')
        owned = request.args.get('owned', 'false').lower() == 'true'
        membership = request.args.get('membership', 'true').lower() == 'true'
        visibility = request.args.get('visibility')
        
        result = gitlab_integration.list_projects(
            search=search if search else None,
            owned=owned,
            membership=membership,
            visibility=visibility
        )
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Полный анализ Merge Request
@app.route('/api/ai-code-review/mr/<project_id>/<mr_id>/analyze', methods=['POST'])
def api_code_review_analyze_mr(project_id, mr_id):
    """
    Полный анализ Merge Request:
    1. Получает информацию о MR
    2. Получает diff изменений
    3. Анализирует изменения
    4. Формирует резюме с рекомендацией
    5. Опционально оставляет комментарии в GitLab
    """
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json or {}
        auto_comment = data.get('auto_comment', False)  # Автоматически оставить комментарий в GitLab
        
        # 1. Получаем информацию о MR
        mr_info = gitlab_integration.get_merge_request(project_id, mr_id)
        if 'error' in mr_info:
            return jsonify(mr_info), 400
        
        # 2. Получаем diff
        diff_result = gitlab_integration.get_merge_request_diff(project_id, mr_id)
        if 'error' in diff_result:
            return jsonify(diff_result), 400
        
        diffs = diff_result.get('diffs', [])
        if not diffs:
            return jsonify({'error': 'Diff не найден для данного MR'}), 400
        
        # Формируем текст diff
        diff_text_parts = []
        for diff_item in diffs:
            diff_text_parts.append(f"Файл: {diff_item.get('new_path', diff_item.get('old_path', 'unknown'))}")
            diff_text_parts.append(f"---")
            diff_text_parts.append(diff_item.get('diff', ''))
            diff_text_parts.append("")
        
        diff_text = '\n'.join(diff_text_parts)
        
        # 3. Анализируем diff
        description = mr_info.get('description', '')
        analysis = code_analyzer.analyze_diff(
            diff_text=diff_text,
            description=description if description else None,
            mr_id=f"{project_id}/{mr_id}"
        )
        
        if analysis.get('error'):
            return jsonify(analysis), 500
        
        # 4. Генерируем резюме
        merge_recommendation = analysis.get('merge_recommendation', 'needs-review')
        
        # Формируем краткое резюме для комментария
        summary_comment = f"""## 🤖 AI Code Review

**Статус:** {merge_recommendation.upper()}

**Краткое резюме:**
{analysis.get('analysis', '')[:500]}...

---
*Автоматический анализ выполнен AI Code Review Assistant*
"""
        
        # 5. Обновляем локальный статус MR
        labels = []
        if merge_recommendation == 'ready-for-merge':
            labels = ['ready-for-merge', 'ai-reviewed']
        elif merge_recommendation == 'needs-fixes':
            labels = ['needs-fixes', 'ai-reviewed']
        elif merge_recommendation == 'reject':
            labels = ['reject', 'ai-reviewed']
        else:
            labels = ['needs-review', 'ai-reviewed']
        
        gitlab_integration.update_mr_status(
            mr_id=f"{project_id}/{mr_id}",
            status=merge_recommendation,
            labels=labels,
            review_data={
                'analysis': analysis.get('analysis', ''),
                'merge_recommendation': merge_recommendation,
                'analyzed_at': datetime.now().isoformat()
            }
        )
        
        result = {
            'mr_info': mr_info,
            'analysis': analysis,
            'merge_recommendation': merge_recommendation,
            'summary_comment': summary_comment,
            'labels': labels,
            'success': True
        }
        
        # 6. Опционально оставляем комментарий в GitLab
        if auto_comment:
            try:
                comment_result = gitlab_integration.add_mr_discussion(
                    project_id=project_id,
                    mr_id=mr_id,
                    body=summary_comment
                )
                result['comment_posted'] = comment_result.get('success', False)
                result['comment_id'] = comment_result.get('discussion', {}).get('id')
            except Exception as e:
                result['comment_error'] = str(e)
        
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': f'Ошибка при анализе MR: {str(e)}'}), 500

# Получить информацию о MR с diff
@app.route('/api/ai-code-review/mr/<project_id>/<mr_id>', methods=['GET'])
def api_code_review_get_mr_full(project_id, mr_id):
    """Получает полную информацию о MR включая diff"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        # Получаем информацию о MR
        mr_info = gitlab_integration.get_merge_request(project_id, mr_id)
        if 'error' in mr_info:
            return jsonify(mr_info), 400
        
        # Получаем diff
        diff_result = gitlab_integration.get_merge_request_diff(project_id, mr_id)
        if 'error' in diff_result:
            return jsonify(diff_result), 400
        
        # Получаем локальный статус
        local_status = gitlab_integration.get_mr_status(f"{project_id}/{mr_id}")
        
        return jsonify({
            'mr_info': mr_info,
            'diff': diff_result,
            'local_status': local_status if 'error' not in local_status else None,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Добавить комментарий к MR в GitLab
@app.route('/api/ai-code-review/mr/<project_id>/<mr_id>/comment', methods=['POST'])
def api_code_review_add_mr_comment(project_id, mr_id):
    """Добавляет комментарий к MR в GitLab"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        comment = data.get('comment', '')
        
        if not comment:
            return jsonify({'error': 'Комментарий обязателен'}), 400
        
        result = gitlab_integration.add_mr_discussion(
            project_id=project_id,
            mr_id=mr_id,
            body=comment
        )
        
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Список MR для проекта
@app.route('/api/ai-code-review/gitlab/projects/<project_id>/merge-requests', methods=['GET'])
def api_code_review_project_mrs(project_id):
    """Получает список MR для конкретного проекта"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        state = request.args.get('state', 'opened')
        result = gitlab_integration.list_merge_requests(project_id=project_id, state=state)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Обучающие подсказки
@app.route('/api/ai-code-review/educational/hints', methods=['POST'])
def api_code_review_educational_hints():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        issue = data.get('issue', '')
        developer_level = data.get('developer_level', 'junior')
        
        if not issue:
            return jsonify({'error': 'Проблема не указана'}), 400
        
        hints = review_generator.generate_educational_hints(issue, developer_level)
        return jsonify(hints)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Объяснение проблемы
@app.route('/api/ai-code-review/educational/explain', methods=['POST'])
def api_code_review_educational_explain():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        issue = data.get('issue', '')
        code_snippet = data.get('code_snippet', '')
        developer_level = data.get('developer_level', 'junior')
        
        if not issue:
            return jsonify({'error': 'Проблема не указана'}), 400
        
        explanation = educational_support.generate_explanation(issue, code_snippet, developer_level)
        return jsonify(explanation)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Статистика разработчика
@app.route('/api/ai-code-review/educational/stats/<developer_id>', methods=['GET'])
def api_code_review_educational_stats(developer_id):
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        stats = educational_support.get_developer_stats(developer_id)
        return jsonify(stats)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Отслеживание прогресса
@app.route('/api/ai-code-review/educational/track-progress', methods=['POST'])
def api_code_review_educational_track_progress():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        developer_id = data.get('developer_id')
        issue_type = data.get('issue_type', 'unknown')
        resolved = data.get('resolved', False)
        
        if not developer_id:
            return jsonify({'error': 'developer_id обязателен'}), 400
        
        result = educational_support.track_developer_progress(developer_id, issue_type, resolved)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Частые проблемы
@app.route('/api/ai-code-review/educational/common-issues', methods=['GET'])
def api_code_review_educational_common_issues():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        limit = request.args.get('limit', 10, type=int)
        issues = educational_support.get_common_issues(limit)
        return jsonify({'issues': issues, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Автоматический анализ и исправление файла
@app.route('/api/ai-code-review/auto-fix-file', methods=['POST'])
def api_code_review_auto_fix_file():
    """
    Автоматически анализирует файл, исправляет ошибки и создаёт коммит
    """
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        project_id = data.get('project_id')
        file_path = data.get('file_path')
        branch = data.get('branch', 'main')
        auto_commit = data.get('auto_commit', True)
        auto_apply_fixes = data.get('auto_apply_fixes', True)
        
        if not project_id or not file_path:
            return jsonify({'error': 'project_id и file_path обязательны'}), 400
        
        # 1. Получаем текущий файл
        file_content_result = gitlab_integration.get_file_content(project_id, file_path, branch)
        if 'error' in file_content_result:
            return jsonify(file_content_result), 400
        
        original_code = file_content_result.get('content', '')
        if not original_code:
            return jsonify({'error': 'Файл пуст или не найден'}), 400
        
        # 2. Определяем язык
        detected_language = language_detector.detect_language(original_code, file_path)
        language = detected_language if detected_language != 'Unknown' else 'Python'
        
        # 3. Анализируем код
        analysis = code_analyzer.analyze_code(original_code, language, file_path)
        if analysis.get('error'):
            return jsonify(analysis), 500
        
        # 4. Генерируем исправления
        fixed_code_result = review_generator.generate_code_fixes(
            analysis=analysis,
            original_code=original_code,
            file_path=file_path
        )
        
        if fixed_code_result.get('error'):
            return jsonify(fixed_code_result), 500
        
        fixed_code = fixed_code_result.get('fixed_code', original_code)
        changes = fixed_code_result.get('changes', [])
        summary = fixed_code_result.get('summary', '')
        
        result = {
            'original_code': original_code,
            'fixed_code': fixed_code,
            'analysis': analysis,
            'changes': changes,
            'summary': summary,
            'has_changes': fixed_code != original_code,
            'success': True
        }
        
        # 5. Автоматически применяем исправления и создаём коммит
        if auto_apply_fixes and fixed_code != original_code:
            # Генерируем имя и описание коммита
            commit_info = review_generator.generate_commit_message(
                file_path=file_path,
                changes=changes,
                summary=summary
            )
            
            commit_title = commit_info.get('title', f'AI: Исправления в {file_path.split("/")[-1]}')
            commit_description = commit_info.get('description', summary or 'Автоматические исправления кода')
            commit_message = f"{commit_title}\n\n{commit_description}"
            
            if auto_commit:
                # Обновляем файл в GitLab
                update_result = gitlab_integration.update_file(
                    project_id=project_id,
                    file_path=file_path,
                    content=fixed_code,
                    branch=branch,
                    commit_message=commit_message
                )
                
                if 'error' in update_result:
                    result['commit_error'] = update_result.get('error')
                    result['commit_success'] = False
                else:
                    result['commit_success'] = True
                    result['commit_message'] = commit_message
                    result['commit_info'] = commit_info
            else:
                result['commit_message'] = commit_message
                result['commit_info'] = commit_info
                result['commit_success'] = None  # Не применено
        
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': f'Ошибка при автоматическом исправлении: {str(e)}'}), 500

# Получить файл для редактирования
@app.route('/api/ai-code-review/gitlab/file', methods=['GET'])
def api_code_review_get_file():
    """Получает файл из репозитория для редактирования"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        project_id = request.args.get('project_id')
        file_path = request.args.get('file_path')
        branch = request.args.get('branch', 'main')
        
        if not project_id or not file_path:
            return jsonify({'error': 'project_id и file_path обязательны'}), 400
        
        result = gitlab_integration.get_file_content(project_id, file_path, branch)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Сохранить файл с автоматическим анализом
@app.route('/api/ai-code-review/gitlab/file/save', methods=['POST'])
def api_code_review_save_file():
    """Сохраняет файл с автоматическим анализом и исправлением"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        data = request.json
        project_id = data.get('project_id')
        file_path = data.get('file_path')
        content = data.get('content')
        branch = data.get('branch', 'main')
        auto_fix = data.get('auto_fix', True)
        auto_commit = data.get('auto_commit', True)
        
        if not project_id or not file_path or content is None:
            return jsonify({'error': 'project_id, file_path и content обязательны'}), 400
        
        # Если включено автоисправление, анализируем и исправляем
        if auto_fix:
            # Определяем язык
            detected_language = language_detector.detect_language(content, file_path)
            language = detected_language if detected_language != 'Unknown' else 'Python'
            
            # Анализируем код
            analysis = code_analyzer.analyze_code(content, language, file_path)
            
            if not analysis.get('error'):
                # Генерируем исправления
                fixed_result = review_generator.generate_code_fixes(
                    analysis=analysis,
                    original_code=content,
                    file_path=file_path
                )
                
                if not fixed_result.get('error'):
                    fixed_code = fixed_result.get('fixed_code', content)
                    changes = fixed_result.get('changes', [])
                    summary = fixed_result.get('summary', '')
                    
                    # Если есть исправления, используем исправленный код
                    if fixed_code != content and len(changes) > 0:
                        content = fixed_code
                        
                        # Генерируем коммит сообщение
                        commit_info = review_generator.generate_commit_message(
                            file_path=file_path,
                            changes=changes,
                            summary=summary
                        )
                        commit_title = commit_info.get('title', f'AI: Исправления в {file_path.split("/")[-1]}')
                        commit_description = commit_info.get('description', summary or 'Автоматические исправления кода')
                        commit_message = f"{commit_title}\n\n{commit_description}"
                    else:
                        commit_message = f'Update {file_path}'
                else:
                    commit_message = f'Update {file_path}'
            else:
                commit_message = f'Update {file_path}'
        else:
            commit_message = data.get('commit_message', f'Update {file_path}')
        
        # Сохраняем файл
        if auto_commit:
            result = gitlab_integration.update_file(
                project_id=project_id,
                file_path=file_path,
                content=content,
                branch=branch,
                commit_message=commit_message
            )
            
            if 'error' in result:
                return jsonify(result), 400
            
            return jsonify({
                'success': True,
                'file_path': file_path,
                'commit_message': commit_message,
                'auto_fixed': auto_fix and 'fixed_code' in locals() and fixed_code != data.get('content'),
                'message': 'Файл успешно сохранён' + (' с автоматическими исправлениями' if auto_fix else '')
            })
        else:
            return jsonify({
                'success': True,
                'file_path': file_path,
                'content': content,
                'commit_message': commit_message,
                'auto_fixed': auto_fix and 'fixed_code' in locals() and fixed_code != data.get('content'),
                'message': 'Файл подготовлен к сохранению'
            })
    except Exception as e:
        return jsonify({'error': f'Ошибка при сохранении файла: {str(e)}'}), 500

# Старый endpoint для обратной совместимости
@app.route('/api/ai-code-review', methods=['POST'])
def api_ai_code_review():
    """Старый endpoint, перенаправляет на анализ кода"""
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    data = request.json
    code = data.get('code', '')
    language = data.get('language', 'Python')
    
    if not code:
        return jsonify({'error': 'Код не предоставлен'}), 400
    
    try:
        analysis = code_analyzer.analyze_code(code, language)
        return jsonify({
            'result': analysis.get('analysis', ''),
            'analysis': analysis,
            'success': True
        })
    except Exception as e:
        error_message = str(e)
        if 'API key' in error_message or 'authentication' in error_message.lower():
            return jsonify({'error': 'Ошибка аутентификации с Gemini API. Проверьте ваш API ключ.'}), 500
        return jsonify({'error': f'Ошибка при обработке запроса: {error_message}'}), 500

# ==================== AI Assistant API Endpoints ====================

# Получить приветствие для страницы
@app.route('/api/ai-assistant/greeting', methods=['GET'])
def api_ai_assistant_greeting():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if not ML_MODELS_AVAILABLE or get_assistant is None:
            return jsonify({'error': 'AI Assistant недоступен'}), 503
        
        page_name = request.args.get('page', 'index')
        assistant = get_assistant()
        greeting = assistant.get_greeting(page_name)
        return jsonify({'greeting': greeting, 'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Задать вопрос ассистенту
@app.route('/api/ai-assistant/ask', methods=['POST'])
def api_ai_assistant_ask():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if not ML_MODELS_AVAILABLE or get_assistant is None:
            return jsonify({'error': 'AI Assistant недоступен'}), 503
        
        data = request.json
        page_name = data.get('page', 'index')
        question = data.get('question', '')
        conversation_history = data.get('conversation_history', [])
        user_id = session.get('username', 'default')
        
        if not question:
            return jsonify({'error': 'Вопрос не предоставлен'}), 400
        
        assistant = get_assistant()
        answer = assistant.ask(page_name, question, user_id, conversation_history)
        
        return jsonify({
            'answer': answer,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Получить путеводитель по функции
@app.route('/api/ai-assistant/guide', methods=['GET'])
def api_ai_assistant_guide():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if not ML_MODELS_AVAILABLE or get_assistant is None:
            return jsonify({'error': 'AI Assistant недоступен'}), 503
        
        page_name = request.args.get('page', 'index')
        feature_name = request.args.get('feature')
        
        assistant = get_assistant()
        guide = assistant.get_feature_guide(page_name, feature_name)
        
        return jsonify({
            'guide': guide,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Получить контекст страницы
@app.route('/api/ai-assistant/context', methods=['GET'])
def api_ai_assistant_context():
    if 'username' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    
    try:
        if not ML_MODELS_AVAILABLE or get_assistant is None:
            return jsonify({'error': 'AI Assistant недоступен'}), 503
        
        page_name = request.args.get('page', 'index')
        assistant = get_assistant()
        context = assistant.get_page_context(page_name)
        
        return jsonify({
            'context': context,
            'success': True
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, port=5000)

